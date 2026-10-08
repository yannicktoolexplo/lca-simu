"""Build the offline 2025 comparison from source workbooks and ordinary full runs.

No simulation parameter is changed. ``build_comparison_payload`` reads its inputs
and returns JSON-compatible data; the caller owns publication and provenance.
Days are offsets from 2025-01-01. Mapping keys are article/division and run label.
Compact row contracts are published in ``payload['columns']``. Missing series
remain empty, missing values remain None, and only complete seven-day windows
are aggregated under explicitly selectable conventions. Source MRP plans retain their original Sunday dates and vintage;
their H/I/J/K columns are never confused with executed material movements.
"""
from __future__ import annotations

from collections import defaultdict
import csv
from datetime import date, datetime, timedelta
import hashlib
import json
import math
from pathlib import Path

import openpyxl


ORIGIN = date(2025, 1, 1)
DAYS = 365
ROOT = Path(__file__).resolve().parents[3]
SITES = {'1430': 'M-1430', '1450': 'SDC-1450', '1810': 'M-1810', '1920': 'DC-1920'}
SITE_NAMES = {'1430': 'Gien', '1450': 'Gaillac', '1810': 'Avène', '1920': 'Muret'}
LABELS = {'nominal': 'Nominal conservé',
          'reference_datee': 'Référence datée · stock global conservé',
          'usages_complementaires': 'Autres usages estimés · stock global conservé',
          'reference_appariee': 'Règles sources · référence à aléas appariés',
          'receptions_datees': 'Réceptions retenues à leur échéance',
          'receptions_et_fabrication': 'Réceptions datées + sécurité de fabrication',
          'regles_sources': 'Règles sources · sécurité et lots',
          'regles_disponibilite': 'Règles sources et disponibilité datée',
          'retrait_statique_338929': '338929 : retrait du forçage statique',
          'dynamique_338929': '338929 : mode dynamique explicite'}
TRACE_FIELDS = ('bb_demand_signal_qty', 'bb_demand_signal_raw_qty',
                'target_stock_qty', 'safety_floor_qty', 'bn_qty',
                'recv_prev_future_qty', 'coverage_target_qty')
OPTIONAL_TRACE_FIELDS = (
    'recv_prev_within_cover_qty', 'recv_prev_after_cover_qty', 'receipt_cover_end_day',
    'mrp_receipt_netting_mode', 'mrp_receipt_netting_scope',
    'decision_receipt_cover_end_day', 'decision_recv_prev_future_qty',
    'decision_recv_prev_within_cover_qty', 'decision_recv_prev_after_cover_qty',
    'decision_stock_proj_qty', 'decision_target_with_backlog_qty',
    'decision_held_release_within_cover_qty', 'decision_bn_qty', 'decision_controlled_bn_qty',
    'production_mrp_safety_enabled', 'production_target_signal_qty',
    'production_historical_target_qty', 'production_safety_floor_qty',
    'production_applied_target_qty', 'production_safety_time_days',
    'production_safety_stock_qty', 'production_target_multiplier', 'production_stock_position_qty',
    'mrp_execution_mode', 'recv_prev_undelivered_qty', 'held_pending_availability_qty',
    'available_release_today_qty',
    'dated_action_scope', 'dated_requirement_qty', 'dated_campaign_remaining_qty', 'dated_reserve_requirement_qty',
    'dated_stock_qty', 'dated_firm_held_qty', 'dated_firm_transit_qty', 'dated_firm_production_qty',
    'dated_proposed_qty', 'dated_release_today_qty', 'dated_released_executed_qty',
    'dated_late_qty', 'dated_late_qty_days', 'dated_first_shortage_day',
    'dated_reserve_late_qty', 'dated_requirement_late_qty',
    'dated_min_projected_qty', 'dated_horizon_days', 'dated_lead_days', 'dated_lead_basis',
    'external_demand_qty', 'external_consumed_qty', 'external_backlog_qty', 'dated_external_requirement_qty',
    'sourcing_policy_id', 'sourcing_primary_available_day',
    'sourcing_physical_shortage_before_primary_qty', 'sourcing_backup_proposed_qty',
    'sourcing_backup_released_qty', 'sourcing_primary_released_qty',
    'sourcing_retained_firm_qty', 'sourcing_bridge_surplus_qty',
    'dated_protection_target_qty', 'dated_protection_max_shortfall_qty',
    'dated_protection_first_shortfall_day',
    'dated_source_protection_target_qty', 'dated_coverage_complement_qty',
    'dated_coverage_activation_day',
    'industrial_source_requirement_qty', 'industrial_own_requirement_qty',
    'industrial_reconstructed_external_qty', 'industrial_planning_complement_qty',
    'industrial_planning_requirement_qty', 'industrial_source_covered_days',
    'industrial_source_excess_own_qty', 'industrial_target_rate_qty',
    'industrial_target_window_days', 'industrial_forecast_source_cells',
    'prospective_safety_basis', 'prospective_safety_source_days',
    'prospective_safety_tomorrow_qty', 'prospective_safety_legacy_qty',
    'prospective_safety_dated_days', 'prospective_safety_fallback_days',
    'purchase_need_date_policy', 'anticipated_safety_working_days', 'anticipated_fixed_floor_qty',
    'anticipated_late_qty', 'anticipated_first_shortage_day', 'anticipated_no_extra_time_floor',
    'purchase_rule_version', 'purchase_fixed_floor_mode', 'purchase_grouping_days',
    'purchase_requirement_basis', 'production_projection_floor_qty',
    'mrp_forecast_source', 'mrp_forecast_vintage_day', 'mrp_forecast_cycle_index',
    'mrp_forecast_window_days', 'mrp_forecast_source_days', 'mrp_forecast_fallback_days',
    'mrp_forecast_daily_qty', 'mrp_forecast_source_rows',
    'mrp_forecast_missing_period_policy', 'mrp_forecast_source_vintages',
    'production_execution_policy', 'production_dated_snapshot_day',
    'production_dated_new_release_qty', 'production_dated_late_qty',
    'production_dated_release_proposal_ids',
    'depot_safety_protection_policy', 'depot_safety_protection_qty',
)
SOURCING_ORDER_FIELDS = (
    'sourcing_policy_id', 'sourcing_role', 'sourcing_reason', 'sourcing_purchase_unit_cost',
    'sourcing_price_uom', 'sourcing_currency', 'sourcing_primary_available_day',
    'sourcing_requirement_due_day', 'sourcing_backup_required_qty', 'sourcing_proposal_id',
)
BATCHING_ORDER_FIELDS = (
    'procurement_batch_policy_id', 'procurement_batch_group_id',
    'procurement_batch_first_need_day', 'procurement_batch_window_end_day_exclusive',
    'procurement_batch_net_qty', 'procurement_batch_physical_net_qty',
    'procurement_batch_reserve_net_qty', 'procurement_batch_proposed_qty',
    'procurement_batch_rounding_multiple_qty', 'procurement_batch_grouping_days',
    'procurement_batch_proposal_id',
)
EXTERNAL_COMPONENT_FIELDS = (
    'demand_qty', 'backlog_start_qty', 'required_qty', 'consumed_qty', 'backlog_end_qty',
    'available_before_qty', 'available_after_qty', 'core_consumed_qty', 'total_component_consumed_qty',
)
EXTERNAL_FORECAST_AUDIT_FIELDS = (
    'forecast_selection_mode', 'forecast_vintage_id', 'forecast_known_day',
    'forecast_latest_known_vintage_id', 'forecast_latest_known_day', 'forecast_period_start_day',
    'forecast_source_covered', 'forecast_coverage_reason', 'forecast_source_file', 'forecast_source_cells',
)
FLOW_TYPES = {'lane_receipt': 0, 'opening_purchase_order_receipt': 0,
              'external_procurement_receipt': 0, 'production_consume': 1,
              'production_consume_reference_transition': 1, 'lane_ship': 2}
SITE_FLOW_FIELDS = ('external_receipt', 'production_output', 'transport_receipt',
                    'transport_shipment', 'availability_release', 'core_consumption',
                    'other_consumption', 'simplified_supply')


def site_flow_kind(row):
    """Keep the ledger's business distinctions; boundary supply is not fabrication."""
    event = row['event_type']
    if event == 'external_procurement_receipt':
        return ('simplified_supply' if row.get('source_type') == 'source_boundary_receipt'
                else 'external_receipt')
    if event.startswith(('production_consume', 'opening_production_consume')):
        return 'core_consumption'
    return {'opening_purchase_order_receipt': 'external_receipt',
            'production_output': 'production_output',
            'opening_production_order': 'production_output',
            'lane_receipt': 'transport_receipt', 'lane_ship': 'transport_shipment',
            'stock_availability_release': 'availability_release',
            'external_component_consume': 'other_consumption'}.get(event)


def production_availability_events(rows):
    """Yield production at actual usability, keeping physical ledger events intact.

    A held output enters the physical site at G; only an executed release of
    that same lot contributes here at I. Its scheduled available_day alone
    is never converted into an executed flow. Legacy unheld events keep G=I.
    """
    pending = defaultdict(list)
    unheld_outputs, initial_holds = set(), set()
    for row in rows:
        event = row['event_type']
        key = row['node_id'], row['item_id'], row.get('lot_id', '')
        if event in ('opening_production_order', 'production_output'):
            raw_notes = row.get('notes') or ''
            notes = json.loads(raw_notes) if raw_notes.lstrip().startswith('{') else {}
            if notes.get('availability_state') != 'held':
                initial_key = (int(row['day']), key)
                if initial_key in initial_holds:
                    raise ValueError('Production initially held requires availability metadata')
                unheld_outputs.add(initial_key)
                yield row
                continue
            day, available_day = int(row['day']), notes.get('available_day')
            quantity = _number(row['qty'])
            if (not key[2] or type(available_day) is not int or available_day < day
                    or quantity is None or quantity < 0):
                raise ValueError('Held production requires a lot and a valid availability date/quantity')
            pending[key].append([row, quantity, available_day])
        elif event == 'stock_availability_hold':
            initial_key = (int(row['day']), key)
            initial_holds.add(initial_key)
            if initial_key in unheld_outputs:
                raise ValueError('Production initially held requires availability metadata')
        elif event == 'stock_availability_release' and key in pending:
            quantity = _number(row['qty'])
            if quantity is None or quantity < 0:
                raise ValueError('Production availability release requires a nonnegative quantity')
            for entry in pending[key]:
                physical, remaining, expected_day = entry
                if quantity <= 1e-6:
                    break
                if remaining <= 1e-6:
                    continue
                if row['uom'] != physical['uom'] or int(row['day']) < expected_day:
                    raise ValueError('Production release precedes availability or changes the lot unit')
                released = min(remaining, quantity)
                yield dict(physical, day=row['day'], qty=str(released),
                           availability_event_id=row.get('event_id'))
                entry[1] -= released
                quantity -= released
            if quantity > 1e-6:
                raise ValueError('Production releases exceed the held physical lot')


WEEKLY_CONVENTIONS = {
    'sunday_start': {'label': 'Dimanche repère de début · dimanche–samedi', 'start_offset': 0, 'end_offset': 6},
    'monday_after': {'label': 'Lundi suivant le repère · lundi–dimanche', 'start_offset': 1, 'end_offset': 7},
    'sunday_end': {'label': 'Historique : dimanche repère de fin · lundi–dimanche', 'start_offset': -6, 'end_offset': 0},
}
COLUMNS = {
    'source_movements': ['sunday_label', 'period_start_day', 'period_end_day',
                         'misc_G', 'incoming_H', 'other_signed_I', 'scan3_signed_J', 'excel_row'],
    'physical_weekly': ['period_start_day', 'period_end_day', 'net_qty', 'incoming_qty',
                        'production_qty', 'shipment_qty', 'core_consumption_qty', 'other_consumption_qty'],
    'customer_history': ['period_start_day', 'positive_demand_qty', 'signed_source_qty', 'excel_row'],
    'customer_service_weekly': ['period_start_day', 'demand_qty', 'served_qty', 'backlog_end_qty'],
    'site_flows_daily': ['day', *SITE_FLOW_FIELDS],
    'observed': ['day', 'total_qty', 'excel_row'],
    'mrp': ['target_day', 'incoming_H', 'outgoing_I', 'dated_contribution_J',
            'projected_balance_K', 'excel_row'],
    'stock': ['day', 'available_qty', 'reserved_not_shipped_qty', 'available_plus_reserved_qty',
              'held_not_usable_qty', 'physical_total_qty'],
    'availability_release': ['day', 'released_existing_stock_qty'],
    'initial_purchase_daily': ['day', 'physically_received_qty', 'became_available_qty'],
    'external_component_daily': ['day', *EXTERNAL_COMPONENT_FIELDS],
    'external_component_forecast_audit': ['day', *EXTERNAL_FORECAST_AUDIT_FIELDS],
    'dated_projection': ['sunday_day', 'firm_available_qty', 'proposed_available_qty',
                         'requirement_qty', 'reserve_requirement_qty', 'available_balance_excluding_reserve',
                         'period_start_day', 'period_end_day'],
    'dated_stock_protection': ['sunday_day', 'maintained_target_qty'],
    'trace': ['day', 'used_need_per_day', 'raw_need_per_day', 'target_qty',
              'safety_floor_qty', 'net_need_eod_qty', 'all_future_receipts_qty', 'coverage_target_qty',
              *OPTIONAL_TRACE_FIELDS],
    'weekly': ['sunday_day', 'executed_receipts', 'executed_consumption',
               'executed_shipments', 'produced_available_qty', 'used_need_sum', 'raw_need_sum'],
    'weekly_alignments': ['source_sunday_day', 'executed_receipts', 'executed_consumption',
                          'executed_shipments', 'produced_available_qty', 'used_need_sum',
                          'raw_need_sum', 'period_start_day', 'period_end_day'],
    'production': ['day', 'produced_available_qty'],
    'service': ['day', 'forecast_demand_qty', 'executed_served_qty', 'backlog_eod_qty'],
}


def _hash(path: Path) -> str:
    with path.open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def _day(value) -> int:
    value = value.date() if isinstance(value, datetime) else value
    if not isinstance(value, date):
        value = date.fromisoformat(str(value)[:10])
    return (value - ORIGIN).days


def _unit(value: str) -> tuple[str, float]:
    return {'G': ('KG', .001), 'KG': ('KG', 1.), 'ZUN': ('UN', 1.),
            'UN': ('UN', 1.), 'M': ('M', 1.)}[str(value).strip().upper()]


def _number(value, factor=1., *, physical_unit=None):
    if value is None or value == '':
        return None
    result = float(value) * factor
    if not math.isfinite(result):
        raise ValueError(f'Non-finite quantity: {value}')
    if physical_unit == 'UN' and abs(result - round(result)) > 1e-6:
        raise ValueError(f'Fractional physical units: {result}')
    return result


def _customer_source_rows(history_rows, projection_rows, units, *,
                          history_headers=None, projection_headers=None):
    """Read numbered Excel values in memory, retaining signs, dates and cells."""
    layouts = {}
    for sheet, headers, defaults, quantity in [
        ('Historique', history_headers,
         ('SKU Code', 'SKU Label', 'First day of Week Year Horizon', 'Actual Demand'), 'Actual Demand'),
        ('Projection', projection_headers,
         ('SKU Code', 'SKU Label', 'First day of Week Year Snapshot',
          'First day of Week Year Horizon', 'Forecasted Demand'), 'Forecasted Demand'),
    ]:
        # Defaults preserve callers supplying the historical in-memory layout;
        # workbook callers always pass the actual header row.
        names = [' '.join(str(value or '').split()).casefold()
                 for value in (defaults if headers is None else headers)]
        required = ['SKU Code', 'First day of Week Year Horizon', quantity]
        if sheet == 'Projection':
            required.append('First day of Week Year Snapshot')
        for name in required + ['SKU Label']:
            count = names.count(name.casefold())
            if (name in required and count != 1) or count > 1:
                raise ValueError(f'Invalid customer source header: {sheet}! {name} (found {count})')
        layouts[sheet] = {name: names.index(name.casefold())
                          for name in required + ['SKU Label'] if name.casefold() in names}
        layouts[sheet]['quantity'] = layouts[sheet][quantity]
    products, seen = {}, set()
    for sheet, rows in [('Historique', history_rows), ('Projection', projection_rows)]:
        columns = layouts[sheet]
        for line, row in rows:
            if not any(value is not None for value in row):
                continue
            if len(row) <= max(columns.values()):
                raise ValueError(f'Incomplete customer source row: {sheet}!{line}')
            item = str(row[columns['SKU Code']]).zfill(6)
            if units.get(item) != 'UN':
                raise ValueError(f'Customer unit not established as UN: {item}')
            name = row[columns['SKU Label']] if 'SKU Label' in columns else None
            product = products.setdefault(item, dict(item=item, name=str(name or item), unit='UN',
                history=[], projections={}, unit_basis='Unité UN du même article dans les sources de stock ; absente du classeur client.',
                site_basis='Demande par article, sans site client dans le classeur.',
                source_columns={key: openpyxl.utils.get_column_letter(layouts[tab]['quantity'] + 1)
                                for key, tab in [('history', 'Historique'), ('projection', 'Projection')]}))
            vintage = _day(row[columns['First day of Week Year Snapshot']]) if sheet == 'Projection' else None
            start = _day(row[columns['First day of Week Year Horizon']])
            signed = _number(row[columns['quantity']], physical_unit='UN')
            if signed is None or signed > 0 or (start + ORIGIN.weekday()) % 7:
                raise ValueError(f'Invalid customer quantity or Monday: {sheet}!{line}')
            if vintage is not None and ((vintage + ORIGIN.weekday()) % 7 or start <= vintage):
                raise ValueError(f'Invalid customer forecast chronology: Projection!{line}')
            key = (sheet, item, vintage, start)
            if key in seen:
                raise ValueError(f'Duplicate customer source row: {key}')
            seen.add(key)
            target = product['history'] if vintage is None else product['projections'].setdefault(str(vintage), [])
            target.append([start, -signed, signed, line])
    for product in products.values():
        product['history'].sort()
        for rows in product['projections'].values():
            rows.sort()
    return products


def _movement_source_rows(numbered_rows, pairs):
    """Keep raw signed movements; a Sunday label starts the following Monday."""
    result, seen = {}, set()
    for line, row in numbered_rows:
        if not any(value is not None for value in row):
            continue
        item, division = str(row[0]).zfill(6), str(row[3])
        entry = _entry(pairs, item, division, row[5])
        pair, label = f'{item}/{division}', _day(row[4])
        if (label + ORIGIN.weekday()) % 7 != 6:
            raise ValueError(f'Movement label must be Sunday: row {line}')
        if (pair, label) in seen:
            raise ValueError(f'Duplicate movement week: {pair}, {label}')
        seen.add((pair, label))
        _, factor = _unit(row[5])
        values = [_number(v, factor, physical_unit=entry['unit']) for v in row[6:10]]
        if any(v is None for v in values):
            raise ValueError(f'Missing movement quantity: row {line}')
        result.setdefault(pair, []).append([label, label + 1, label + 7, *values, line])
    for pair, rows in result.items():
        rows.sort()
        photos = {r[0]: r[1] for r in pairs[pair]['observed']}
        balances = []
        for label, start, end, g, h, i, j, line in rows:
            if start in photos and end + 1 in photos:
                delta = photos[end + 1] - photos[start]
                balances.append(dict(start_day=start, delta_qty=delta,
                    signed_net_qty=g + h + i + j, gap_qty=delta - (g + h + i + j),
                    incoming_H=h, excel_row=line))
        # Detect the documented systematic inconsistency, without correcting H.
        positive = [r for r in balances if r['incoming_H'] > 0]
        suspect_h = len(positive) >= 3 and all(
            abs(r['gap_qty'] + r['incoming_H'] / 2) < 1e-6 for r in positive)
        net_only = any(r[5] > 0 or r[6] > 0 or r[4] < 0 for r in rows)
        pairs[pair]['source_movements'] = dict(rows=rows, balances=balances,
            gross_identifiable=not net_only, receipts_comparable=not net_only and not suspect_h,
            incoming_inconsistent=suspect_h, source_sheet='Feuille1',
            calendar='Repère dimanche, intervalle du lundi suivant au dimanche inclus ; étayé par les photos de stock.',
            limits=('Flux signés nets : entrées et sorties brutes non identifiables.' if net_only else
                    'Entrées H incohérentes avec les photos ; valeurs brutes conservées sans correction.' if suspect_h else
                    'Des écarts de bilan peuvent subsister ; les dates journalières des mouvements sont inconnues.'))
    return result


def _physical_week_rows(daily, has_stock):
    """Aggregate executed site ledger on complete Monday–Sunday windows."""
    lookup = {r[0]: r[1:] for r in daily}
    result = []
    for _, start, end in _weekly_windows('monday_after'):
        if not has_stock:
            result.append([start, end, *([None] * 6)])
            continue
        totals = [math.fsum(lookup.get(d, [0.] * 8)[i] for d in range(start, end + 1)) for i in range(8)]
        ext, made, received, shipped, released, core, other, simplified = totals
        incoming = ext + received + simplified + made
        result.append([start, end, incoming - shipped - core - other,
                       incoming, made, shipped, core, other])
    return result


def _customer_week_rows(daily):
    lookup = {r[0]: r for r in daily}
    result = []
    for _, start, end in _weekly_windows('monday_after'):
        if all(d in lookup for d in range(start, end + 1)):
            result.append([start, math.fsum(lookup[d][1] for d in range(start, end + 1)),
                           math.fsum(lookup[d][2] for d in range(start, end + 1)), lookup[end][3]])
        else:
            result.append([start, None, None, None])
    return result


CUSTOMER_RESPONSE_COLUMNS = {
    'observed': ['start_day', 'end_day', 'history_demand_qty', 'stock_open_qty',
                 'stock_close_qty', 'stock_delta_qty', 'movement_net_qty',
                 'balance_residual_qty', 'history_excel_row', 'opening_photo_excel_row',
                 'closing_photo_excel_row', 'movement_excel_row'],
    'simulated': ['start_day', 'end_day', 'demand_qty', 'dc_customer_shipped_qty',
                  'dc_received_qty', 'customer_received_qty', 'served_qty',
                  'transit_end_qty', 'backlog_end_qty', 'dc_stock_open_qty',
                  'dc_stock_close_qty', 'dc_stock_delta_qty', 'customer_stock_end_qty',
                  'dc_balance_residual_qty'],
}
CUSTOMER_RESPONSE_ITEMS = ('268091', '268967')


def _customer_response_run(service_rows, availability_rows, event_rows):
    """Aggregate executed registers in memory, without inventing missing states.

    Event absence means zero only for an item with a documented client service
    and DC stock. An unmatched inbound shipment makes transit unknown, rather
    than assuming that the initial pipeline was empty. The caller supplies the
    complete event ledger, including the days preceding the first full week.
    """
    services, stocks, clients = {}, {}, defaultdict(set)

    def quantity(row, field, *, physical=True):
        unit, factor = _unit(row['uom'])
        if unit != 'UN' or factor != 1:
            raise ValueError('Customer response requires UN registers')
        value = _number(row[field], physical_unit='UN' if physical else None)
        if value is None or value < 0:
            raise ValueError(f'Invalid customer response quantity: {field}')
        return value

    for row in service_rows:
        item, day = row['item_id'].removeprefix('item:'), int(row['day'])
        if item not in CUSTOMER_RESPONSE_ITEMS or not 0 <= day < DAYS:
            continue
        key = item, row['node_id'], day
        if key in services:
            raise ValueError(f'Duplicate customer response service: {key}')
        services[key] = [quantity(row, f, physical=f == 'served_qty') for f in
                         ('demand_qty', 'served_qty', 'backlog_end_qty')]
        clients[item].add(row['node_id'])
    for row in availability_rows:
        item, day, node = row['item_id'].removeprefix('item:'), int(row['day']), row['node_id']
        if item not in CUSTOMER_RESPONSE_ITEMS or not 0 <= day < DAYS:
            continue
        if node != 'DC-1920' and node not in clients[item]:
            continue
        key = item, node, day
        if key in stocks:
            raise ValueError(f'Duplicate customer response stock: {key}')
        stocks[key] = quantity(row, 'physical_qty')

    # Keep event identities and shipment linkage: a factory arrival at Muret is
    # not a client receipt, and stock release is not a new physical arrival.
    flows = defaultdict(lambda: defaultdict(lambda: [0., 0., 0., 0.]))
    shipments, receipt_events, event_ids = {}, [], set()
    for row in event_rows:
        item, day, node = row['item_id'].removeprefix('item:'), int(row['day']), row['node_id']
        if item not in CUSTOMER_RESPONSE_ITEMS or not 0 <= day < DAYS:
            continue
        if node != 'DC-1920' and node not in clients[item]:
            continue
        if row['event_id'] in event_ids:
            raise ValueError('Duplicate customer response lot event')
        event_ids.add(row['event_id'])
        event, qty = row['event_type'], quantity(row, 'qty')
        values = flows[item][day]
        if event == 'lane_ship' and node == 'DC-1920':
            destinations = [client for client in clients[item]
                            if f'_TO_{client}_' in row.get('source_id', '')]
            if not destinations:
                continue  # Not a documented DC-to-customer lane.
            if len(destinations) != 1 or not row.get('shipment_id'):
                raise ValueError('Ambiguous customer response shipment destination')
            key = item, row['shipment_id']
            shipment = shipments.setdefault(key, dict(node=destinations[0], qty=0., received=0.))
            if shipment['node'] != destinations[0]:
                raise ValueError('Customer shipment changes destination')
            shipment['qty'] += qty
            values[0] += qty
        elif node == 'DC-1920' and event in (
                'lane_receipt', 'opening_purchase_order_receipt', 'external_procurement_receipt'):
            values[1] += qty
        elif node in clients[item] and event == 'lane_receipt':
            receipt_events.append((item, row.get('shipment_id'), node, day, qty))
        elif node in clients[item] and event == 'demand_service':
            values[3] += qty
    unknown_pipeline = set()
    for item, shipment_id, node, day, qty in receipt_events:
        shipment = shipments.get((item, shipment_id))
        if shipment is None:
            unknown_pipeline.add(item)
            continue
        if shipment['node'] != node:
            raise ValueError('Customer receipt changes shipment destination')
        shipment['received'] += qty
        if shipment['received'] > shipment['qty']:
            raise ValueError('Customer receipts exceed dispatched shipment')
        flows[item][day][2] += qty

    result = {}
    for item in CUSTOMER_RESPONSE_ITEMS:
        documented_days = {d for d in range(DAYS) if clients[item]
                           and (item, 'DC-1920', d) in stocks
                           and all((item, client, d) in services for client in clients[item])}
        cumulative_transit, transit = 0., {}
        prefix_documented = True
        for day in range(DAYS):
            shipped, _, received, served = flows[item][day]
            cumulative_transit += shipped - received
            if cumulative_transit < 0:
                raise ValueError('Customer receipt precedes its dispatch')
            prefix_documented = prefix_documented and day in documented_days
            transit[day] = (cumulative_transit if prefix_documented and item not in unknown_pipeline
                            else None)
            if all((item, client, day) in services for client in clients[item]) and clients[item]:
                recorded_service = math.fsum(services[item, client, day][1] for client in clients[item])
                if served != recorded_service:
                    raise ValueError('Customer service CSV differs from executed lot register')
        weekly = []
        for _, start, end in _weekly_windows('monday_after'):
            demand = served = backlog = None
            if clients[item] and all((item, client, d) in services for client in clients[item]
                                      for d in range(start, end + 1)):
                demand = math.fsum(services[item, client, d][0] for client in clients[item]
                                   for d in range(start, end + 1))
                served = math.fsum(services[item, client, d][1] for client in clients[item]
                                   for d in range(start, end + 1))
                backlog = math.fsum(services[item, client, end][2] for client in clients[item])
            week_documented = all(d in documented_days for d in range(start, end + 1))
            totals = ([math.fsum(flows[item][d][i] for d in range(start, end + 1))
                       for i in range(3)] if week_documented else [None] * 3)
            if item in unknown_pipeline:
                totals[2] = None  # Receipt origin is not fully established.
            opening, closing = stocks.get((item, 'DC-1920', start - 1)), stocks.get((item, 'DC-1920', end))
            delta = closing - opening if opening is not None and closing is not None else None
            residual = delta - (totals[1] - totals[0]) if delta is not None and week_documented else None
            client_stock = (math.fsum(stocks[item, client, end] for client in clients[item])
                            if clients[item] and all((item, client, end) in stocks for client in clients[item])
                            else None)
            weekly.append([start, end, demand, *totals, served, transit[end], backlog,
                           opening, closing, delta, client_stock, residual])
        result[item] = weekly
    return result


def _customer_response_payload(customer_demand, pairs, run_info):
    """Join only observed weekly demand, photos and signed source movements."""
    if customer_demand is None:
        return None
    products = {}
    for item in CUSTOMER_RESPONSE_ITEMS:
        product = customer_demand['products'].get(item)
        entry = pairs.get(f'{item}/1920')
        if product is None or entry is None:
            continue
        history = {r[0]: r for r in product['history']}
        photos = {r[0]: r for r in entry['observed']}
        movements = {r[1]: r for r in entry.get('source_movements', {}).get('rows', [])}
        observed = []
        for _, start, end in _weekly_windows('monday_after'):
            h, op, cl, mv = history.get(start), photos.get(start), photos.get(end + 1), movements.get(start)
            if mv is not None and mv[2] != end:
                raise ValueError('Customer response movement week does not match photo interval')
            opening, closing = op[1] if op else None, cl[1] if cl else None
            delta = closing - opening if opening is not None and closing is not None else None
            net = math.fsum(mv[3:7]) if mv is not None else None
            residual = delta - net if delta is not None and net is not None else None
            observed.append([start, end, h[1] if h else None, opening, closing, delta, net, residual,
                             h[3] if h else None, op[2] if op else None, cl[2] if cl else None,
                             mv[7] if mv else None])
        products[item] = dict(item=item, unit='UN', dc_pair=f'{item}/1920', observed=observed,
            simulations={key: info.get('customer_response', {}).get(item, []) for key, info in run_info.items()},
            limits=['La demande historique ne donne pas les dates industrielles de prise de commande ni de livraison.',
                    'Les mouvements réels sont nets signés (G + H + I + J) ; aucune réception ou expédition brute n’en est déduite.',
                    'Le résidu réel = variation des photos − mouvements nets ; il reste visible, sans correction des sources.',
                    'Stocks client, transit et commandes non servies sont simulés ; aucun équivalent industriel n’est fourni.',
                    'Le transit est reconstitué depuis les départs et réceptions documentés dans le registre. Ce diagnostic ne couvre pas un transit client déjà engagé avant J0 dont le registre ne donne pas l’origine.'])
    return dict(schema_version=1, calendar='monday_sunday_complete_2025',
                columns=CUSTOMER_RESPONSE_COLUMNS, products=products,
                date_basis='Photos aux lundis de début et de fin ; simulation aux clôtures de la veille. Les flux couvrent lundi–dimanche inclus.',
                missing_basis='Une absence est null ; une quantité explicitement nulle reste 0. Les événements absents du registre complet valent zéro uniquement pour un circuit client documenté.')


FACTORY_DISPATCH_FACTORIES = {'268091': ('M-1810', 'Avène'), '268967': ('M-1430', 'Gien')}


def _factory_dispatch_quantity(value, *, signed=False):
    """Keep physical UN integers without turning a missing source into zero."""
    result = _number(value, physical_unit='UN')
    if result is None:
        return None
    if not signed and result < 0:
        raise ValueError('Negative factory dispatch physical quantity')
    return int(round(result))


def _factory_dispatch_observed(product, entry):
    def indexed(rows, column, kind):
        result = {}
        for row in rows:
            if row[column] in result:
                raise ValueError(f'Duplicate factory dispatch {kind}: {row[column]}')
            result[row[column]] = row
        return result

    history = indexed(product.get('history', []), 0, 'history')
    photos = indexed(entry.get('observed', []), 0, 'photo')
    movements = indexed(entry.get('source_movements', {}).get('rows', []), 1, 'movement')
    observed = []
    for _, start, end in _weekly_windows('monday_after'):
        h, op, cl, mv = history.get(start), photos.get(start), photos.get(end + 1), movements.get(start)
        if mv is not None and (mv[0] != start - 1 or mv[2] != end):
            raise ValueError('Factory dispatch source movement interval mismatch')
        opening = _factory_dispatch_quantity(op[1]) if op is not None else None
        closing = _factory_dispatch_quantity(cl[1]) if cl is not None else None
        outbound = _factory_dispatch_quantity(h[1]) if h is not None else None
        other = _factory_dispatch_quantity(mv[3], signed=True) if mv is not None else None
        delta = closing - opening if opening is not None and closing is not None else None
        raw = delta + outbound - other if all(v is not None for v in (delta, outbound, other)) else None
        status = 'missing_source' if raw is None else 'negative_reconstruction' if raw < 0 else 'reconstructed'
        observed.append(dict(
            arrival_start_day=start, arrival_end_day=end,
            arrival_start_date=(ORIGIN + timedelta(days=start)).isoformat(),
            arrival_end_date=(ORIGIN + timedelta(days=end)).isoformat(),
            stock_open_qty=opening, stock_close_qty=closing, stock_delta_qty=delta,
            history_outbound_assumption_qty=outbound, other_net_qty=other,
            reconstructed_receipt_raw_qty=raw,
            reconstructed_receipt_qty=raw if raw is not None and raw >= 0 else None,
            status=status, history_excel_row=h[3] if h is not None else None,
            history_date=(ORIGIN + timedelta(days=h[0])).isoformat() if h is not None else None,
            opening_photo_excel_row=op[2] if op is not None else None,
            opening_photo_date=(ORIGIN + timedelta(days=op[0])).isoformat() if op is not None else None,
            closing_photo_excel_row=cl[2] if cl is not None else None,
            closing_photo_date=(ORIGIN + timedelta(days=cl[0])).isoformat() if cl is not None else None,
            movement_excel_row=mv[7] if mv is not None else None,
            movement_label_date=(ORIGIN + timedelta(days=mv[0])).isoformat() if mv is not None else None))
    return observed


def _factory_dispatch_run(availability_rows, event_rows, edges, *, has_event_ledger=True):
    """Read executed factory-to-DC shipments using graph endpoints, in memory.

    Legacy edge identifiers can name DC-1910 while their actual destination is
    DC-1920. Never infer the destination from an identifier substring. Daily
    zero requires an event ledger and documented factory/DC availability. Delay
    evidence comes from linked executed receipts, never their planned dates.
    """
    dc = 'DC-1920'
    routes = {item: {} for item in FACTORY_DISPATCH_FACTORIES}
    for edge in edges:
        for item, (factory, _) in FACTORY_DISPATCH_FACTORIES.items():
            if (edge.get('type') == 'transport' and edge.get('from') == factory
                    and edge.get('to') == dc and f'item:{item}' in edge.get('items', [])):
                if edge['id'] in routes[item]:
                    raise ValueError(f'Duplicate factory dispatch route: {edge["id"]}')
                routes[item][edge['id']] = edge
    stocks, shipments, receipts, event_ids = set(), {}, [], set()

    def quantity(row, field):
        unit, factor = _unit(row['uom'])
        if unit != 'UN' or factor != 1:
            raise ValueError('Factory dispatch requires physical UN registers')
        value = _factory_dispatch_quantity(row[field])
        if value is None:
            raise ValueError(f'Missing factory dispatch register quantity: {field}')
        return value

    for row in availability_rows:
        item = row['item_id'].removeprefix('item:')
        if item not in routes:
            continue
        day, node = int(row['day']), row['node_id']
        if not 0 <= day < DAYS or node not in (FACTORY_DISPATCH_FACTORIES[item][0], dc):
            continue
        key = item, node, day
        if key in stocks:
            raise ValueError(f'Duplicate factory dispatch availability: {key}')
        quantity(row, 'physical_qty')
        stocks.add(key)
    daily = defaultdict(lambda: defaultdict(int))
    for row in event_rows:
        item, kind = row['item_id'].removeprefix('item:'), row['event_type']
        if item not in routes or kind not in ('lane_ship', 'lane_receipt'):
            continue
        day, node, route = int(row['day']), row['node_id'], row.get('source_id')
        if not 0 <= day < DAYS or route not in routes[item]:
            continue
        expected_node = FACTORY_DISPATCH_FACTORIES[item][0] if kind == 'lane_ship' else dc
        if node != expected_node:
            raise ValueError('Factory dispatch event disagrees with graph route endpoint')
        identity, shipment_id = row.get('event_id'), row.get('shipment_id')
        if not identity or identity in event_ids or not shipment_id:
            raise ValueError('Missing or duplicate factory dispatch event/shipment identity')
        event_ids.add(identity)
        qty = quantity(row, 'qty')
        key = item, shipment_id
        if kind == 'lane_receipt':
            receipts.append((key, day, route, qty, identity))
            continue
        shipment = shipments.setdefault(key, dict(shipment_id=shipment_id, route_id=route,
            departure_day=day, shipped_qty=0, received_qty=0, departure_event_ids=[],
            receipt_event_ids=[], receipt_days=set()))
        if shipment['departure_day'] != day or shipment['route_id'] != route:
            raise ValueError('Factory dispatch shipment changes day or route')
        shipment['shipped_qty'] += qty
        shipment['departure_event_ids'].append(identity)
        daily[item][day] += qty
    unmatched_receipts = defaultdict(list)
    for key, day, route, qty, identity in receipts:
        shipment = shipments.get(key)
        if shipment is None:
            unmatched_receipts[key[0]].append(dict(shipment_id=key[1], receipt_day=day,
                                                  received_qty=qty, event_id=identity))
            continue
        if shipment['route_id'] != route or day < shipment['departure_day']:
            raise ValueError('Factory receipt precedes dispatch or changes route')
        shipment['received_qty'] += qty
        if shipment['received_qty'] > shipment['shipped_qty']:
            raise ValueError('Factory receipts exceed dispatched shipment')
        shipment['receipt_event_ids'].append(identity)
        shipment['receipt_days'].add(day)
    result = {}
    for item, (factory, _) in FACTORY_DISPATCH_FACTORIES.items():
        documented_days = sorted(day for day in range(DAYS)
            if (item, factory, day) in stocks and (item, dc, day) in stocks)
        covered = set(documented_days) if has_event_ledger and routes[item] else set()
        matched, incomplete, delays = [], [], set()
        for (shipment_item, _), shipment in shipments.items():
            if shipment_item != item:
                continue
            evidence = {**shipment, 'receipt_days': sorted(shipment['receipt_days'])}
            if (shipment['received_qty'] == shipment['shipped_qty']
                    and shipment['shipped_qty'] > 0 and shipment['receipt_days']):
                evidence['transport_delay_days'] = sorted(day - shipment['departure_day']
                                                         for day in shipment['receipt_days'])
                delays.update(evidence['transport_delay_days'])
                matched.append(evidence)
            else:
                incomplete.append(evidence)
        result[item] = dict(
            daily=[dict(day=day, shipped_qty=daily[item][day] if day in covered else None)
                   for day in range(DAYS)],
            documented_days=documented_days, route_ids=sorted(routes[item]),
            transport_delay_days=next(iter(delays)) if len(delays) == 1 else None,
            transport_delay_evidence=dict(observed_delays_days=sorted(delays),
                matched_shipments=matched, incomplete_shipments=incomplete,
                unmatched_receipts=unmatched_receipts[item],
                basis='Écart entre lane_ship et lane_receipt exécutés, appariés par article, shipment_id et liaison du graphe ; quantités reçues entièrement rapprochées.'),
            has_event_ledger=has_event_ledger,
            missing_basis='Zéro seulement si la liaison usine–dépôt, le registre des événements et les états quotidiens des deux sites sont documentés ; sinon null.')
    return result


def build_factory_dispatch_comparison(customer_demand, pairs, runs: dict[str, Path]):
    """Augment a frozen comparison payload without rerunning or changing it.

    The returned observed data is a weekly *hypothesis*, not measured industrial
    dispatches. Callers shift each arrival interval backwards by a selectable
    integer lag and sum simulated departures on that same shifted interval.
    The source weeks are never distributed into invented daily observations.
    """
    if customer_demand is None:
        return None
    products = {}
    for item, (factory, factory_name) in FACTORY_DISPATCH_FACTORIES.items():
        product, entry = customer_demand.get('products', {}).get(item), pairs.get(f'{item}/1920')
        if product is None or entry is None:
            continue
        if product.get('unit', 'UN') != 'UN' or entry.get('unit', 'UN') != 'UN':
            raise ValueError('Factory dispatch source requires UN')
        products[item] = dict(item=item, name=product.get('name', item), unit='UN',
            factory_id=factory, factory_name=factory_name, dc_id='DC-1920', dc_name='Muret',
            observed=_factory_dispatch_observed(product, entry), simulations={})
    provenance = {}
    for label, folder in runs.items():
        run, files = Path(folder), {}

        def recorded(path):
            files[str(path)] = _hash(path)
            return path

        def rows(name):
            path = run / 'data' / name
            if not path.is_file():
                return []
            with recorded(path).open(encoding='utf-8-sig', newline='') as stream:
                return list(csv.DictReader(stream))

        summary = json.loads(recorded(run / 'summaries/first_simulation_summary.json').read_text(encoding='utf-8'))
        if summary.get('warmup_days') != 0:
            raise ValueError('Factory dispatch comparison requires no warm-up')
        graph_path = Path(summary['input_file'])
        if not graph_path.is_absolute():
            graph_path = ROOT / graph_path
        if _hash(recorded(graph_path)) != summary['input_sha256']:
            raise ValueError(f'Graph changed since factory dispatch simulation: {graph_path}')
        graph = json.loads(graph_path.read_text(encoding='utf-8'))
        if graph.get('meta', {}).get('opening_open_orders', {}).get('snapshot_date') != ORIGIN.isoformat():
            raise ValueError('Factory dispatch graph does not establish 2025-01-01 as day zero')
        event_path = run / 'data/production_lot_events.csv'
        has_ledger = event_path.is_file() and bool(summary.get('policy', {}).get('lot_trace_enabled'))
        simulations = _factory_dispatch_run(rows('production_stock_availability_daily.csv'),
            rows(event_path.name), graph.get('edges', []), has_event_ledger=has_ledger)
        for item, product in products.items():
            product['simulations'][label] = simulations[item]
        for path, digest in files.items():
            if _hash(Path(path)) != digest:
                raise ValueError(f'Input changed during factory dispatch comparison: {path}')
        provenance[label] = dict(run_directory=str(run), sim_days=summary['sim_days'], inputs=files)
    simulations = [sim for product in products.values() for sim in product['simulations'].values()]
    delays = {sim['transport_delay_days'] for sim in simulations}
    default_lag = next(iter(delays)) if simulations and len(delays) == 1 and None not in delays else None
    return dict(schema_version=1, origin=ORIGIN.isoformat(), products=products,
        calendar='monday_sunday_complete_2025', default_transport_lag_days=default_lag,
        transport_lag_basis=('Délai constant constaté entre départ et réception simulés ; son application aux flux industriels reste une hypothèse réglable.'
                             if default_lag is not None else 'Aucun délai constant commun démontré ; choisir une hypothèse de délai explicite.'),
        formula='stock_close_qty - stock_open_qty + history_outbound_assumption_qty - other_net_qty',
        hypothesis='HYPOTHÈSE : Actual Demand de Flow_Data_Customer_Demand.xlsx / Historique représente les sorties clients du dépôt pendant la semaine. Cette équivalence n’est pas prouvée.',
        comparison_basis='La réception hebdomadaire reconstituée est décalée en bloc vers le passé du délai choisi. Les départs usine simulés sont sommés sur exactement cette même fenêtre décalée, bornes incluses.',
        source_columns=dict(history='Historique / Actual Demand (valeur absolue)',
                            other_net='Feuille1 / G Divers signé ; ni I net ni somme G+H+I+J'),
        limits=['Aucun départ industriel journalier n’est observé ; les photos hebdomadaires et la demande ne prouvent pas les dates de transport.',
                'Une reconstruction négative reste dans le tableau brut et est exclue de la courbe, sans écrêtage à zéro.',
                'Une source manquante reste null. Les 51 semaines complètes vont du 6 janvier au 28 décembre 2025.',
                'Réservations, fabrication, libérations et départs dépôt–clients sont exclus des départs usine.'],
        provenance=provenance)


CUSTOMER_FORECAST_FIELDS = (
    'period_start_day', 'period_end_day', 'represented_start_day', 'represented_end_day',
    'weekly_source_qty', 'used_future_qty', 'source_vintage_days', 'source_rows',
    'source_days', 'fallback_days', 'fallback_source_vintage_days', 'fallback_source_rows', 'status', 'policy',
)
COLUMNS['customer_forecast'] = list(CUSTOMER_FORECAST_FIELDS)


def _customer_forecast_rows(rows):
    result, seen, clients = {}, set(), {}
    numeric = set(CUSTOMER_FORECAST_FIELDS[:6]) | {'source_days', 'fallback_days'}
    for row in rows:
        day, item, node = int(row['day']), row['item_id'].removeprefix('item:'), row['node_id']
        unit, factor = _unit(row['uom'])
        if unit != 'UN' or factor != 1:
            raise ValueError('Customer forecast export must be in UN')
        if item in clients and clients[item] != node:
            raise ValueError(f'Customer source quantity cannot be duplicated across clients: {item}')
        clients[item] = node
        values = [_number(row.get(field)) if field in numeric else row.get(field, '')
                  for field in CUSTOMER_FORECAST_FIELDS]
        start, end, represented_start, represented_end = values[:4]
        if None in (start, end, represented_start, represented_end) or not (
                start <= represented_start <= represented_end <= end and represented_start > day):
            raise ValueError('Invalid customer forecast represented interval')
        key = item, day, start
        if key in seen:
            raise ValueError(f'Duplicate customer forecast decision: {key}')
        seen.add(key)
        result.setdefault(item, {}).setdefault(str(day), []).append(values)
    for decisions in result.values():
        for rows in decisions.values():
            rows.sort(key=lambda r: r[0])
    return result


def _trace_values(row, factor):
    """Keep the seven legacy quantities first; optional dates are never scaled.

    Old runs and days without a decision snapshot retain missing values. A zero
    is meaningful only when it was actually written by the simulation.
    """
    result = [_number(row[field], factor) for field in TRACE_FIELDS]
    for field in OPTIONAL_TRACE_FIELDS:
        value = row.get(field)
        if field in ('mrp_receipt_netting_mode', 'mrp_receipt_netting_scope', 'mrp_execution_mode', 'dated_lead_basis', 'dated_action_scope', 'sourcing_policy_id', 'industrial_forecast_source_cells', 'prospective_safety_basis', 'purchase_need_date_policy', 'purchase_rule_version', 'purchase_fixed_floor_mode', 'purchase_requirement_basis',
                     'mrp_forecast_source', 'mrp_forecast_source_rows', 'mrp_forecast_missing_period_policy',
                     'mrp_forecast_source_vintages', 'production_execution_policy', 'production_dated_release_proposal_ids',
                     'depot_safety_protection_policy'):
            result.append(value or None)
        else:
            result.append(_number(value, factor if field.endswith(('_qty', '_qty_days')) else 1.))
    return result


def _metrics(values):
    if not values:
        return None
    count = len(values)
    observed = sum(o for o, _ in values) / count
    errors = [s - o for o, s in values]
    mae = sum(abs(e) for e in errors) / count
    return {'n': count, 'observed_mean': observed,
            'predicted_mean': sum(s for _, s in values) / count,
            'mae': mae, 'bias': sum(errors) / count,
            'normalized_mae_pct': 100 * mae / observed if observed else None}


def _weekly_windows(convention):
    """Sunday labels with seven daily observations wholly inside the first year."""
    definition = WEEKLY_CONVENTIONS[convention]
    first_sunday = (6 - ORIGIN.weekday()) % 7
    for sunday in range(first_sunday, DAYS, 7):
        start, end = sunday + definition['start_offset'], sunday + definition['end_offset']
        if 0 <= start <= end < DAYS:
            yield sunday, start, end


def _projection_event_buckets(events, pair, decision):
    """Separate dated flows from audit quantities and stock protection levels."""
    opening = [qty for kind, _, qty, _ in events if kind == 'opening_available']
    if len(opening) != 1:
        raise ValueError(f'Projection must have exactly one opening state: {pair}, {decision}')
    buckets = defaultdict(lambda: [0., 0., 0., 0.])
    protection_levels = {}
    for kind, target, qty, _ in events:
        if kind == 'opening_available':
            continue
        if kind == 'industrial_complement_delta':
            # Signed reconciliation evidence, never an additional requirement
            # or receipt. Other audit quantities retain their nonnegative rule.
            if qty is None or not math.isfinite(qty):
                raise ValueError(f'Invalid signed industrial planning audit: {pair}, {decision}, {kind}')
            continue
        if kind in {'industrial_source', 'industrial_own', 'industrial_prior_external',
                    'industrial_complement', 'industrial_own_excess', 'anticipated_planning_requirement',
                    'industrial_weekly_complement_before_netting', 'industrial_temporal_overlap_removed',
                    'industrial_original_own_requirement', 'industrial_own_planning_allocation',
                    'industrial_own_planned', 'industrial_planning', 'industrial_own_advanced'}:
            # Audit decomposition only: actual planning requirements are
            # exported once as requirement. Neither before-netting quantities
            # nor removed overlap represent an additional physical flow.
            if qty is None or qty < 0:
                raise ValueError(f'Invalid industrial planning audit: {pair}, {decision}, {kind}')
            continue
        if kind == 'stock_protection':
            if qty is None or qty < 0:
                raise ValueError(f'Invalid stock protection level: {pair}, {decision}, {target}')
            if target in protection_levels and protection_levels[target] != qty:
                raise ValueError(f'Conflicting stock protection levels: {pair}, {decision}, {target}')
            protection_levels[target] = qty
            continue
        position = {'requirement': 2, 'reserve': 3, 'proposal': 1}.get(kind)
        if kind.startswith('firm_'):
            position = 0
        if position is None:
            raise ValueError(f'Unknown dated projection event: {kind}')
        target = max(decision, target)
        buckets[target][position] += qty
    return opening[0], buckets, protection_levels


def _projection_rows(opening, buckets, decision, end_day):
    """Weekly projected balances, including future dates but never reserve consumption."""
    balance, result = opening, []
    for start in range(decision, end_day + 1, 7):
        end = min(start + 6, end_day)
        totals = [sum(buckets.get(day, (0., 0., 0., 0.))[i]
                      for day in range(start, end + 1)) for i in range(4)]
        balance += totals[0] + totals[1] - totals[2]
        result.append([start, *totals, balance, start, end])
    return result


def _projection_protection_levels(levels, decision, end_day):
    """Each checkpoint persists until the next one; aggregate by maximum, not sum."""
    points = iter(sorted(levels.items()))
    pending = next(points, None)
    level, result = None, []
    for start in range(decision, end_day + 1, 7):
        values = []
        for day in range(start, min(start + 6, end_day) + 1):
            while pending is not None and pending[0] <= day:
                level = pending[1]
                pending = next(points, None)
            if level is not None:
                values.append(level)
        result.append([start, max(values) if values else None])
    return result


def _weekly_rows(daily_movements, daily_production, daily_trace, has_stock, alignment):
    """Aggregate a complete event ledger and optional daily series, in memory.

    A missing event means no executed movement only when a stock series exists.
    Missing production/need observations stay None; partial weeks are excluded.
    Rows use COLUMNS['weekly_alignments'], including the seven-day bounds.
    """
    result = []
    for sunday, start, end in _weekly_windows(alignment):
        days = range(start, end + 1)
        flows = [sum(daily_movements.get(d, (0., 0., 0.))[i] for d in days)
                 if has_stock else None for i in range(3)]
        production = (sum(daily_production[d] for d in days)
                      if all(d in daily_production for d in days) else None)
        needs = [sum(daily_trace[d][i] for d in days)
                 if all(d in daily_trace for d in days) else None for i in range(2)]
        result.append([sunday, *flows, production, *needs, start, end])
    return result


def _entry(pairs, item, division, unit, name=None):
    item, division = str(item).zfill(6), str(division)
    key = f'{item}/{division}'
    standard, _ = _unit(unit)
    entry = pairs.setdefault(key, {
        'item': item, 'division': division, 'site': SITES[division],
        'site_name': name or SITE_NAMES[division], 'unit': standard,
        'source_unit': str(unit), 'observed': [], 'mrp': {}, 'mrp_meta': {},
        'simulations': {}, 'rules': {'source': {}, 'model': {}}, 'warnings': [],
    })
    if entry['unit'] != standard:
        raise ValueError(f'Incompatible units for {key}: {unit}')
    return entry


def _read_sources(inventory_path, mrp_path):
    pairs, rule_rows, lot_rows = {}, [], []
    book = openpyxl.load_workbook(inventory_path, read_only=True, data_only=True)
    try:
        keys = set()
        for line, row in enumerate(book['Stocks'].iter_rows(min_row=2, values_only=True), 2):
            item, _, division, name, quantity, _, unit, photo = row
            entry = _entry(pairs, item, division, unit, name)
            day = _day(photo)
            if not 0 <= day < DAYS:
                raise ValueError(f'Stock photo outside 2025, row {line}')
            key = (entry['item'], entry['division'], day)
            if key in keys:
                raise ValueError(f'Duplicate stock photo {key}')
            keys.add(key)
            _, factor = _unit(unit)
            entry['observed'].append([day, _number(quantity, factor, physical_unit=entry['unit']), line])
        for line, row in enumerate(book['Politique de stock MRP'].iter_rows(min_row=2, values_only=True), 2):
            item, _, division, _, safety_days, safety_qty, unit, _ = row
            entry = _entry(pairs, item, division, unit)
            _, factor = _unit(unit)
            rule = {'safety_days': _number(safety_days), 'safety_qty': _number(safety_qty, factor),
                    'calendar': 'Lundi–vendredi', 'unit': entry['unit'], 'excel_row': line}
            entry['rules']['source'].update(rule)
            rule_rows.append({'pair': f'{entry["item"]}/{entry["division"]}', **rule})
        for line, row in enumerate(book['Taile de Lot'].iter_rows(min_row=2, values_only=True), 2):
            item, _, division, _, fixed, maximum, minimum, _ = row
            item, division = str(item).zfill(6), str(division)
            known_units = {e['unit'] for e in pairs.values() if e['item'] == item}
            # This sheet has no unit column. Preserve values and state inference.
            source_units = {e['source_unit'] for e in pairs.values() if e['item'] == item}
            source_unit = next(iter(source_units)) if len(source_units) == 1 else None
            unit, factor = _unit(source_unit) if source_unit else (None, 1.)
            lot_rows.append({'item': item, 'division': division, 'site': SITES[division],
                'fixed_qty': _number(fixed, factor), 'max_qty': _number(maximum, factor),
                'min_qty': _number(minimum, factor), 'unit': unit if len(known_units) == 1 else None,
                'unit_basis': 'Unité inférée des stocks du même article ; absente de la feuille des lots.',
                'excel_row': line})
    finally:
        book.close()
    trajectories = defaultdict(list)
    metadata = defaultdict(lambda: {'receipt_days': set(), 'source_rows': 0})
    mrp_keys = set()
    book = openpyxl.load_workbook(mrp_path, read_only=True, data_only=True)
    try:
        for line, row in enumerate(book['Feuille1'].iter_rows(min_row=2, values_only=True), 2):
            snapshot, item, _, division, lead, unit, target, incoming, outgoing, physical, end = row
            entry = _entry(pairs, item, division, unit)
            key, vintage, day = f'{entry["item"]}/{entry["division"]}', _day(snapshot), _day(target)
            if not 0 <= vintage < DAYS or day < vintage:
                raise ValueError(f'Invalid MRP chronology, row {line}')
            if (key, vintage, day) in mrp_keys:
                raise ValueError(f'Duplicate MRP row: {key}, {vintage}, {day}')
            mrp_keys.add((key, vintage, day))
            _, factor = _unit(unit)
            values = [_number(v, factor) for v in (incoming, outgoing, physical, end)]
            if any(v is None for v in values):
                raise ValueError(f'Missing MRP quantity, row {line}')
            trajectories[key, vintage].append([day, *values, line])
            metadata[key, vintage]['receipt_days'].add(_number(lead))
            metadata[key, vintage]['source_rows'] += 1
    finally:
        book.close()
    residual, displayed = 0., 0
    for (key, vintage), rows in trajectories.items():
        rows.sort()
        previous = 0.
        for _, incoming, outgoing, contribution, balance, _ in rows:
            residual = max(residual, abs(balance - (previous + incoming - outgoing + contribution)))
            previous = balance
        visible = [r for r in rows if r[0] < DAYS]
        pairs[key]['mrp'][str(vintage)] = visible
        pairs[key].setdefault('mrp_full', {})[str(vintage)] = rows
        displayed += len(visible)
        meta = metadata[key, vintage]
        pairs[key]['mrp_meta'][str(vintage)] = {
            'receipt_days': sorted(meta['receipt_days']), 'source_rows': meta['source_rows'],
            'displayed_rows': len(visible), 'full_plan_last_day': rows[-1][0],
            'sum_j_full_plan': sum(r[3] for r in rows), 'initial_j': rows[0][3],
            'current_week_j': next((r[3] for r in rows if r[0] == vintage), None),
            'future_dated_j': sum(r[3] for r in rows if r[0] > vintage),
        }
    if residual > 1e-6:
        raise ValueError(f'MRP recurrence failed: {residual}')
    for entry in pairs.values():
        entry['observed'].sort()
        entry['rules']['source']['lot_rules_for_article'] = [r for r in lot_rows if r['item'] == entry['item']]
        if not entry['observed']:
            entry['warnings'].append('Aucune photo de stock dans le classeur 2025 pour ce couple.')
        elif len(entry['observed']) < 53:
            entry['warnings'].append(f'Seulement {len(entry["observed"])} photos : couverture annuelle partielle.')
    anomaly = pairs.get('002612/1450')
    if anomaly and any(d == 5 and q == 1250414 for d, q, _ in anomaly['observed']):
        anomaly['warnings'].append('Photo du 06/01 : 1 250 414 KG, contre 1 664 KG le 13/01. Valeur suspecte conservée sans correction.')
    return pairs, {'photos': len(keys), 'observed_pairs': sum(bool(e['observed']) for e in pairs.values()),
        'source_pairs': len(pairs), 'mrp_rows': len(mrp_keys), 'mrp_rows_displayed_2025': displayed,
        'mrp_trajectories': len(trajectories), 'mrp_recurrence_max_error': residual,
        'complete_flow_weeks': 51}, lot_rows


def _event_unit_evidence(rows, missing_pairs):
    """Resolve export units from explicit ledger rows, never from source stocks."""
    evidence = {}
    for line, row in enumerate(rows, 2):
        pair = row['node_id'], row['item_id']
        if pair not in missing_pairs or not row.get('uom'):
            continue
        current = _unit(row['uom'])
        if pair in evidence and _unit(evidence[pair]['uom']) != current:
            raise ValueError(f'Conflicting ledger units for dynamic pair: {pair}')
        evidence.setdefault(pair, {'uom': row['uom'], 'evidence_file': 'production_lot_events.csv',
                                   'csv_line': line})
    return evidence


def _ledger_quantity_matches(reported, recorded, *, unit, report_factor, event_factor_sum, event_count):
    """Compare one rounded daily export with a sum of independently rounded events.

    Both engine CSV exports round to six decimals in their own unit. A mass
    converted from G therefore has a thousandth of the KG rounding allowance.
    Physical UN are exact integers and receive no mass-rounding allowance.
    """
    if unit == 'UN':
        return reported == recorded
    rounding = .5e-6 * (report_factor + event_factor_sum)
    arithmetic = (event_count + 2) * math.ulp(max(abs(reported), abs(recorded), 1.))
    return abs(reported - recorded) <= rounding + arithmetic


def _external_component_rows(rows, source_units):
    """Read the optional estimated-use ledger without manufacturing missing days."""
    result = defaultdict(dict)
    physical = {'consumed_qty', 'available_before_qty', 'available_after_qty',
                'core_consumed_qty', 'total_component_consumed_qty'}
    for line, row in enumerate(rows, 2):
        pair, day = (row['node_id'], row['item_id']), int(row['day'])
        if pair not in source_units or not 0 <= day < DAYS:
            continue
        unit, factor = _unit(row['uom'])
        if unit != source_units[pair]:
            raise ValueError(f'External component unit mismatch: {pair}, line {line}')
        values = {field: _number(row[field], factor, physical_unit=unit if field in physical else None)
                  for field in EXTERNAL_COMPONENT_FIELDS}
        if any(value is None or value < -1e-6 for value in values.values()):
            raise ValueError(f'Invalid external component quantity: {pair}, line {line}')
        checks = (
            values['required_qty'] - values['demand_qty'] - values['backlog_start_qty'],
            values['backlog_end_qty'] - values['required_qty'] + values['consumed_qty'],
            values['available_after_qty'] - values['available_before_qty'] + values['consumed_qty'],
            values['total_component_consumed_qty'] - values['core_consumed_qty'] - values['consumed_qty'],
        )
        if any(abs(value) > 2e-6 for value in checks):
            raise ValueError(f'External component balance mismatch: {pair}, line {line}')
        if day in result[pair]:
            raise ValueError(f'Duplicate external component day: {pair}, {day}')
        result[pair][day] = [day, *(values[field] for field in EXTERNAL_COMPONENT_FIELDS)]
    for pair, days in result.items():
        for day, row in days.items():
            if day - 1 in days and abs(row[2] - days[day - 1][5]) > 2e-6:
                raise ValueError(f'External component backlog discontinuity: {pair}, {day}')
    return {pair: [row for _, row in sorted(days.items())] for pair, days in result.items()}


def _external_component_scope(config, pair):
    """Preserve pair-specific forecast provenance without adding its vintages."""
    key = pair[1].removeprefix('item:') + '/' + pair[0].rsplit('-', 1)[-1]
    scope = {
        'evidence_file': 'external_component_demand_daily.csv',
        'method': 'Estimated incremental use outside the modelled products; global stock retained.',
        'semantics': config.get('semantics'), 'origin': config.get('origin'),
        'schema_version': config.get('schema_version'),
        'repeat_period_days': config.get('repeat_period_days'),
        'estimation_audit': next((row for row in config.get('estimation_audit', [])
                                  if row.get('pair') == key), None),
        'sources': [row for row in config.get('rows', [])
                    if (row.get('node_id'), row.get('item_id')) == pair],
    }
    versions = [row for row in config.get('versioned_series', [])
                if (row.get('node_id'), row.get('item_id')) == pair]
    if len(versions) > 1:
        raise ValueError(f'Duplicate versioned component series: {pair}')
    if versions:
        scope.update(versioned_series=versions[0],
                     repeat_period_days=versions[0].get('repeat_period_days'),
                     forecast_update_mode='rolling_vintages_future_replacement_current_week_frozen')
    return scope


def _external_forecast_audit(rows):
    """Read optional decision provenance; absence never means observed zero."""
    result = defaultdict(dict)
    dates = {'forecast_known_day', 'forecast_latest_known_day', 'forecast_period_start_day'}
    for row in rows:
        day = int(row['day'])
        if not 0 <= day < DAYS or not any(field in row for field in EXTERNAL_FORECAST_AUDIT_FIELDS):
            continue
        if not all(field in row for field in EXTERNAL_FORECAST_AUDIT_FIELDS):
            raise ValueError('Incomplete external forecast audit columns')
        pair = row['node_id'], row['item_id']
        if day in result[pair]:
            raise ValueError(f'Duplicate external forecast audit day: {pair}, {day}')
        values = []
        for field in EXTERNAL_FORECAST_AUDIT_FIELDS:
            raw = row[field]
            value = None if raw in ('', None) else raw
            if value is not None and (field in dates or field == 'forecast_source_covered'):
                value = int(value)
                if field.endswith('known_day') and value > day:
                    raise ValueError(f'Forecast not yet known on audit day: {pair}, {day}')
                if field == 'forecast_source_covered' and value not in (0, 1):
                    raise ValueError('Invalid external forecast source coverage')
            values.append(value)
        result[pair][day] = [day, *values]
    return {pair: [row for _, row in sorted(days.items())] for pair, days in result.items()}


def _sourcing_price(value, price_uom, target_unit):
    price = _number(value)
    if price is None:
        return None
    if not price_uom:
        raise ValueError('Sourcing price has no explicit unit')
    unit, factor = _unit(price_uom)
    if unit != target_unit or price < 0:
        raise ValueError('Invalid sourcing price unit or amount')
    # A price per gram becomes a price per kilogram by division, not multiplication.
    return price / factor


def _sourcing_order(row, target_unit, quantity_factor):
    if not row.get('sourcing_policy_id'):
        return {}
    if any(field not in row for field in SOURCING_ORDER_FIELDS):
        raise ValueError('Incomplete sourcing order audit')
    result = {field: row[field] or None for field in SOURCING_ORDER_FIELDS}
    result['sourcing_purchase_unit_cost'] = _sourcing_price(
        row['sourcing_purchase_unit_cost'], row['sourcing_price_uom'], target_unit)
    result['sourcing_price_uom'] = target_unit if row.get('sourcing_price_uom') else None
    for field in ('sourcing_primary_available_day', 'sourcing_requirement_due_day'):
        result[field] = _number(row[field])
    result['sourcing_backup_required_qty'] = _number(row['sourcing_backup_required_qty'], quantity_factor)
    return result


def _sourcing_scope(config, pair, target_unit):
    rows = [row for row in config.get('rows', [])
            if (row.get('node_id'), row.get('item_id')) == pair]
    if len(rows) > 1:
        raise ValueError(f'Duplicate sourcing policy for {pair}')
    if not rows:
        return None
    rule = rows[0]
    return {**rule, 'offers': [{**offer,
        'display_purchase_unit_cost': _sourcing_price(
            offer.get('purchase_unit_cost'), offer.get('price_uom'), target_unit),
        'display_price_uom': target_unit} for offer in rule.get('offers', [])]}


def _batching_scope(config, pair, target_unit):
    rows = [row for row in config.get('rows', [])
            if (row.get('node_id'), row.get('item_id')) == pair]
    if len(rows) > 1:
        raise ValueError(f'Duplicate procurement batching policy for {pair}')
    if not rows:
        return None
    rule = rows[0]
    unit, factor = _unit(rule.get('uom'))
    if unit != target_unit:
        raise ValueError(f'Procurement batching unit mismatch: {pair}')
    return {**rule, 'display_uom': unit,
            'display_rounding_multiple_qty': _number(
                rule.get('rounding_multiple_qty'), factor, physical_unit=unit)}


def _internal_component_scope(config, pair, target_unit):
    rules = [row for row in config.get('rows', [])
             if (row.get('node_id'), row.get('item_id')) == pair]
    if len(rules) > 1:
        raise ValueError(f'Duplicate internal component policy for {pair}')
    if not rules:
        return None
    rule = rules[0]
    result = {**rule, 'display_uom': target_unit}
    if rule.get('transfer_multiple_qty') is not None:
        unit, factor = _unit(rule.get('uom'))
        if unit != target_unit:
            raise ValueError(f'Internal component policy unit mismatch: {pair}')
        result['display_transfer_multiple_qty'] = _number(
            rule['transfer_multiple_qty'], factor, physical_unit=unit)
    return result


def _batching_order(row, quantity_factor):
    if not row.get('procurement_batch_policy_id'):
        return {}
    if any(field not in row for field in BATCHING_ORDER_FIELDS):
        raise ValueError('Incomplete procurement batching order audit')
    return {field: (_number(row[field], quantity_factor if field.endswith('_qty') else 1)
                    if field.endswith(('_qty', '_day', '_days', '_day_exclusive'))
                    else row[field] or None) for field in BATCHING_ORDER_FIELDS}


def _managed_inventory_scopes(payload, states):
    """Keep physical coverage distinct from an unconfigured replenishment policy."""
    if payload is None:
        return {}
    if not isinstance(payload, dict) or not isinstance(payload.get('rows'), list):
        raise ValueError('Managed inventory scope requires explicit rows')
    result = {}
    for row in payload['rows']:
        if not isinstance(row, dict):
            raise ValueError('Invalid managed inventory scope row')
        node, item = row.get('node_id'), row.get('item_id')
        if not isinstance(node, str) or not node or not isinstance(item, str) or not item:
            raise ValueError('Managed inventory scope requires an explicit article/site')
        pair = node, item
        if pair not in states or pair in result:
            raise ValueError('Managed inventory scope must name a unique modeled stock')
        status = row.get('supply_status')
        if status not in ('unconfigured', 'configured_candidate'):
            raise ValueError('Unknown managed inventory supply status')
        if status == 'configured_candidate':
            policy = row.get('supply_policy')
            if not isinstance(policy, dict) or policy.get('status') != 'candidate_not_inferred_ERP':
                raise ValueError('Candidate managed supply requires an explicit experimental policy')
        result[pair] = {
            'supply_status': status,
            'simulation_scope': 'partial_simulation',
            'supply_limit': (('Stock, usages et nouveaux achats simulés selon une politique candidate. '
                             'Les paramètres locaux inférés restent à valider ; ce statut ne certifie pas '
                             'la reproduction du MRP industriel.') if status == 'configured_candidate' else
                            ('Stock initial, usages et engagements connus représentés ; '
                             'approvisionnement automatique non paramétré. '
                             'Aucun nouvel achat n’est généré pour ce couple article/site.')),
        }
    return result


def _run_payload(label, run, pairs):
    files = {}

    def recorded(path):
        files[str(path)] = _hash(path)
        return path

    def rows(name):
        with recorded(run / 'data' / name).open(encoding='utf-8-sig', newline='') as stream:
            yield from csv.DictReader(stream)

    summary = json.loads(recorded(run / 'summaries/first_simulation_summary.json').read_text(encoding='utf-8'))
    if summary.get('warmup_days') != 0 or int(summary['sim_days']) < DAYS:
        raise ValueError('2025 comparison requires at least 365 days and no warm-up.')
    if not summary.get('policy', {}).get('lot_trace_enabled'):
        raise ValueError('Executed flow comparison requires the lot event ledger.')
    graph_path = Path(summary['input_file'])
    if not graph_path.is_absolute():
        graph_path = ROOT / graph_path
    if _hash(graph_path) != summary['input_sha256']:
        raise ValueError(f'Graph changed since simulation: {graph_path}')
    graph = json.loads(recorded(graph_path).read_text(encoding='utf-8'))
    if graph.get('meta', {}).get('opening_open_orders', {}).get('snapshot_date') != ORIGIN.isoformat():
        raise ValueError('Model opening snapshot does not establish 2025-01-01 as day zero.')
    states = {(n['id'], s['item_id']): s for n in graph['nodes']
              for s in n.get('inventory', {}).get('states', [])}
    managed_scopes = _managed_inventory_scopes(graph.get('meta', {}).get('managed_inventory_pairs'), states)
    units = {p: s['uom'] for p, s in states.items()}
    lookup = {(e['site'], 'item:' + e['item']): key for key, e in pairs.items()}
    missing_units = set(lookup) - set(units)
    dynamic_unit_evidence = (_event_unit_evidence(rows('production_lot_events.csv'), missing_units)
                             if missing_units else {})
    units.update({pair: evidence['uom'] for pair, evidence in dynamic_unit_evidence.items()})
    daily, trace, produced = defaultdict(dict), defaultdict(dict), defaultdict(dict)
    availability = defaultdict(dict)
    reservations, movements = [], defaultdict(lambda: defaultdict(lambda: [0., 0., 0.]))
    site_flows = defaultdict(lambda: defaultdict(lambda: [0.] * len(SITE_FLOW_FIELDS)))
    production_without_process = defaultdict(lambda: defaultdict(float))
    production_sites, production_item, service_item = defaultdict(set), defaultdict(dict), defaultdict(dict)
    openings, stock_sources, policy = {}, {}, defaultdict(lambda: defaultdict(set))
    initial_orders = defaultdict(list)
    initial_purchase_execution = defaultdict(list)
    supplier_purchase_execution = defaultdict(list)
    response_events, response_stocks, response_service = [], [], []
    projected_events = defaultdict(lambda: defaultdict(list))
    initial_purchase_lots, availability_events = set(), defaultdict(list)
    initial_purchase_daily = defaultdict(lambda: defaultdict(lambda: [0., 0.]))
    external_consumption = defaultdict(lambda: defaultdict(float))
    core_component_consumption = defaultdict(lambda: defaultdict(float))
    consumption_rounding = defaultdict(lambda: defaultdict(lambda: [0., 0]))
    external_file = run / 'data/external_component_demand_daily.csv'
    external_source_rows = list(rows(external_file.name)) if external_file.is_file() else []
    external_forecast_audit = _external_forecast_audit(external_source_rows)
    external_daily = (_external_component_rows(external_source_rows,
                       {pair: pairs[key]['unit'] for pair, key in lookup.items()}) if external_file.is_file() else {})
    external_factors = {(row['node_id'], row['item_id'], int(row['day'])): _unit(row['uom'])[1]
                        for row in external_source_rows if 0 <= int(row['day']) < DAYS}
    external_config = graph.get('meta', {}).get('external_component_demands') or {}
    sourcing_config = summary.get('policy', {}).get('sourcing_policy') or {}
    batching_config = summary.get('policy', {}).get('procurement_batching') or {}
    planning_horizon = summary.get('policy', {}).get('mrp_planning_horizon') or {}
    internal_policy = summary.get('policy', {}).get('internal_component_policy') or {}
    if (external_config.get('rows') or external_config.get('versioned_series')) and not external_file.is_file():
        raise ValueError(f'External component demands require their daily audit: {external_file}')
    dated_execution = summary.get('policy', {}).get('mrp_execution_mode') in ('availability', 'dated')

    def conversion(pair):
        if pair not in units:
            raise ValueError(f'No explicit model unit for exported pair: {pair}')
        standard, factor = _unit(units[pair])
        if pair in lookup and pairs[lookup[pair]]['unit'] != standard:
            raise ValueError(f'Model/source units differ: {pair}')
        return standard, factor

    for filename in ('production_input_stocks_daily.csv', 'production_output_products_daily.csv',
                     'production_dc_stocks_daily.csv'):
        for row in rows(filename):
            day, pair = int(row['day']), (row['node_id'], row['item_id'])
            if not 0 <= day < DAYS:
                continue
            # Historical replacement references (e.g. EX-344135) can appear in
            # input stocks without an active graph state. Only source pairs are
            # compared here; production outputs remain available as BOM context.
            if filename != 'production_output_products_daily.csv' and pair not in lookup:
                continue
            unit, factor = conversion(pair)
            if filename == 'production_output_products_daily.csv':
                q = _number(row['produced_qty'], factor, physical_unit=unit)
                produced[pair][day] = q
                production_sites[pair[1]].add(pair[0])
                production_item[pair[1]][day] = production_item[pair[1]].get(day, 0.) + q
            if pair in lookup:
                if day in daily[pair]:
                    raise ValueError(f'Duplicate stock series: {pair}, {day}')
                daily[pair][day] = _number(row['stock_end_of_day'], factor, physical_unit=unit)
                stock_sources[pair] = filename + ':stock_end_of_day'
    for row in rows('mrp_trace_daily.csv'):
        day, pair = int(row['day']), (row['node_id'], row['item_id'])
        if pair not in lookup or not 0 <= day < DAYS:
            continue
        unit, factor = conversion(pair)
        if day in trace[pair]:
            raise ValueError(f'Duplicate MRP trace: {pair}, {day}')
        trace[pair][day] = _trace_values(row, factor)
        for field in ('gross_requirement_basis', 'safety_time_source_days', 'review_period_days',
                      'safety_time_calendar', 'mrp_demand_signal_source'):
            policy[pair][field].add(row[field])
        policy[pair]['safety_stock_qty'].add(_number(row['safety_stock_qty'], factor))
        # This non-process inventory state has no dedicated daily stock export.
        if pair == ('SDC-1450', 'item:693055'):
            daily[pair][day] = _number(row['stock_proj_qty'], factor, physical_unit=unit)
            stock_sources[pair] = 'mrp_trace_daily.csv:stock_proj_qty'
    for row in rows('initialization_observed_stock.csv'):
        pair = row['node_id'], row['item_id']
        if pair in lookup:
            unit, factor = _unit(row['uom'])
            if unit != pairs[lookup[pair]]['unit']:
                raise ValueError(f'Opening stock unit mismatch: {pair}')
            openings[pair] = _number(row['opening_stock_qty'], factor, physical_unit=unit)
    initial_pipeline_file = run / 'data/initialization_pipeline.csv'
    if initial_pipeline_file.is_file():
        for line, row in enumerate(rows(initial_pipeline_file.name), 2):
            pair = row['node_id'], row['item_id']
            if pair not in lookup or pair not in units:
                continue
            unit, factor = conversion(pair)
            initial_orders[pair].append({
                'qty': _number(row['seeded_pipeline_qty'], factor, physical_unit=unit),
                'unit': unit, 'category': row['category'],
                'planning_element': row.get('planning_element'),
                'physical_delivery_day': _number(row.get('physical_delivery_day')),
                'usable_day': _number(row.get('usable_day')),
                'receipt_release_days': _number(row.get('receipt_release_days')),
                'lane_src': row.get('lane_src'), 'source_file': row.get('source_file'),
                'evidence_file': initial_pipeline_file.name, 'csv_line': line,
            })
    order_file = run / 'data/mrp_orders_daily.csv'
    if dated_execution and order_file.is_file():
        for line, row in enumerate(rows(order_file.name), 2):
            pair = row['node_id'], row['item_id']
            order_type = row.get('order_type')
            if pair not in lookup or order_type not in ('opening_purchase_order', 'external_procurement_supplier_delivery'):
                continue
            unit, factor = conversion(pair)
            execution = {
                'order_id': row.get('mrp_order_id'),
                'qty': _number(row['planned_receipt_qty'], factor, physical_unit=unit),
                'unit': unit,
                **{field: _number(row.get(field)) for field in (
                    'physical_delivery_day', 'available_day',
                    'actual_physical_receipt_day', 'actual_available_day')},
                'status_end_of_run': row.get('order_status_end_of_run'),
                'source_file': row.get('source_file'), 'source_row': row.get('source_row'),
                'evidence_file': order_file.name, 'csv_line': line,
            }
            if order_type == 'opening_purchase_order':
                initial_purchase_execution[pair].append(execution)
            else:
                release_day = _number(row.get('release_day'))
                if release_day is not None and 0 <= release_day < DAYS:
                    execution.update({
                        'release_day': release_day,
                        'ordered_qty': _number(row['release_qty'], factor, physical_unit=unit),
                        'supplier': row.get('src_node_id'), 'release_status': row.get('release_status'),
                        **{field: _number(row.get(field)) for field in (
                            'delivery_reference_days', 'delivery_sampled_days', 'receipt_source_days')},
                        'receipt_calendar': row.get('receipt_calendar'),
                        **({'empirical_sample_id': row.get('empirical_sample_id'),
                            'empirical_status': row['empirical_status'],
                            **{field: _number(row.get(field)) for field in (
                                'empirical_known_day', 'empirical_pool_size', 'empirical_delta_days',
                                'empirical_unclipped_lead_days')}} if row.get('empirical_status') else {}),
                        **_sourcing_order(row, unit, factor),
                        **_batching_order(row, factor),
                    })
                    supplier_purchase_execution[pair].append(execution)
    projection_file = run / 'data/mrp_dated_plan_weekly.csv'
    if projection_file.is_file():
        for row in rows(projection_file.name):
            pair = row['node_id'], row['item_id']
            decision = int(row['decision_day'])
            if pair not in lookup or not 0 <= decision < DAYS:
                continue
            unit, factor = _unit(row['uom'])
            if unit != pairs[lookup[pair]]['unit']:
                raise ValueError(f'Projection unit mismatch: {pair}')
            projected_events[pair][decision].append((row['kind'], int(row['target_day']),
                                                     _number(row['qty'], factor), row['action_scope']))
    for row in rows('production_lot_events.csv'):
        day, pair = int(row['day']), (row['node_id'], row['item_id'])
        if pair[1].removeprefix('item:') in CUSTOMER_RESPONSE_ITEMS and 0 <= day < DAYS:
            response_events.append(row)
        if pair not in lookup or not 0 <= day < DAYS:
            continue
        event = row['event_type']
        core_event = event in ('production_consume', 'production_consume_reference_transition',
                              'opening_production_consume', 'opening_production_consume_reference_transition')
        if event not in FLOW_TYPES and not core_event and event not in ('shipment_reserve', 'opening_production_order', 'production_output', 'stock_availability_release', 'external_component_consume'):
            continue
        unit, factor = _unit(row['uom'])
        if unit != pairs[lookup[pair]]['unit']:
            raise ValueError(f'Event unit mismatch: {pair}')
        quantity = _number(row['qty'], factor, physical_unit=unit)
        flow_kind = site_flow_kind(row)
        if flow_kind:
            site_flows[pair][day][SITE_FLOW_FIELDS.index(flow_kind)] += quantity
        if core_event:
            core_component_consumption[pair][day] += quantity
        if event == 'external_component_consume':
            external_consumption[pair][day] += quantity
        if core_event or event == 'external_component_consume':
            rounding = consumption_rounding[pair][(day, 'core' if core_event else 'external')]
            rounding[0] += factor
            rounding[1] += 1
        if event == 'opening_purchase_order_receipt':
            initial_purchase_lots.add(row['lot_id'])
            initial_purchase_daily[pair][day][0] += quantity
        elif event == 'stock_availability_release':
            availability_events[pair].append((day, row['lot_id'], quantity))
        # Opening orders may supply an article with no active process (693055).
        # For ordinary outputs the production CSV already includes these orders.
        if event in ('opening_production_order', 'production_output') and pair not in produced:
            production_without_process[pair][day] += quantity
        if event in FLOW_TYPES:
            movements[pair][day][FLOW_TYPES[event]] += quantity
        if event in ('shipment_reserve', 'lane_ship'):
            reservations.append((day, row['event_id'], pair, row['lot_id'], row['shipment_id'], event, quantity))
    for pair, events in availability_events.items():
        for day, lot, quantity in events:
            if lot in initial_purchase_lots:
                initial_purchase_daily[pair][day][1] += quantity
    for pair, values in production_without_process.items():
        production_sites[pair[1]].add(pair[0])
        for day in range(DAYS):
            quantity = values.get(day, 0.)
            produced[pair][day] = quantity
            production_item[pair[1]][day] = production_item[pair[1]].get(day, 0.) + quantity
    by_shipment, by_pair, reserved = defaultdict(float), defaultdict(float), defaultdict(dict)
    reservations.sort(key=lambda r: (r[0], r[1]))
    position = 0
    for day in range(DAYS):
        while position < len(reservations) and reservations[position][0] == day:
            _, _, pair, lot, shipment, event, quantity = reservations[position]
            key = pair, lot, shipment
            if event == 'shipment_reserve':
                by_shipment[key] += quantity
                by_pair[pair] += quantity
            elif by_shipment[key] > 1e-6:
                if quantity > by_shipment[key] + 1e-6:
                    raise ValueError(f'Shipment exceeds prior reservation: {key}')
                by_shipment[key] -= quantity
                by_pair[pair] -= quantity
            position += 1
        for pair in lookup:
            if by_pair[pair] < -1e-6:
                raise ValueError(f'Negative reservation: {pair}, {day}')
            reserved[pair][day] = max(0., by_pair[pair])
    availability_file = run / 'data/production_stock_availability_daily.csv'
    has_availability = availability_file.is_file()
    if has_availability:
        for row in rows(availability_file.name):
            day, pair = int(row['day']), (row['node_id'], row['item_id'])
            if pair[1].removeprefix('item:') in CUSTOMER_RESPONSE_ITEMS and 0 <= day < DAYS:
                response_stocks.append(row)
            if pair not in lookup or not 0 <= day < DAYS:
                continue
            unit, factor = _unit(row['uom'])
            if unit != pairs[lookup[pair]]['unit']:
                raise ValueError(f'Availability unit mismatch: {pair}')
            if day in availability[pair]:
                raise ValueError(f'Duplicate availability: {pair}, {day}')
            values = [_number(row[field], factor, physical_unit=unit) for field in
                      ('available_qty', 'held_qty', 'reserved_qty', 'physical_qty', 'released_qty')]
            if any(value is None or value < -1e-6 for value in values):
                raise ValueError(f'Invalid availability quantity: {pair}, {day}')
            available, held, reserved_qty, physical, _ = values
            if abs(physical - available - held - reserved_qty) > 2e-6:
                raise ValueError(f'Physical stock partition mismatch: {pair}, {day}')
            if day in daily[pair] and abs(available - daily[pair][day]) > 2e-6:
                raise ValueError(f'Legacy available stock mismatch: {pair}, {day}')
            if abs(reserved_qty - reserved[pair].get(day, 0.)) > 2e-6:
                raise ValueError(f'Ledger reservation mismatch: {pair}, {day}')
            availability[pair][day] = values
            daily[pair][day] = available
            reserved[pair][day] = reserved_qty
            stock_sources[pair] = availability_file.name + ':available_qty'
    for row in rows('production_demand_service_daily.csv'):
        day, pair = int(row['day']), (row['node_id'], row['item_id'])
        if not 0 <= day < DAYS:
            continue
        unit, factor = conversion(pair)
        if pair[1].removeprefix('item:') in CUSTOMER_RESPONSE_ITEMS:
            # Service CSV has no unit column; use the validated graph state.
            response_service.append({**row, 'uom': units[pair]})
        target = service_item[pair[1]].setdefault(day, [0., 0., 0.])
        for index, field in enumerate(('demand_qty', 'served_qty', 'backlog_end_qty')):
            target[index] += _number(row[field], factor, physical_unit=unit if field == 'served_qty' else None)

    products, downstream = {}, defaultdict(set)
    for node in graph['nodes']:
        for process in node.get('processes', []):
            outputs = {o['item_id'] for o in process.get('outputs', [])}
            for output in outputs:
                products[output] = (node['id'], process)
            for component in process.get('inputs', []):
                downstream[component['item_id']].update(outputs)

    def contexts(item):
        distances, todo = {item: 0}, [item]
        while todo:
            current = todo.pop()
            for output in downstream[current]:
                if output not in distances:
                    distances[output] = distances[current] + 1
                    todo.append(output)
        result = []
        for target, distance in sorted(distances.items()):
            if not production_item.get(target) and not service_item.get(target):
                continue
            candidate = next((p for p in units if p[1] == target), None)
            unit, _ = _unit(units[candidate])
            result.append({'item': target.removeprefix('item:'), 'unit': unit,
                'sites': sorted(production_sites[target]),
                'relation': ('article sélectionné' if distance == 0 else
                             'composant direct de la nomenclature' if distance == 1 else 'aval de la nomenclature'),
                'production': [[d, q] for d, q in sorted(production_item[target].items())],
                'service': [[d, *v] for d, v in sorted(service_item[target].items())],
                'service_scope': 'Tous les clients simulés de cet article, hors site sélectionné.',
                'relation_scope': 'Lien par article dans la nomenclature du modèle ; ne prouve pas la consommation d’un lot du site sélectionné.',
                'production_scope': 'Quantité devenue disponible en usine ; ordres initiaux compris. Pas une production industrielle observée.'})
        return result

    # A campaign commitment, its execution and its release are different flows.
    # Keep this catalogue independent of the selected component/BOM context.
    production_catalog = {}
    for item, values in production_item.items():
        candidate = next(p for p in units if p[1] == item)
        unit, _ = conversion(candidate)
        production_catalog[item] = {
            'item': item.removeprefix('item:'), 'unit': unit,
            'sites': sorted(production_sites[item]),
            'production': [[d, q] for d, q in sorted(values.items())],
            'service': [[d, *v] for d, v in sorted(service_item[item].items())],
            'campaigns': [], 'daily_flows': {},
            'mode': 'process' if item in products else 'boundary',
            'source_pairs': [key for key, entry in pairs.items()
                             if 'item:' + entry['item'] == item],
        }
    campaign_file = run / 'data/production_campaigns.csv'
    seen_campaigns = set()
    for row in rows(campaign_file.name) if campaign_file.is_file() else []:
        item = row['output_item_id']
        if item not in production_catalog or row.get('record_type') != 'campaign':
            continue
        campaign_id = row['campaign_id']
        if campaign_id in seen_campaigns:
            raise ValueError(f'Duplicate campaign commitment: {campaign_id}')
        seen_campaigns.add(campaign_id)
        day = int(float(row['campaign_started_day']))
        if not 0 <= day < DAYS:
            continue
        unit, factor = conversion((row['node_id'], item))
        quantity = _number(row['started_qty'], factor, physical_unit=unit)
        entry = production_catalog[item]
        entry['campaigns'].append({'id': campaign_id, 'day': day,
                                   'site': row['node_id'], 'qty': quantity})
        entry['daily_flows'].setdefault(day, [0.] * 6)[0] += quantity
    for row in rows('production_output_products_daily.csv'):
        item, day = row['item_id'], int(row['day'])
        if item not in production_catalog or not 0 <= day < DAYS:
            continue
        unit, factor = conversion((row['node_id'], item))
        values = production_catalog[item]['daily_flows'].setdefault(day, [0.] * 6)
        # Absence of the execution column is not evidence of zero execution.
        executed = row.get('executed_qty')
        values[1] = (None if executed in (None, '') or values[1] is None else
                     values[1] + _number(executed, factor))
        values[2] += _number(row['produced_qty'], factor, physical_unit=unit)
    for row in production_availability_events(rows('production_lot_events.csv')):
        item, day = row['item_id'], int(row['day'])
        if item not in production_catalog or not 0 <= day < DAYS:
            continue
        event = row['event_type']
        unit, factor = _unit(row['uom'])
        if unit != production_catalog[item]['unit']:
            raise ValueError(f'Production event unit mismatch: {item}')
        quantity = _number(row['qty'], factor, physical_unit=unit)
        values = production_catalog[item]['daily_flows'].setdefault(day, [0.] * 6)
        values[3 if event == 'opening_production_order' else 4] += quantity
    # Simplified boundary supply is not a production output and never goes
    # through the production-availability iterator above.
    for row in rows('production_lot_events.csv'):
        item, day = row['item_id'], int(row['day'])
        if (item not in production_catalog or not 0 <= day < DAYS
                or site_flow_kind(row) != 'simplified_supply'):
            continue
        unit, factor = _unit(row['uom'])
        if unit != production_catalog[item]['unit']:
            raise ValueError(f'Production event unit mismatch: {item}')
        production_catalog[item]['daily_flows'].setdefault(day, [0.] * 6)[5] += _number(
            row['qty'], factor, physical_unit=unit)
    for item, entry in production_catalog.items():
        flow_rows = []
        for day in range(DAYS):
            values = entry['daily_flows'].get(day, [0.] * 6)
            if entry['mode'] == 'boundary':
                values[0] = values[1] = None  # no manufacturing process is modelled
                values[2] = None
            elif not campaign_file.is_file():
                values[0] = None
            flow_rows.append([day, *values])
        entry['daily_flows'] = flow_rows
        entry['flow_columns'] = ['day', 'launched_qty', 'executed_qty', 'available_qty',
                                 'opening_available_qty', 'new_available_qty', 'boundary_receipt_qty']

    for pair, key in lookup.items():
        entry, ts, stock = pairs[key], trace[pair], daily[pair]
        external_rows = external_daily.get(pair, [])
        external_by_day = {row[0]: row for row in external_rows}
        if external_rows and set(external_by_day) != set(range(DAYS)):
            raise ValueError(f'Incomplete external component daily coverage: {label}, {pair}')
        for day in set(external_by_day) | set(external_consumption[pair]):
            row = external_by_day.get(day)
            if row is None:
                raise ValueError(f'External component CSV/lot-ledger mismatch: {label}, {pair}, {day}')
            for kind, column, ledger in [('external', 4, external_consumption),
                                         ('core', 8, core_component_consumption)]:
                factor_sum, count = consumption_rounding[pair][(day, kind)]
                ledger_qty = ledger[pair].get(day, 0.)
                if not _ledger_quantity_matches(row[column], ledger_qty, unit=entry['unit'],
                        report_factor=external_factors[(*pair, day)],
                        event_factor_sum=factor_sum, event_count=count):
                    raise ValueError(f'{kind.title()} component CSV/lot-ledger mismatch: '
                                     f'{label}, {pair}, {day}, CSV={row[column]}, ledger={ledger_qty}')
        if stock and set(stock) != set(range(DAYS)):
            raise ValueError(f'Incomplete daily stock coverage: {label}, {pair}')
        if ts and set(ts) != set(range(DAYS)):
            raise ValueError(f'Incomplete daily MRP coverage: {label}, {pair}')
        av = availability[pair]
        projections = {}
        for decision, events in sorted(projected_events[pair].items()):
            opening, buckets, protection_levels = _projection_event_buckets(events, pair, decision)
            projected_rows = _projection_rows(opening, buckets, decision, DAYS - 1)
            projections[str(decision)] = {
                'rows': projected_rows, 'opening_available_qty': opening,
                'source_event_count': len(events), 'last_event_day': max(day for _, day, _, _ in events),
                'action_scope': sorted({scope for _, _, _, scope in events}),
                'balance_basis': 'available_projection_excluding_reserve_requirements',
                'receipt_date_basis': 'availability_I_not_physical_delivery_G',
            }
            if protection_levels:
                projections[str(decision)]['stock_protection_levels'] = _projection_protection_levels(
                    protection_levels, decision, DAYS - 1)
                projections[str(decision)]['stock_protection_basis'] = 'checkpoint_level_persists_to_horizon_weekly_maximum_not_sum_or_consumption'
            if planning_horizon.get('configured_days'):
                end_day = decision + int(planning_horizon['configured_days'])
                projections[str(decision)].update(
                    horizon_end_day=end_day,
                    horizon_rows=_projection_rows(opening, buckets, decision, end_day))
                if protection_levels:
                    projections[str(decision)]['horizon_stock_protection_levels'] = _projection_protection_levels(
                        protection_levels, decision, end_day)
        if has_availability and stock and set(av) != set(range(DAYS)):
            raise ValueError(f'Incomplete daily availability coverage: {label}, {pair}')
        result = {'stock': [[d, q, reserved[pair][d], q + reserved[pair][d],
                            av[d][1] if d in av else None, av[d][3] if d in av else None]
                           for d, q in sorted(stock.items())],
                  'availability_release': [[d, v[4]] for d, v in sorted(av.items())],
                  'site_flows_daily': [[d, *v] for d, v in sorted(site_flows[pair].items())],
                  'site_flows_source': 'production_lot_events.csv',
                  'has_availability': bool(av),
                  'stock_scope': {
                      **managed_scopes.get(pair, {}),
                      'model_scope': ('outside_simulated_stock' if not stock else
                                      'opening_orders_only' if pair not in states else 'graph_inventory'),
                      'model_unit_evidence': dynamic_unit_evidence.get(pair),
                      'unavailable_stock_documented': bool(av),
                      'represented_physical_basis': ('available_reserved_held' if av else 'available_reserved_only'),
                      'limit': (('Hors périmètre des stocks simulés : aucune série de stock pour cet article/site. Les éventuels ordres documentés ne reconstituent pas son inventaire industriel. ' if not stock else
                                 'Périmètre partiel : article/site absent des stocks initiaux du graphe ; seuls les ordres connus et leurs états sont représentés. Le zéro simulé avant réception ne prouve pas un stock industriel nul. Le stock industriel initial reste non renseigné. ' if pair not in states else '') +
                                ('Le registre documente les états représentés par ce scénario ; il ne prouve pas que toutes les indisponibilités industrielles sont modélisées.' if av else
                                 'Stock indisponible non représenté séparément dans ce run. Disponible + réservé décrit seulement le physique représenté. Le stock total simulé comparable à l’inventaire industriel n’est pas entièrement reconstitué.')),
                  },
                  'trace': [[d, *v] for d, v in sorted(ts.items())], 'weekly': [],
                  'weekly_alignments': {},
                  'initial_orders': initial_orders[pair],
                  'initial_purchase_execution': initial_purchase_execution[pair],
                  'supplier_purchase_execution': supplier_purchase_execution[pair],
                  'supplier_purchase_documented': dated_execution and order_file.is_file(),
                  'sourcing_policy': _sourcing_scope(sourcing_config, pair, entry['unit']),
                  'procurement_batching': _batching_scope(batching_config, pair, entry['unit']),
                  'internal_component_policy': _internal_component_scope(internal_policy, pair, entry['unit']),
                  'dated_projections': projections,
                  'initial_purchase_daily': ([[d, *initial_purchase_daily[pair].get(d, [0., 0.])]
                                               for d in range(DAYS)] if dated_execution and stock else []),
                  'external_component_daily': external_rows,
                  'external_component_forecast_audit': external_forecast_audit.get(pair, []),
                  'external_component_scope': (_external_component_scope(external_config, pair)
                                               if external_rows else None),
                  'initial_orders_documented': initial_pipeline_file.is_file() and pair in units,
                  'opening': openings.get(pair), 'stock_source': stock_sources.get(pair), 'metrics': {},
                  'opening_basis': 'Stock initial avant opérations et avant séparation de disponibilité.',
                  'production_context': contexts(pair[1]),
                  'production': [[d, q] for d, q in sorted(production_item.get(pair[1], {}).items())],
                  'production_site': sorted(production_sites[pair[1]]),
                  'service': [[d, *v] for d, v in sorted(service_item.get(pair[1], {}).items())],
                  'service_scope': 'Tous les clients simulés de cet article, hors site sélectionné.'}
        for convention in WEEKLY_CONVENTIONS:
            result['weekly_alignments'][convention] = _weekly_rows(
                movements[pair], produced[pair], ts, bool(stock), convention)
        # Keep the original row contract and calendar for existing consumers.
        result['weekly'] = [row[:7] for row in result['weekly_alignments']['sunday_end']]
        for name, offset in (('previous_day', -1), ('same_day', 0)):
            photos = [(d + offset, q) for d, q, _ in entry['observed'] if d > 0 and d + offset in stock]
            result['metrics'][name] = _metrics([(q, stock[d]) for d, q in photos])
            result['metrics'][name + '_onsite'] = _metrics([(q, stock[d] + reserved[pair][d]) for d, q in photos])
            result['metrics'][name + '_physical'] = _metrics([(q, av[d][3]) for d, q in photos if d in av])
        rule = {f: sorted(v) for f, v in policy[pair].items()}
        rule['graph_mrp_policy'] = states.get(pair, {}).get('mrp_policy')
        if pair[1] in products:
            factory, process = products[pair[1]]
            lot = dict(process.get('lot_sizing', {}))
            if lot.get('uom'):
                unit, factor = _unit(lot['uom'])
                lot = {k: _number(v, factor) if k.endswith('_qty') else v for k, v in lot.items()}
                lot['uom'] = unit
            rule['production_lot'] = {'site': factory, **lot}
        entry['rules']['model'][label] = rule
        entry['simulations'][label] = result
    manifest = run / 'run_manifest.json'
    if manifest.is_file():
        recorded(manifest)
    customer_forecast_path = run / 'data/customer_forecast_projection_weekly.csv'
    customer_forecasts = (_customer_forecast_rows(rows(customer_forecast_path.name))
                          if customer_forecast_path.is_file() else {})
    return {'label': LABELS.get(label, label), 'path': str(run), 'days': summary['sim_days'],
            'customer_response': _customer_response_run(response_service, response_stocks, response_events),
            'customer_forecasts': customer_forecasts,
            'production_catalog': list(production_catalog.values()),
            'chain_bom_268967': chain_bom_factors(graph),
            'kind': 'nominal' if label == 'nominal' else 'diagnostic', 'warmup_days': 0,
            'common_random_numbers': summary.get('policy', {}).get('common_random_numbers'),
            'mrp_receipt_netting_mode': summary.get('policy', {}).get('mrp_receipt_netting_mode'),
            'mrp_execution_mode': summary.get('policy', {}).get('mrp_execution_mode'),
            'opening_purchase_availability': summary.get('policy', {}).get('opening_purchase_availability'),
            'production_mrp_safety_targets': summary.get('policy', {}).get('production_mrp_safety_targets'),
            'production_mrp_safety_scope': summary.get('policy', {}).get('production_mrp_safety_scope'),
            'chain_planning_policy': summary.get('policy', {}).get('chain_planning_policy'),
            'customer_demand_policy': summary.get('policy', {}).get('customer_demand_policy'),
            'demand_execution_contract': summary.get('demand_execution_contract'),
            'has_availability': has_availability,
            'initial_stock_availability': summary.get('initial_stock_availability'),
            'external_component_demands': ({key: value for key, value in
                (summary.get('policy', {}).get('external_component_demands') or external_config).items()
                if key not in ('rows', 'versioned_series')} if external_config else None),
            'sourcing_policy': {key: value for key, value in sourcing_config.items() if key != 'rows'} or None,
            'procurement_batching': {key: value for key, value in batching_config.items() if key != 'rows'} or None,
            'mrp_planning_horizon': planning_horizon or None,
            'internal_component_policy': {key: value for key, value in internal_policy.items() if key != 'rows'} or None,
            'opening_production_order_bom_issue_mode': summary.get('policy', {}).get(
                'initialization_policy', {}).get('opening_production_order_bom_issue_mode'),
            'origin_basis': 'Date du carnet initial du graphe, sans période de chauffe.',
            'scenario_id': summary['scenario_id']}, files


def chain_bom_factors(graph):
    """Convert documented chain stocks to theoretical PF equivalents, not allocation."""
    def ratio(node_id, output_item, input_item, output_unit):
        node = next((n for n in graph.get('nodes', []) if n['id'] == node_id), {})
        processes = [p for p in node.get('processes', [])
                     if any(o.get('item_id') == output_item for o in p.get('outputs', []))]
        if len(processes) != 1:
            return None
        process = processes[0]
        inputs = [r for r in process.get('inputs', []) if r.get('item_id') == input_item]
        if len(inputs) != 1:
            return None
        raw_unit, raw_factor = _unit(inputs[0]['ratio_unit'])
        batch_unit, batch_factor = _unit(process['batch_size_unit'])
        if raw_unit != 'KG' or batch_unit != output_unit:
            return None
        qty, batch = float(inputs[0]['ratio_per_batch']) * raw_factor, float(process['batch_size']) * batch_factor
        if not math.isfinite(qty) or not math.isfinite(batch) or qty <= 0 or batch <= 0:
            return None
        return qty / batch

    raw = ratio('SDC-1450', 'item:773474', 'item:021081', 'KG')
    intermediate = ratio('M-1430', 'item:268967', 'item:773474', 'UN')
    if raw is None or intermediate is None:
        return None
    return {'raw_kg_per_intermediate_kg': raw, 'intermediate_kg_per_pf': intermediate,
            'pf_per_stock_unit': {'021081/1450': 1 / (raw * intermediate),
                '773474/1450': 1 / intermediate, '773474/1430': 1 / intermediate, '268967/1920': 1.}}


def chain_stock_equivalence(pairs, runs):
    """Sum distinct documented stock positions; missing positions remain unknown."""
    positions = ('021081/1450', '773474/1450', '773474/1430', '268967/1920')
    days = sorted({row[0] for key in positions for row in pairs.get(key, {}).get('observed', [])})
    observed = {key: {r[0]: r[1] for r in pairs.get(key, {}).get('observed', [])} for key in positions}
    physical = {(key, run): {r[0]: r[5] for r in pairs.get(key, {}).get('simulations', {}).get(run, {}).get('stock', [])}
                for key in positions for run in runs}

    def total(values, factors):
        if not factors or any(not isinstance(values.get(key), (int, float)) or not math.isfinite(values[key]) for key in positions):
            return None
        return sum(values[key] * factors['pf_per_stock_unit'][key] for key in positions)

    rows = []
    for day in days:
        rows.append({'day': day,
            'observed_pf_equivalent': total({key: observed[key].get(day) for key in positions}, runs['nominal'].get('chain_bom_268967')),
            'simulations': {run: total({key: physical[key, run].get(day - 1) for key in positions}, info.get('chain_bom_268967'))
                            for run, info in runs.items()}})
    return {'title': 'Stocks documentés de la chaîne 268967 · équivalent PF théorique',
        'description': 'Conversion BOM des matières à Gaillac, du 773474 à Gaillac et Gien, et des PF au dépôt. Photos comparées à la clôture simulée de la veille.',
        'positions': list(positions), 'factors_by_run': {key: value.get('chain_bom_268967') for key, value in runs.items()},
        'limitations': ['Stocks PF à l’usine, transit et encours exclus : ce tableau ne couvre pas toutes les positions physiques.',
            'Les stocks sources peuvent servir plusieurs produits. L’équivalent suppose leur conversion théorique vers 268967, sans prouver cette allocation ni la disponibilité des autres composants.',
            'La cible d’un an concerne la chaîne entière ; aucun seuil par site ni couverture industrielle en jours n’est déduit de ces seuls stocks.',
            'Une photo, un état physique ou une nomenclature absents rendent le total inconnu. Le 1er janvier n’a pas de clôture simulée de la veille.'],
        'rows': rows}


def build_comparison_payload(inventory_path: Path, mrp_path: Path,
                             runs: dict[str, Path], *, customer_demand_path: Path | None = None,
                             movements_path: Path | None = None) -> dict:
    """Read full runs and both Excel files; return the self-contained UI payload.

    Each run needs daily stock/production/service/MRP CSVs, initialization stock,
    the lot event ledger, its summary and the unchanged input graph. At least the
    ``nominal`` run is mandatory. No historical artifact is used at runtime.
    """
    if 'nominal' not in runs:
        raise ValueError('The unchanged nominal run must be supplied.')
    inventory_path, mrp_path = Path(inventory_path), Path(mrp_path)
    source_hashes = {str(p): _hash(p) for p in (inventory_path, mrp_path)}
    pairs, coverage, lot_rules = _read_sources(inventory_path, mrp_path)
    customer_demand = None
    if movements_path is not None:
        movements_path = Path(movements_path)
        source_hashes[str(movements_path)] = _hash(movements_path)
        book = openpyxl.load_workbook(movements_path, read_only=True, data_only=True)
        try:
            _movement_source_rows(enumerate(book['Feuille1'].iter_rows(min_row=2, values_only=True), 2), pairs)
        finally:
            book.close()
    if customer_demand_path is not None:
        customer_demand_path = Path(customer_demand_path)
        source_hashes[str(customer_demand_path)] = _hash(customer_demand_path)
        units = {}
        for entry in pairs.values():
            if entry['item'] in units and units[entry['item']] != entry['unit']:
                raise ValueError(f'Conflicting customer article unit: {entry["item"]}')
            units[entry['item']] = entry['unit']
        book = openpyxl.load_workbook(customer_demand_path, read_only=True, data_only=True)
        try:
            products = _customer_source_rows(
                enumerate(book['Historique'].iter_rows(min_row=2, values_only=True), 2),
                enumerate(book['Projection'].iter_rows(min_row=2, values_only=True), 2), units,
                history_headers=next(book['Historique'].iter_rows(max_row=1, values_only=True)),
                projection_headers=next(book['Projection'].iter_rows(max_row=1, values_only=True)))
        finally:
            book.close()
        customer_demand = dict(source_file=customer_demand_path.name, products=products,
            calendar='Lundi–dimanche, semaines complètes pour les comparaisons physiques 2025.',
            snapshot_convention='Repère mensuel de version, pas horodatage industriel exact de publication.',
            scope='Actual Demand représente une demande réalisée, pas une preuve de livraison au client.')
    run_info, csv_hashes = {}, {}
    for label, run in runs.items():
        run_info[label], csv_hashes[label] = _run_payload(label, Path(run), pairs)
        if movements_path is not None:
            for entry in pairs.values():
                sim = entry['simulations'][label]
                sim['physical_weekly'] = _physical_week_rows(sim['site_flows_daily'], bool(sim['stock']))
        if customer_demand is not None:
            for context in run_info[label]['production_catalog']:
                context['customer_service_weekly'] = _customer_week_rows(context.get('service', []))
    for path, digest in source_hashes.items():
        if _hash(Path(path)) != digest:
            raise ValueError(f'Workbook changed during comparison: {path}')
    coverage['runs'] = {label: {
        'simulated_pairs': sum(bool(e['simulations'][label]['stock']) for e in pairs.values()),
        'missing_stock_pairs': [k for k, e in pairs.items() if not e['simulations'][label]['stock']],
        'weekly_photos_compared': sum((e['simulations'][label]['metrics']['previous_day'] or {}).get('n', 0)
                                      for e in pairs.values()),
    } for label in runs}
    coverage['weekly_conventions'] = {key: {
        'complete_weeks': len(windows := list(_weekly_windows(key))),
        'first_start_day': windows[0][1] if windows else None,
        'last_end_day': windows[-1][2] if windows else None,
    } for key in WEEKLY_CONVENTIONS}
    return {'schema_version': 7 if customer_demand or movements_path else 6, 'origin': ORIGIN.isoformat(), 'end': '2025-12-31',
        'customer_demand': customer_demand,
        'customer_response': _customer_response_payload(customer_demand, pairs, run_info),
        'factory_dispatch': build_factory_dispatch_comparison(customer_demand, pairs, runs),
        'movements_source': Path(movements_path).name if movements_path else None,
        'columns': COLUMNS, 'runs': run_info, 'pairs': dict(sorted(pairs.items())),
        'weekly_conventions': WEEKLY_CONVENTIONS, 'default_weekly_alignment': 'sunday_start',
        'coverage': coverage, 'source_lot_rules': lot_rules,
        'site_overviews': {'1450': {'chain_coverage': chain_stock_equivalence(pairs, run_info)}},
        'provenance': {'source_workbooks': source_hashes, 'run_inputs': csv_hashes,
                       'extractor_sha256': _hash(Path(__file__))},
        'conventions': [
            'Exécution et photos comparées sur la première année seulement : J0–J364, du 01/01 au 31/12/2025, sans chauffe. Un tableau prévisionnel peut prolonger les plans connus sur 52 semaines, sans créer une exécution en 2026.',
            'Photo du lundi rapprochée par défaut de la clôture du dimanche précédent ; heure de photo inconnue. Autre alignement disponible.',
            'Stock total observé, disponible simulé et disponible + réservé non expédié ont des périmètres distincts. Lorsqu’un registre de disponibilité existe, le physique total simulé = disponible + réservé non expédié + détenu non utilisable ; il exclut le transit. Sans ce registre, détenu et physique total restent non renseignés.',
            'Une libération de stock détenu modifie sa disponibilité, sans réception, production ni nouvel approvisionnement. Les motifs industriels précis, qualité ou autre, ne sont pas tous documentés.',
            'Réservations reconstituées par shipment_reserve moins lane_ship. Elles ne sont pas une réception future.',
            'MRP source : 52 versions conservées séparément ; H entrée prévue, I sortie prévue, J stock déjà détenu réparti selon sa date de disponibilité, K solde projeté. J futur est du stock présent qui devient utilisable plus tard, jamais une nouvelle livraison.',
            'Pour une version MRP, J de la semaine courante + toutes les contributions J futures décrit le stock détenu réparti dans ce plan ; ce total peut différer de la photo du lundi suivant. L’écart inexpliqué ne constitue pas automatiquement du stock en contrôle qualité.',
            'Récurrence K = K précédent + H − I + J vérifiée sur toutes les lignes sources. Les courbes restent limitées aux cibles 2025 ; les lignes futures du plan source demeurent disponibles dans le tableau prévisionnel étendu.',
            'Les dates MRP du dimanche restent les repères bruts sources. Trois rattachements sont disponibles : dimanche–samedi à partir du repère, lundi–dimanche suivant, ou lundi–dimanche se terminant au repère (historique). Seules les fenêtres de sept jours entièrement en 2025 sont agrégées ; aucune absence source remplacée par zéro et aucune annualisation.',
            'Le choix par défaut est dimanche repère de début. Les 14 entrées positives H de 021081 du plan du 5 janvier concordent avec les 23 dates de livraison du carnet dans la semaine suivante. Ce rapprochement ne tranche pas la frontière dimanche/lundi et ne confirme pas la convention des autres articles ou des sorties I.',
            'Les soldes K et disponibilités J conservent le repère hebdomadaire source : celui-ci ne prouve pas une clôture quotidienne ni une date exacte de libération. Le sélecteur de convention agit sur les agrégations hebdomadaires de flux et de besoins, pas sur les photos ni les clôtures quotidiennes.',
            'Besoins MRP utilisés/bruts et cibles sont des signaux prévisionnels, distincts des consommations exécutées du registre matière.',
            'Besoin net MRP relevé en fin de journée après les commandes : sa valeur seule ne reconstitue pas la décision de début de journée. Les disponibilités futures agrégées peuvent comprendre des achats non livrés et des achats déjà sur site, encore indisponibles, lorsque le scénario distingue ces états.',
            'Le mode daté représente les achats du carnet initial par une réception physique à G, puis une libération à I, sur le même lot. La libération ne crée ni seconde réception ni transport fournisseur documenté. Les dates exécutées sont celles de la simulation ; elles ne prouvent pas une exécution industrielle.',
            'Dans le mode d’achat agrégé, le délai fournisseur FIA est compté une seule fois entre commande et livraison, puis la réception sépare livraison et disponibilité. Le détail fabrication et transport amont reste inconnu : aucun trajet ni camion fournisseur n’est déduit de ce délai. Le calendrier de réception lundi–vendredi sans fériés est une convention candidate, distincte du calendrier de sécurité confirmé.',
            'Les règles industrielles ne sont pas entièrement identifiées : l’anticipation des besoins, le regroupement des commandes, le calendrier de réception et l’assiette du stock de sécurité restent à vérifier. Une quantité standard FIA ne constitue pas à elle seule un minimum, un multiple ou une capacité journalière.',
            'Si les traces le documentent, les réceptions attendues sont séparées selon leur date prévue de disponibilité : dans l’horizon de couverture existant ou après sa borne incluse. Cette distinction concerne le réapprovisionnement par transport ; elle ne constitue pas encore un calcul MRP complet semaine par semaine.',
            'Les snapshots avant commande sont distincts des positions de fin de journée. Un champ de décision absent signifie absence de relevé à cette date, notamment hors revue ou hors réapprovisionnement par transport ; il ne vaut ni besoin nul ni commande nulle.',
            'La cible utilisée pour une fabrication est relevée séparément de la cible MRP diagnostique. Elle ne devient visible que si le run exporte ce relevé ; une cible absente ne vaut pas zéro et une cible supérieure ne prouve pas que la quantité a pu être fabriquée.',
            'Production = quantité devenue disponible, ordres initiaux inclus ; service = clients agrégés par article. Ces contextes sont hors du site sélectionné et gardent leur propre unité.',
            'Moyennes des écarts calculées seulement aux photos disponibles, hors ouverture du 1er janvier ; absence de simulation affichée comme absence, jamais comme zéro.',
            'Unités normalisées G → KG et ZUN → UN. Pas de somme entre articles ou unités hétérogènes.',
            'Règles des nouveaux classeurs affichées pour comparaison, sans les appliquer au nominal. Période de validité à confirmer.',
            'Variantes séparées du nominal. Même graine ne garantit pas les mêmes aléas après changement des ordres ; essais exploratoires non validés comme remplacement du modèle.',
        ]}
