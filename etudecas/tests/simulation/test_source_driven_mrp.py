"""Hand-computed planning oracles; memory only, no filesystem fixtures."""
from copy import deepcopy
from collections import defaultdict
from datetime import date, timedelta
import math

import pytest

from etudecas.simulation.engine import run_first_simulation as engine
from etudecas.simulation_prep.inject_mrp_seed_data_v2 import build_customer_demand_payloads


def customer_import_rows():
    graph = {"nodes": [{"id": "C-test", "type": "customer", "inventory": {
        "states": [{"item_id": "item:268091", "uom": "UN"}]}}]}
    history = [{"SKU Code": "268091", "First day of Week Year Horizon": "2024-12-30",
                "Actual Demand": "-13300"},
               {"SKU Code": "268091", "First day of Week Year Horizon": "2025-01-06",
                "Actual Demand": "0"}]
    projection = [{"SKU Code": "268091", "First day of Week Year Snapshot": "2024-12-30",
                   "First day of Week Year Horizon": (date(2025, 1, 6) + timedelta(weeks=w)).isoformat(),
                   "Forecasted Demand": -329532 if w == 0 else -700}
                  for w in range(52)]
    return graph, history, projection


def test_customer_import_preserves_calendar_sign_zero_cells_and_graph():
    graph, history, projection = customer_import_rows()
    before = deepcopy((graph, history, projection))
    payload = build_customer_demand_payloads(graph, history, projection, source_file="customer.xlsx", source_sha256="abc")
    assert (graph, history, projection) == before
    first, second = payload["customer_demand_history"]["rows"]
    assert (first["period_start_day"], first["qty"], first["source_row"]) == (-2, 13300, "Historique!A2:D2")
    assert second["qty"] == 0
    forecast = payload["mrp_forecasts"]["rows"][0]
    assert (forecast["vintage_day"], forecast["period_start_day"], forecast["qty"]) == (-2, 5, 329532)
    assert forecast["source_row"] == "Projection!A2:E2"
    assert payload["mrp_forecasts"]["calendar"] == "monday_start"


@pytest.mark.parametrize("fault", ["positive", "fraction", "missing", "nan", "duplicate", "sunday", "gap", "short", "unknown", "unit", "future_vintage"])
def test_customer_import_rejects_ambiguous_source_in_memory(fault):
    graph, history, projection = customer_import_rows()
    if fault in {"positive", "fraction", "missing", "nan"}:
        history[0]["Actual Demand"] = {"positive": 1, "fraction": -1.5, "missing": None, "nan": "nan"}[fault]
    elif fault == "duplicate": history.append(dict(history[0]))
    elif fault == "sunday": history[0]["First day of Week Year Horizon"] = "2024-12-29"
    elif fault == "gap": history[1]["First day of Week Year Horizon"] = "2025-01-13"
    elif fault == "short": projection.pop()
    elif fault == "unknown": history[0]["SKU Code"] = "other"
    elif fault == "unit": graph["nodes"][0]["inventory"]["states"][0]["uom"] = "KG"
    elif fault == "future_vintage": projection[0]["First day of Week Year Snapshot"] = "2025-01-06"
    with pytest.raises(ValueError):
        build_customer_demand_payloads(graph, history, projection, source_file="customer.xlsx", source_sha256="abc")


RECEIPT_HOLIDAYS_2025 = [
    "2025-01-01", "2025-04-21", "2025-05-01", "2025-05-08", "2025-05-29",
    "2025-06-09", "2025-07-14", "2025-08-15", "2025-11-01", "2025-11-11", "2025-12-25",
]


def receipt_closure_row(pair=("SDC-1450", "item:021081")):
    return {"node_id": pair[0], "item_id": pair[1], "receipt_days": 75,
            "source_file": "Flow_Data_MRP_results.xlsx", "source_cells": "Feuille1!E277",
            "nonworking_dates": ["2025-08-04"], "calendar_id": "dated_receipt_candidate",
            "calendar_status": "conditional_source_supported_not_universal",
            "calendar_source": "Explicit receipt-only experimental dates"}


def receipt_closure_payload(rows):
    return {"schema_version": 1, "calendar": "monday_friday", "rows": rows}


@pytest.mark.parametrize("physical,days,expected", [(2,1,5), (3,1,5), (0,13,19), (364,1,365)])
def test_receipt_closure_absent_preserves_weekdays_and_next_year(physical, days, expected):
    assert engine.supplier_receipt_available_day(physical, days, origin_date="2025-01-01") == expected
    assert engine.supplier_receipt_available_day(
        physical, days, origin_date="2025-01-01", nonworking_dates=frozenset()) == expected


@pytest.mark.parametrize("source_row,physical,expected", [
    (23,"2025-04-16","2025-08-21"), (25,"2025-05-20","2025-09-19"),
    (30,"2025-05-12","2025-09-11"), (39,"2025-04-16","2025-08-21"),
    (42,"2025-05-05","2025-09-05"),
])
def test_receipt_closure_matches_five_gaillac_source_dates(source_row, physical, expected):
    # Extract_En_cours.xlsx/Sheet1 G/H/I; independent source audit identifies
    # these five rows as the only discriminants among 104 opening-book rows.
    closure = [(date(2025,8,4)+timedelta(days=n)).isoformat() for n in range(14)]
    dates = engine.parse_supplier_receipt_nonworking_dates(sorted(set(RECEIPT_HOLIDAYS_2025 + closure)))
    origin = date(2025,1,1)
    available = engine.supplier_receipt_available_day(
        (date.fromisoformat(physical)-origin).days, 75, origin_date=origin.isoformat(), nonworking_dates=dates)
    assert origin+timedelta(days=available) == date.fromisoformat(expected), source_row


@pytest.mark.parametrize("physical,days,expected", [
    ("2025-04-30",1,"2025-05-02"), ("2025-12-24",2,"2025-12-29"),
    ("2025-12-25",0,"2025-12-25"), ("2025-12-31",1,"2026-01-01"),
    ("2026-08-03",1,"2026-08-04"),
])
def test_receipt_closure_counts_absolute_dates_once_and_does_not_recur(physical, days, expected):
    origin = date(2025,1,1)
    dates = engine.parse_supplier_receipt_nonworking_dates(RECEIPT_HOLIDAYS_2025 + ["2025-08-04"])
    available = engine.supplier_receipt_available_day(
        (date.fromisoformat(physical)-origin).days, days, origin_date=origin.isoformat(), nonworking_dates=dates)
    assert origin+timedelta(days=available) == date.fromisoformat(expected)


@pytest.mark.parametrize("invalid", [
    None, True, "2025-08-04", [None], [True], [date(2025,8,4)],
    ["20250804"], ["2025-8-4"], ["2025-02-30"], ["2025-08-04 "],
    ["2025-08-04T00:00:00"], ["2025-08-04", "2025-08-04"],
])
def test_receipt_closure_rejects_noncanonical_or_missing_dates(invalid):
    with pytest.raises(ValueError):
        engine.parse_supplier_receipt_nonworking_dates(invalid)


@pytest.mark.parametrize("invalid", [["2025-08-04"], None, frozenset({"2025-08-04"}), frozenset({True})])
def test_receipt_closure_helper_requires_compiled_date_values(invalid):
    with pytest.raises(ValueError):
        engine.supplier_receipt_available_day(0,1,origin_date="2025-01-01",nonworking_dates=invalid)


def test_receipt_closure_resolver_is_pair_scoped_and_does_not_mutate_source():
    gaillac = ("SDC-1450", "item:021081")
    avene = ("M-1810", "item:001757")
    row = receipt_closure_row(gaillac)
    other = {"node_id": avene[0], "item_id": avene[1], "receipt_days":13,
             "source_file":"Flow_Data_MRP_results.xlsx", "source_cells":"Feuille1!E2"}
    payload = receipt_closure_payload([row,other])
    before = deepcopy(payload)
    result = engine.resolve_supplier_receipt_nonworking_dates(
        payload,execution_mode="dated",eligible_pairs={gaillac,avene})
    assert result == {gaillac:frozenset({date(2025,8,4)})}
    assert payload == before
    # At the same physical receipt, only the opted-in receiving site skips Monday.
    assert engine.supplier_receipt_available_day(212,1,origin_date="2025-01-01",nonworking_dates=result[gaillac]) == 216
    assert engine.supplier_receipt_available_day(212,1,origin_date="2025-01-01") == 215


@pytest.mark.parametrize("mode", ["historical", "availability"])
def test_receipt_closure_rejects_inactive_execution_modes(mode):
    row=receipt_closure_row()
    with pytest.raises(ValueError):
        engine.resolve_supplier_receipt_nonworking_dates(
            receipt_closure_payload([row]),execution_mode=mode,eligible_pairs={(row['node_id'],row['item_id'])})


def test_receipt_closure_legacy_metadata_has_no_effect_or_required_extension():
    payload=receipt_closure_payload([{"node_id":"M-1810","item_id":"item:001757","receipt_days":13}])
    for mode in ["historical","availability","dated"]:
        assert engine.resolve_supplier_receipt_nonworking_dates(payload,execution_mode=mode,eligible_pairs=set()) == {}
    assert engine.resolve_supplier_receipt_nonworking_dates(None,execution_mode="historical",eligible_pairs=set()) == {}


@pytest.mark.parametrize("field,value", [
    ("calendar_id",None), ("calendar_status",""), ("calendar_source"," "), ("nonworking_dates",None),
])
def test_receipt_closure_requires_complete_calendar_provenance(field,value):
    row=receipt_closure_row()
    pair=(row['node_id'],row['item_id'])
    for missing in [True,False]:
        changed=deepcopy(row)
        if missing:
            changed.pop(field)
        else:
            changed[field]=value
        with pytest.raises(ValueError):
            engine.resolve_supplier_receipt_nonworking_dates(
                receipt_closure_payload([changed]),execution_mode="dated",eligible_pairs={pair})


def test_receipt_closure_rejects_internal_pairs_duplicates_and_conflicting_ids():
    row=receipt_closure_row()
    pair=(row['node_id'],row['item_id'])
    with pytest.raises(ValueError):
        engine.resolve_supplier_receipt_nonworking_dates(
            receipt_closure_payload([row]),execution_mode="dated",eligible_pairs=set())
    with pytest.raises(ValueError):
        engine.resolve_supplier_receipt_nonworking_dates(
            receipt_closure_payload([row,row]),execution_mode="dated",eligible_pairs={pair})
    other=receipt_closure_row(("M-1810","item:001757"))
    other['nonworking_dates']=["2025-08-05"]
    with pytest.raises(ValueError):
        engine.resolve_supplier_receipt_nonworking_dates(
            receipt_closure_payload([row,other]),execution_mode="dated",
            eligible_pairs={pair,(other['node_id'],other['item_id'])})


def test_fia_provenance_matches_real_excel_rows_in_both_readers():
    from pathlib import Path
    from openpyxl import load_workbook
    from etudecas.knowledge_graph import update_supply_graph_from_case_data as source

    directory = Path(__file__).resolve().parents[2] / "data" / "source"
    checked = 0
    for name in ("268091.xlsx", "268967.xlsx", "773474.xlsx"):
        path = directory / name
        book = load_workbook(path, read_only=True, data_only=True)
        expected = {row: values for row, values in enumerate(
            book["FIA"].iter_rows(min_row=2, values_only=True), start=2)
            if values[0] not in (None, "")}
        book.close()
        for reader in (source.load_workbook_rows, source.load_workbook_rows_from_zip):
            _, rows = reader(path, "FIA")
            rows = [row for row in rows if row.get("Numéro d'article") not in (None, "")]
            assert len(rows) == len(expected)
            for row in rows:
                physical = expected[row["_source_row"]]
                assert str(row["Numéro d'article"]) == str(physical[0])
                assert row["Numéro de compte fournisseur"] == physical[1]
                edge = {"id": "lane", "attrs": {"source_sheet": "Relations_acteurs", "source_row": 999}}
                report = {"updated_edges": []}
                source.update_edge_from_fia(edge, row, name, name[:-5], report)
                assert edge["attrs"]["source_row"] == row["_source_row"]
                assert edge["attrs"]["source_workbook"] == name
                assert edge["attrs"]["source_sheet"] == "FIA"
                assert report["updated_edges"][0]["source_row"] == row["_source_row"]
                assert edge["order_terms"]["sell_price"] == physical[2]
                assert edge["order_terms"]["price_base"] == physical[3]
                checked += 1
    assert checked == 70


def test_fia_without_source_row_does_not_keep_an_unrelated_cell_reference():
    from etudecas.knowledge_graph import update_supply_graph_from_case_data as source

    edge = {"id": "lane", "attrs": {"source_row": 8}}
    source.update_edge_from_fia(edge, {"Montant": 5.43, "Base de prix": 1},
                                "268091.xlsx", "268091", {"updated_edges": []})
    assert edge["attrs"]["source_row"] is None
    assert edge["order_terms"]["sell_price"] == 5.43


PF = ("C", "item:PF")
OTHER = ("C", "item:OTHER")


def forecast_row(vintage, start, quantity, *, pair=PF, source="I10"):
    return {"node_id": pair[0], "item_id": pair[1], "vintage_day": vintage,
            "period_start_day": start, "period_days": 7, "qty": quantity,
            "uom": "UN", "source_row": source}


def forecast_payload(rows):
    return {"schema_version": 1, "origin_date": "2025-01-01", "calendar": "sunday_start",
            "repeat_period_days": 365, "current_bucket_policy": "excluded_unresolved_carry_over", "rows": rows}


def make_forecast(rows, *, pairs=None, excluded=None):
    payload = forecast_payload(rows)
    if excluded is not None:
        payload["excluded_current_bucket"] = excluded
    return engine.RollingMrpForecast(payload, demand_pairs=pairs or {PF}, origin_date="2025-01-01")


def test_multilevel_forecast_explodes_ten_finished_into_twenty_semi_and_sixty_raw():
    lanes = [{"src": "F", "dst": "C", "item_id": "item:PF"},
             {"src": "S", "dst": "F", "item_id": "item:SEMI"},
             {"src": "V", "dst": "S", "item_id": "item:RAW"}]
    bom = {("F", "item:PF"): [(("F", "item:SEMI"), 2)],
           ("S", "item:SEMI"): [(("S", "item:RAW"), 3)]}
    signal = engine.propagate_supply_demand_rates({PF: 10}, lanes, bom)
    assert signal == {PF: 10, ("F", "item:PF"): 10,
                      ("F", "item:SEMI"): 20, ("S", "item:SEMI"): 20,
                      ("S", "item:RAW"): 60, ("V", "item:RAW"): 60}
    components, outputs = engine.lotified_mps_component_signal(
        {("F", "item:PF"): 10, ("S", "item:SEMI"): 20}, day=0,
        process_input_requirements_by_output_pair=bom, production_lot_policy_by_pair={},
        process_capacity_by_output_pair={pair: 1_000_000 for pair in bom},
        mps_open_campaign_qty_by_pair={}, mps_started_lots_by_week_pair={})
    assert outputs == {("F", "item:PF"): 10, ("S", "item:SEMI"): 20}
    assert components == {("F", "item:SEMI"): 20, ("S", "item:RAW"): 60}


def test_zero_forecast_does_not_turn_available_capacity_into_a_requirement():
    bom = {("F", "item:PF"): [(("F", "item:RAW"), 3)]}
    assert engine.propagate_supply_demand_rates({("F", "item:PF"): 0}, [], bom) == {}
    components, outputs = engine.lotified_mps_component_signal(
        {("F", "item:PF"): 0}, day=0, process_input_requirements_by_output_pair=bom,
        production_lot_policy_by_pair={}, process_capacity_by_output_pair={("F", "item:PF"): 1_000_000},
        mps_open_campaign_qty_by_pair={}, mps_started_lots_by_week_pair={})
    assert components == outputs == {}


def test_complete_multilevel_program_does_not_explode_same_site_twice():
    bom = {("F", "item:PF"): [(("F", "item:SEMI"), 2)],
           ("F", "item:SEMI"): [(("F", "item:RAW"), 3)]}
    components, outputs = engine.lotified_mps_component_signal(
        {("F", "item:PF"): 10, ("F", "item:SEMI"): 20}, day=0,
        process_input_requirements_by_output_pair=bom, production_lot_policy_by_pair={},
        process_capacity_by_output_pair={pair: 1_000_000 for pair in bom},
        mps_open_campaign_qty_by_pair={}, mps_started_lots_by_week_pair={}, expand_internal_outputs=False)
    assert outputs == {("F", "item:PF"): 10, ("F", "item:SEMI"): 20}
    assert components == {("F", "item:SEMI"): 20, ("F", "item:RAW"): 60}


def test_historical_recursive_mps_still_explodes_an_unexpanded_program():
    bom = {("F", "item:PF"): [(("F", "item:SEMI"), 2)],
           ("F", "item:SEMI"): [(("F", "item:RAW"), 3)]}
    components, outputs = engine.lotified_mps_component_signal(
        {("F", "item:PF"): 10}, day=0, process_input_requirements_by_output_pair=bom,
        production_lot_policy_by_pair={}, process_capacity_by_output_pair={pair: 1_000_000 for pair in bom},
        mps_open_campaign_qty_by_pair={}, mps_started_lots_by_week_pair={})
    assert outputs == {("F", "item:PF"): 10, ("F", "item:SEMI"): 20}
    assert components == {("F", "item:RAW"): 60}


def test_standard_order_reference_is_not_a_forced_pack_multiple():
    # Source H712 = 144631 UN; a 5000-unit supplier standard is not its minimum.
    assert engine.mrp_purchase_order_quantity(144631, 200000, 5000, binding=False, uom="UN") == 144631
    assert engine.mrp_purchase_order_quantity(144631, 200000, 5000, binding=True, uom="UN") == 145000


def test_nonbinding_standard_can_use_stock_below_one_reference_order():
    assert engine.mrp_purchase_order_quantity(144631, 10000, 20000, binding=False, uom="UN") == 10000
    assert engine.mrp_purchase_order_quantity(144631, 10000, 20000, binding=True, uom="UN") == 0


def test_nonbinding_purchase_still_enforces_integer_physical_units():
    assert engine.mrp_purchase_order_quantity(10.9, 100, 5000, binding=False, uom="UN") == 10
    assert engine.mrp_purchase_order_quantity(10.9, 9.9, 5000, binding=False, uom="UN") == 9
    assert engine.mrp_purchase_order_quantity(10.9, 100, 5000, binding=False, uom="KG") == 10.9


@pytest.mark.parametrize("value", [None, math.nan, math.inf, -1])
def test_nonbinding_purchase_does_not_impute_invalid_need(value):
    with pytest.raises((TypeError, ValueError)):
        engine.mrp_purchase_order_quantity(value, 10000, 5000, binding=False, uom="UN")


def test_rolling_forecast_cannot_see_a_future_vintage():
    forecast = make_forecast([forecast_row(4, 18, 70, source="I10"),
                              forecast_row(11, 18, 210, source="I20")])
    value, audit = forecast.window(PF, decision_day=10, target_day=18, fallback_daily_values=[2] * 7)
    assert value == 10
    assert (audit["vintage_day"], audit["source_days"], audit["fallback_days"], audit["source_rows"]) == (4, 7, 0, "I10")
    value, audit = forecast.window(PF, decision_day=11, target_day=18, fallback_daily_values=[2] * 7)
    assert value == 30
    assert audit["vintage_day"] == 11


def test_rolling_forecast_before_first_version_uses_only_nominal_fallback():
    forecast = make_forecast([forecast_row(4, 11, 700)])
    value, audit = forecast.window(PF, decision_day=0, target_day=11, fallback_daily_values=[1, 2, 3, 4, 5, 6, 7])
    assert value == 4
    assert audit["vintage_day"] == ""
    assert (audit["source_days"], audit["fallback_days"]) == (0, 7)


def test_missing_week_does_not_reuse_old_vintage_or_become_zero():
    forecast = make_forecast([forecast_row(4, 25, 700, source="I10"),
                              forecast_row(11, 18, 210, source="I20")])
    value, audit = forecast.window(PF, decision_day=11, target_day=25, fallback_daily_values=[1, 2, 3, 4, 5, 6, 7])
    assert value == 4
    assert (audit["vintage_day"], audit["source_days"], audit["fallback_days"], audit["source_rows"]) == (11, 0, 7, "")


def test_explicit_forecast_zero_stays_zero_and_keeps_provenance():
    forecast = make_forecast([forecast_row(11, 25, 0, source="I20")])
    value, audit = forecast.window(PF, decision_day=11, target_day=25, fallback_daily_values=[100] * 7)
    assert value == 0
    assert (audit["source_days"], audit["fallback_days"], audit["source_rows"]) == (7, 0, "I20")


def test_current_bucket_backlog_is_not_injected_as_a_new_forecast():
    excluded = [forecast_row(4, 4, 637074, source="I545")]
    forecast = make_forecast([forecast_row(4, 11, 188520, source="I546")], excluded=excluded)
    value, audit = forecast.window(PF, decision_day=4, target_day=4, fallback_daily_values=[1900] * 7)
    assert value == 1900
    assert (audit["source_days"], audit["fallback_days"], audit["source_rows"]) == (0, 7, "")
    assert excluded[0]["qty"] == 637074  # Source retained for audit, never erased.


def test_unresolved_current_bucket_cannot_accidentally_enter_runtime_rows():
    with pytest.raises(ValueError):
        make_forecast([forecast_row(4, 4, 637074, source="I545")])


def test_mixed_forecast_window_counts_partial_source_and_fallback_days():
    forecast = make_forecast([forecast_row(4, 11, 70, source="I10")])
    value, audit = forecast.window(PF, decision_day=4, target_day=16, fallback_daily_values=[1, 2, 3, 4])
    # Days 16,17: 10 each; days 18,19: explicit fallback 3 and 4.
    assert value == 27 / 4
    assert (audit["source_days"], audit["fallback_days"]) == (2, 2)


def test_latest_vintage_is_global_and_does_not_mix_pairs_from_older_versions():
    forecast = make_forecast([forecast_row(4, 18, 700, pair=OTHER),
                              forecast_row(11, 18, 210)], pairs={PF, OTHER})
    value, audit = forecast.window(OTHER, decision_day=11, target_day=18, fallback_daily_values=[3] * 7)
    assert value == 3
    assert (audit["vintage_day"], audit["source_days"], audit["fallback_days"]) == (11, 0, 7)


def test_annual_repeat_keeps_last_known_year_end_plan_until_next_release():
    forecast = make_forecast([forecast_row(4, 11, 70, source="I10"),
                              forecast_row(361, 368, 14, source="I90")])
    value, audit = forecast.window(PF, decision_day=368, target_day=368, fallback_daily_values=[99] * 7)
    assert value == 2
    assert (audit["vintage_day"], audit["cycle_index"]) == (361, 0)
    value, audit = forecast.window(PF, decision_day=369, target_day=376, fallback_daily_values=[99] * 7)
    assert value == 10
    assert (audit["vintage_day"], audit["cycle_index"]) == (369, 1)


def test_forecast_fractional_planning_does_not_mutate_physical_demand_inputs():
    source_rows = [forecast_row(4, 11, 1)]
    original = deepcopy(source_rows)
    nominal = [1900] * 7
    forecast = make_forecast(source_rows)
    value, audit = forecast.window(PF, decision_day=4, target_day=11, fallback_daily_values=nominal)
    assert value == pytest.approx(1 / 7)
    assert source_rows == original
    assert nominal == [1900] * 7
    assert audit["source_days"] == 7


def test_overlapping_forecast_periods_are_rejected_instead_of_double_counted():
    with pytest.raises(ValueError, match="Overlapping"):
        make_forecast([forecast_row(4, 11, 70), forecast_row(4, 14, 70, source="I11")])


@pytest.mark.parametrize("quantity", [None, math.nan, math.inf, -1, True])
def test_invalid_forecast_quantities_are_rejected_without_zero_imputation(quantity):
    with pytest.raises(ValueError):
        make_forecast([forecast_row(4, 11, quantity)])


@pytest.mark.parametrize("residual", [0, 1e-12, 1e-6])
def test_numerical_work_residual_cannot_issue_a_whole_packaging_unit(residual):
    # The observed stalled campaign issued one unit on each of 78 days despite
    # zero accepted WIP progress. No accepted work must mean zero BOM issue.
    issues = [engine.required_component_quantity(
        engine.production_wip_execution_quantity(residual), 1, "UN")
        for _ in range(78)]
    assert sum(issues) == 0


@pytest.mark.parametrize("work", [1e-5, 8331.034483])
def test_real_fractional_work_is_preserved_inside_an_unfinished_batch(work):
    assert engine.production_wip_execution_quantity(work) == work
    assert engine.required_component_quantity(work, 1, "UN") == math.ceil(work)


@pytest.mark.parametrize("work", [None, math.nan, math.inf, -1])
def test_invalid_work_cannot_be_silently_accepted(work):
    with pytest.raises((ValueError, TypeError)):
        engine.production_wip_execution_quantity(work)


def test_aggregate_boundary_nets_stock_and_existing_orders_before_one_lot_rounding():
    # Source opening state: 590k G available and 1.2M G released later. A
    # hypothetical 1.8M G target leaves 10k G, hence ONE 600k G order.
    assert engine.source_boundary_order_quantity(1_800_000, 590_000, 1_200_000, 600_000) == 600_000
    assert engine.source_boundary_order_quantity(1_800_000, 590_000, 1_800_000, 600_000) == 0


def test_aggregate_boundary_existing_three_receipt_origins_prevent_duplicate_order():
    standard, external, opening = 400_000, 300_000, 500_000
    assert engine.source_boundary_order_quantity(1_200_000, 0, standard + external + opening, 600_000) == 0


def test_aggregate_boundary_neither_creates_zero_need_nor_drops_a_second_required_lot():
    assert engine.source_boundary_order_quantity(0, 0, 0, 600_000) == 0
    assert engine.source_boundary_order_quantity(600_000, 0, 0, 600_000) == 600_000
    assert engine.source_boundary_order_quantity(600_001, 0, 0, 600_000) == 1_200_000


@pytest.mark.parametrize("arguments", [
    (math.nan, 0, 0, 600_000), (1, -1, 0, 600_000),
    (1, 0, math.inf, 600_000), (1, 0, 0, 0), (1, 0, 0, -1),
])
def test_aggregate_boundary_rejects_unknown_or_invalid_inputs(arguments):
    with pytest.raises((ValueError, TypeError)):
        engine.source_boundary_order_quantity(*arguments)


@pytest.mark.parametrize('cover,bridge,expected', [(125, 66, 125), (92, 14, 92), (5, 9, 5), (0, 0, 0)])
def test_initial_order_quantity_does_not_shorten_forecast_cover(cover, bridge, expected):
    assert engine.mrp_opening_order_cover_days(cover, bridge, quantity_netting_only=True) == expected


@pytest.mark.parametrize('cover,bridge,expected', [(125, 66, 59), (92, 14, 78), (5, 9, 0), (0, 0, 0)])
def test_historical_mode_retains_its_previous_cover_convention(cover, bridge, expected):
    assert engine.mrp_opening_order_cover_days(cover, bridge, quantity_netting_only=False) == expected


def test_initial_receipt_changes_stock_position_without_discounting_need_twice():
    # Hand-computable scenario, not industrial observed quantities: 10,000
    # units/day for 125 days; 200,000 on hand, 629,000 already ordered.
    days = engine.mrp_opening_order_cover_days(125, 66, quantity_netting_only=True)
    target = 10_000 * days
    before_receipt = max(0, target - 200_000 - 629_000)
    after_receipt = max(0, target - 829_000)
    assert target == 1_250_000
    assert before_receipt == after_receipt == 421_000


@pytest.mark.parametrize('cover,bridge', [
    (None, 1), (1, None), (math.nan, 1), (1, math.nan),
    (math.inf, 1), (1, math.inf), (-1, 1), (1, -1),
])
def test_invalid_cover_is_rejected_instead_of_imputed(cover, bridge):
    with pytest.raises(ValueError):
        engine.mrp_opening_order_cover_days(cover, bridge, quantity_netting_only=True)


def opening_pack_case(orders=((1, 0, 2, 200),), *, closed=()):
    """Memory-only explicit BOM/OF; no source file is created or changed."""
    pair = ("M-test", "PF")
    units = {"PF": "UN", "tube": "UN", "box": "UN", "carton": "UN", "bulk": "KG"}
    components = [(pair[0], item) for item in ("tube", "box", "carton")]
    payload = dict(schema_version=1, mode="pack_at_source_G_v1", late_release_policy="source_G_I_workdays_v1",
        rows=[dict(node_id=pair[0], item_id=pair[1], pack_components=[dict(item_id=p[1], role="Pack",
            source_file="BOM source", source_cells=f"BOM!D{i}") for i, p in enumerate(components, 1)])])
    rows = [dict(node_id=pair[0], item_id=pair[1], order_type="opening_production_order",
        source_marker=f"OF:{row}", source_row=row, physical_delivery_day=g, arrival_day=i,
        receipt_qty=quantity) for row, g, i, quantity in orders]
    kwargs = dict(opening_orders=rows, nodes=[dict(id=pair[0], processes=[dict(
        outputs=[dict(item_id=pair[1])], capacity=dict(max_rate=150))])],
        bom_by_pair={pair: [(components[0], 1.0), (components[1], 1.0), (components[2], 0.011),
                           ((pair[0], "bulk"), 0.035)]}, item_unit_map=units,
        origin_date="2025-01-01", dated_execution=True, lot_trace=True, bom_issue_mode="wip",
        closed_days_by_pair={pair: frozenset(closed)})
    return pair, units, components, payload, kwargs


def opening_pack_runtime(orders=((1, 0, 2, 200),), *, carton=10, boxes=500, closed=()):
    pair, units, components, payload, kwargs = opening_pack_case(orders, closed=closed)
    book = engine.OpeningPackOrders(payload, **kwargs)
    stock = defaultdict(float, {components[0]: 500, components[1]: boxes,
                               components[2]: carton, (pair[0], "bulk"): 50})
    ledger = engine.LotLedger(enabled=True, scenario_id="memory_pack")
    ledger.seed_opening_stock(day=-1, stock=stock, item_unit_map=units)
    availability = engine.OpeningPurchaseAvailability([], item_unit_map=units)
    pipeline, commitments, in_transit = defaultdict(list), defaultdict(dict), defaultdict(float)
    calendar = engine.DatedReceiptCalendar()
    def consume(component_pair, qty, fragment):
        assert qty <= stock[component_pair]
        stock[component_pair] -= qty
        return ledger.consume(day=fragment["day"], node_id=component_pair[0], item_id=component_pair[1],
            qty=qty, event_type="production_consume", source_id=fragment["fragment_id"], uom=units[component_pair[1]])
    def materialize(source, fragment, parents):
        return engine.materialize_opening_pack_fragment(source, fragment, parents, ledger=ledger,
            availability=availability, stock=stock, pipeline=pipeline, in_transit=in_transit,
            commitments=commitments, receipt_calendar=calendar)
    def execute(day, capacity):
        effective = engine.manufacturing_day_capacity(capacity, day=day, closed_days=frozenset(closed))
        return engine.execute_opening_pack_fifo(book, pair, day=day, capacity=effective,
            available={p: stock[p] for p in components}, consume=consume, materialize=materialize)
    return dict(pair=pair, components=components, book=book, stock=stock, ledger=ledger,
        availability=availability, pipeline=pipeline, commitments=commitments, in_transit=in_transit,
        calendar=calendar, execute=execute)


def test_opening_pack_no_carton_no_pf_no_debit_and_future_requirement_retained():
    state = opening_pack_runtime(carton=0)
    before = dict(state["stock"])
    remaining, fragments = state["execute"](0, 150)
    assert (remaining, fragments) == (150, [])
    assert dict(state["stock"]) == before
    assert not state["book"].fragments and not state["pipeline"]
    assert not [e for e in state["ledger"].event_rows if e["event_type"] == "production_consume"]
    needs, firms = state["book"].planning(0)
    assert [needs[p][0].qty for p in state["components"]] == [200, 200, 3]
    assert all(r.due_day == 1 for rows in needs.values() for r in rows)
    assert sum(r.qty for r in firms[state["pair"]]) == 200


def test_opening_pack_integer_fragments_cumulative_ceil_and_exact_genealogy():
    state = opening_pack_runtime()
    assert state["execute"](0, 100)[0] == 0
    assert state["execute"](1, 100)[0] == 0
    fragments = state["book"].fragments
    assert [f["qty"] for f in fragments] == [100, 100]
    assert [f["components"][state["components"][2]] for f in fragments] == [2, 1]
    assert [f["available_day"] for f in fragments] == [2, 5]  # Wed->Fri; Thu + 2 weekdays->Mon.
    assert [state["stock"][p] for p in state["components"]] == [300, 300, 7]
    assert state["stock"][state["pair"][0], "bulk"] == 50  # Never issued a second time.
    assert state["stock"][state["pair"]] == 0
    assert state["availability"].held_by_pair[state["pair"]] == 200
    children = [state["ledger"].lots[f["lot_id"]] for f in fragments]
    assert len({child["business_batch_id"] for child in children}) == 1
    assert all(child["trace_status"] == "partially_traced_mixed_occurrence" for child in children)
    for f in fragments:
        debits = [e for e in state["ledger"].event_rows
                  if e["event_type"] == "production_consume" and e["source_id"] == f["fragment_id"]]
        assert {e["item_id"]: e["qty"] for e in debits} == {p[1]: q for p, q in f["components"].items()}
        links = [r for r in state["ledger"].genealogy_rows if r["child_lot_id"] == f["lot_id"]]
        assert sum(r["parent_qty"] for r in links) == sum(f["components"].values())
    assert engine.opening_pack_order_status(state["book"], "OF:1", last_day=1)["actual_available_day"] == ""
    assert engine.opening_pack_order_status(state["book"], "OF:1", last_day=5)["actual_available_day"] == 5


def test_opening_pack_two_orders_share_fifo_stock_and_daily_capacity_with_new_work():
    state = opening_pack_runtime(((2, 0, 2, 80), (1, 0, 2, 100)), boxes=130)
    remaining, fragments = state["execute"](0, 150)
    assert [(f["order_id"], f["qty"]) for f in fragments] == [("OF:1", 100), ("OF:2", 30)]
    assert remaining == 20
    assert state["stock"][state["components"][1]] == 0
    # The caller passes only this remainder into normal new fabrication.
    new_work = min(remaining, 50)
    assert sum(f["qty"] for f in fragments) + new_work == 150
    assert state["book"].rows["OF:2"]["completed_qty"] == 30
    assert engine.opening_pack_order_status(state["book"], "OF:2", last_day=10)["actual_physical_receipt_day"] == ""


def test_opening_pack_held_fragment_plus_pending_residual_netted_once():
    state = opening_pack_runtime()
    unrelated = engine.FirmReceipt("other", 8, 17)
    state["commitments"][state["pair"]][unrelated.receipt_id] = (8, 17, "production")
    state["execute"](0, 100)
    needs = engine.sync_opening_pack_commitments(state["book"], decision_day=0, commitments=state["commitments"])
    snapshot = deepcopy(state["commitments"])
    engine.sync_opening_pack_commitments(state["book"], decision_day=0, commitments=state["commitments"])
    assert state["commitments"] == snapshot
    firm = state["commitments"][state["pair"]]
    assert sorted((r[1], r[2]) for k, r in firm.items() if k != "other") == [(100, "held"), (100, "production")]
    assert sum(r[1] for r in firm.values()) == 217
    assert [needs[p][0].qty for p in state["components"]] == [100, 100, 1]
    assert set(needs) == set(state["components"])


def test_opening_pack_closure_prevents_all_conditioning_and_shifts_pending_need():
    state = opening_pack_runtime(closed=(0, 1))
    assert state["execute"](0, 150) == (0, [])
    assert not state["book"].fragments
    needs, firms = state["book"].planning(0)
    assert {r.due_day for rows in needs.values() for r in rows} == {2}
    assert firms[state["pair"]][0].available_day == 6  # Fri 3 Jan + 2 weekdays = Tue 7 Jan.


def test_opening_pack_source_I_preserved_and_late_actual_source_workdays_shifted():
    origin = date(2025, 1, 1)
    g = (date(2025, 4, 28) - origin).days
    i = (date(2025, 5, 14) - origin).days
    pair, units, components, payload, kwargs = opening_pack_case(((74, g, i, 100),))
    book = engine.OpeningPackOrders(payload, **kwargs)
    row = book.rows["OF:74"]
    assert row["source_workdays"] == 12  # Actual G/I, not the announced 10-day duration.
    assert book.availability(row, g) == i
    assert origin + timedelta(days=book.availability(row, g + 1)) == date(2025, 5, 15)
    assert book.candidate(row, day=g-1, capacity=100, available={p: 500 for p in components})["qty"] == 0


def test_opening_pack_no_materialization_without_exact_pack_parents():
    state = opening_pack_runtime()
    row = state["book"].rows["OF:1"]
    fragment = state["book"].candidate(row, day=0, capacity=100, available=state["stock"])
    before = len(state["ledger"].lots)
    with pytest.raises(RuntimeError, match="exact consumed Pack"):
        engine.materialize_opening_pack_fragment(row, fragment, [], ledger=state["ledger"],
            availability=state["availability"], stock=state["stock"], pipeline=state["pipeline"],
            in_transit=state["in_transit"], commitments=state["commitments"])
    assert len(state["ledger"].lots) == before


def test_opening_pack_single_release_of_same_lot_without_second_consumption():
    state = opening_pack_runtime(((1, 0, 2, 100),))
    state["execute"](0, 150)
    fragment = state["book"].fragments[0]
    availability = state["availability"]
    qty = fragment["qty"]
    before_debits = sum(e["qty"] for e in state["ledger"].event_rows if e["event_type"] == "production_consume")
    availability.release(fragment["fragment_id"], 2, qty, ledger=state["ledger"])
    assert availability.held_by_pair[state["pair"]] == 0
    assert state["ledger"].lots[fragment["lot_id"]]["qty_remaining"] == 100
    assert sum(e["qty"] for e in state["ledger"].event_rows if e["event_type"] == "production_consume") == before_debits
    with pytest.raises(ValueError, match="single physical receipt"):
        availability.release(fragment["fragment_id"], 2, qty, ledger=state["ledger"])


@pytest.mark.parametrize("fault", ["schema_bool", "no_role", "not_in_bom", "no_capacity", "no_trace", "fractional_pf"])
def test_opening_pack_rejects_ambiguous_roles_and_unphysical_contracts(fault):
    pair, units, components, payload, kwargs = opening_pack_case()
    if fault == "schema_bool": payload["schema_version"] = True
    elif fault == "no_role": payload["rows"][0]["pack_components"][0]["role"] = "unknown"
    elif fault == "not_in_bom": payload["rows"][0]["pack_components"][0]["item_id"] = "absent"
    elif fault == "no_capacity": kwargs["nodes"][0]["processes"][0]["capacity"] = {}
    elif fault == "no_trace": kwargs["lot_trace"] = False
    elif fault == "fractional_pf": kwargs["opening_orders"][0]["receipt_qty"] = 1.5
    with pytest.raises(ValueError):
        engine.OpeningPackOrders(payload, **kwargs)
