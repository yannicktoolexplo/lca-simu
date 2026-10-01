"""Prepare, run and compare isolated component forecasts and purchase policies.

The first MRP vintage estimates other-use shares; optional revisions replace
future demand at its knowledge date. H, J, K and inventory photographs never
fit that demand. Missing weeks stay unknown. Purchase policies are explicit
experiments, not identified ERP rules.
"""
from __future__ import annotations

import argparse
from collections import defaultdict
from copy import deepcopy
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
import subprocess
import sys
import time

ROOT = Path(__file__).resolve().parents[2]
ORIGIN = datetime(2025, 1, 1)
SITES = {'1430': 'M-1430', '1450': 'SDC-1450', '1810': 'M-1810', '1920': 'DC-1920'}
DEFAULT_REFERENCE = ROOT / 'artifacts/testing/mrp_execution_20260928/simulation_v6'
DEFAULT_WORK = ROOT / 'artifacts/testing/shared_materials_20260929/study'


def unit(value):
    value = str(value).strip().upper().rstrip('.')
    if value == 'G':
        return 'KG', .001
    if value in ('UN', 'ZUN', 'UNIT', 'UNITS'):
        return 'UN', 1.
    if value in ('KG', 'M'):
        return value, 1.
    raise ValueError(f'Unsupported unit: {value}')


def sha(path):
    with Path(path).open('rb') as stream:
        return hashlib.file_digest(stream, 'sha256').hexdigest()


def write(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open('x', encoding='utf8') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, allow_nan=False)


def read_first_plan(path):
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    first = None
    records = []
    try:
        for line, row in enumerate(book['Feuille1'].iter_rows(min_row=2, values_only=True), 2):
            vintage = row[0]
            if not isinstance(vintage, datetime):
                raise ValueError(f'Invalid vintage A{line}')
            if first is None or vintage < first:
                first, records = vintage, []
            if vintage != first:
                continue
            canonical, factor = unit(row[5])
            if row[8] is None:
                raise ValueError(f'Missing gross requirement I{line}; absence is not zero')
            qty = float(row[8]) * factor
            if not math.isfinite(qty) or qty < 0:
                raise ValueError(f'Invalid gross requirement I{line}')
            records.append({'item': str(row[1]).zfill(6), 'division': str(row[3]),
                            'type': str(row[2]), 'unit': canonical, 'qty': qty,
                            'day': (row[6] - ORIGIN).days, 'row': line})
    finally:
        book.close()
    return (first - ORIGIN).days, records


def consumption_network(graph):
    """Coefficients in canonical units, including upstream transformation BOMs."""
    producers, states = {}, {}
    for node in graph['nodes']:
        for state in node.get('inventory', {}).get('states', []):
            states[(node['id'], state['item_id'])] = state
        for process in node.get('processes', []):
            outputs = process.get('outputs', [])
            if len(outputs) != 1:
                raise ValueError('Study requires a single output per production process')
            item = outputs[0]['item_id']
            if item in producers:
                raise ValueError(f'Ambiguous production site for {item}')
            producers[item] = (node['id'], process)

    def expand(item, trail=()):
        if item not in producers:
            return {}
        if item in trail:
            raise ValueError('Cyclic BOM')
        node, process = producers[item]
        _, base_factor = unit(process['batch_size_unit'])
        base = float(process['batch_size']) * base_factor
        if base <= 0:
            raise ValueError('Nonpositive BOM base')
        result = defaultdict(float)
        for component in process['inputs']:
            child = component['item_id']
            _, factor = unit(component['ratio_unit'])
            ratio = float(component['ratio_per_batch']) * factor / base
            result[(node, child)] += ratio
            for pair, coefficient in expand(child, (*trail, item)).items():
                result[pair] += ratio * coefficient
        return dict(result)
    return states, producers, expand


def estimate_first_plan(graph, known_day, records, threshold=1.25):
    """Pure estimator. Additional downstream use is deducted from upstream residuals."""
    if threshold <= 1 or not math.isfinite(threshold):
        raise ValueError('Screening threshold must exceed one')
    states, producers, expand = consumption_network(graph)
    start = known_day + 7  # Current bucket may include old backlog; do not invent new demand.
    grouped = defaultdict(list)
    for row in records:
        grouped[(SITES[row['division']], 'item:' + row['item'])].append(row)
    for rows in grouped.values():
        days = [r['day'] for r in rows]
        if len(set(days)) != len(days):
            raise ValueError('Duplicate pair/week in first vintage')
    pf = {pair[1]: rows for pair, rows in grouped.items()
          if rows[0]['type'] == 'PF' and pair[1] in producers}
    coefficients = {item: expand(item) for item in pf}
    direct_pairs = set().union(*(set(c) for c in coefficients.values()))
    # A component output drives its upstream ingredients; visit it first.
    depth = {pair: len(expand(pair[1])) for pair in direct_pairs}
    ordered = sorted(grouped, key=lambda pair: (-depth.get(pair, -1), pair))
    schedules, audit, selected = [], [], []
    for pair in ordered:
        rows = sorted(grouped[pair], key=lambda row: row['day'])
        info = {'pair': rows[0]['item'] + '/' + rows[0]['division'],
                'node_id': pair[0], 'item_id': pair[1], 'unit': rows[0]['unit'],
                'article_type': rows[0]['type'], 'known_day': known_day,
                'status': 'not_selected', 'reason': None}
        audit.append(info)
        roots = [item for item, c in coefficients.items() if pair in c]
        if pair not in states or not roots:
            info['reason'] = 'no_modelled_consumption_at_this_site_or_no_stock_state'
            continue
        canonical, engine_factor = unit(states[pair]['uom'])
        if canonical != rows[0]['unit'] or any(r['unit'] != canonical for r in rows):
            raise ValueError(f'Unit disagreement for {pair}')
        ends = []
        for item in roots:
            days = sorted(r['day'] for r in pf[item] if known_day <= r['day'] < 365)
            if not days or days != list(range(known_day, days[-1] + 1, 7)):
                ends = []
                break
            ends.append(min(365, days[-1] + 7))
        if not ends:
            info['reason'] = 'finished_product_forecast_has_missing_weeks'
            continue
        end = min(ends)
        def sum_in_window(series):
            return sum(r['qty'] * min(7, end - r['day']) / 7
                       for r in series if start <= r['day'] < end)
        own_by_pf = {item: sum_in_window(pf[item]) * coefficients[item][pair]
                     for item in roots}
        own = sum(own_by_pf.values())
        induced = 0.
        for child in selected:
            coefficient = expand(child['item_id']).get(pair, 0.)
            for row in child['canonical_schedule']:
                overlap = max(0, min(end, row['day'] + row['period_days']) - max(start, row['day']))
                induced += coefficient * row['qty'] * overlap / row['period_days']
        total = sum_in_window(rows)
        represented = own + induced
        ratio = total / represented if represented > 0 else None
        info.update({'horizon_start_day': start, 'horizon_end_exclusive': end,
                     'source_reported_requirement': total, 'own_bom_requirement': own,
                     'own_requirement_by_pf': own_by_pf, 'already_induced_by_downstream': induced,
                     'source_to_represented_ratio': ratio,
                     'reported_source_weeks': sum(start <= r['day'] < end for r in rows),
                     'source_cells': [f"I{r['row']}" for r in rows if start <= r['day'] < end]})
        if ratio is None or ratio <= threshold or info['reported_source_weeks'] < 3:
            info['reason'] = 'no_sufficient_positive_residual_on_full_common_horizon'
            continue
        fraction = (total - represented) / total
        info.update(status='estimated_complement', reason='first_vintage_excess_after_full_horizon_bom_and_upstream_deduction',
                    estimated_other_use_fraction=fraction,
                    shared_use_status='user_suspected' if info['pair'] in ('042342/1430', '002612/1810', '007923/1810') else 'hypothesis_only')
        canonical_schedule = []
        for row in rows:
            if not start <= row['day'] < end:
                continue
            days = min(7, end - row['day'])
            theoretical = row['qty'] * fraction * days / 7
            physical = theoretical / engine_factor
            if canonical == 'UN':
                physical = math.floor(physical + .5)
            if physical <= 0:
                continue
            schedules.append({'demand_id': f"shared:{row['item']}:{row['division']}:{row['row']}",
                              'node_id': pair[0], 'item_id': pair[1], 'known_day': known_day,
                              'period_start_day': row['day'], 'period_days': days,
                              'qty': physical, 'uom': states[pair]['uom'],
                              'source_file': 'Flow_Data_MRP_results.xlsx',
                              'source_cells': f"Feuille1!I{row['row']}",
                              'estimation_basis': json.dumps({'method': 'first_plan_aggregate_residual_share',
                                                   'fraction': fraction, 'unrounded_canonical_qty': theoretical,
                                                   'screening_ratio': threshold, 'evidence_pair': info['pair']}, sort_keys=True)})
            canonical_schedule.append({'day': row['day'], 'period_days': days, 'qty': physical * engine_factor})
        selected.append({'item_id': pair[1], 'canonical_schedule': canonical_schedule})
        info['scheduled_canonical_qty'] = sum(r['qty'] for r in canonical_schedule)
    return schedules, sorted(audit, key=lambda row: row['pair'])


def protected(plan):
    for path, expected in plan['protected_inputs'].items():
        if sha(path) != expected:
            raise ValueError(f'Protected input changed: {path}')


def revision_series(records, vintages, pair, fraction, *, planning_horizon_days=None, output_uom='KG'):
    """Replace future forecasts; never inject the current, possibly overdue bucket."""
    if planning_horizon_days is not None and (type(planning_horizon_days) is not int or planning_horizon_days != 364):
        raise ValueError('Rolling source forecasts require exactly 52 weeks')
    if output_uom not in ('KG', 'G', 'UN'):
        raise ValueError('Revised forecasts require KG, G or UN output')
    output_factor = 1000. if output_uom == 'G' else 1.
    item, division = pair.split('/')
    versions = []
    for known in sorted(set(vintages)):
        rows = []
        for record in records:
            end = known + planning_horizon_days + 1 if planning_horizon_days is not None else 365
            if record['known_day'] != known or record['day'] <= known or record['day'] >= end:
                continue
            if record['unit'] != ('UN' if output_uom == 'UN' else 'KG'):
                raise ValueError('Revised forecasts require matching canonical mass or unit quantities')
            length = 7 if planning_horizon_days is not None else min(7, 365 - record['day'])
            quantity = record['qty'] * fraction * length / 7 * output_factor
            basis = {'method': 'fixed_first_plan_share_revised_future', 'fraction': fraction}
            if output_uom == 'UN':
                # Keep the fractional forecast as evidence; the complementary
                # physical weekly scenario follows estimate_first_plan rounding.
                basis['unrounded_canonical_qty'] = quantity
                quantity = math.floor(quantity + .5)
            rows.append({'demand_id': f'shared-revision:{item}:{division}:{record["day"]}',
                         'period_start_day': record['day'], 'period_days': length,
                         'qty': quantity,
                         'source_file': 'Flow_Data_MRP_results.xlsx',
                         'source_cells': f'Feuille1!I{record["row"]}',
                         'estimation_basis': json.dumps(basis, sort_keys=True)})
        if len({r['period_start_day'] for r in rows}) != len(rows):
            raise ValueError('Duplicate target week in a revision')
        versions.append({'vintage_id': f'mrp:{known}', 'known_day': known,
                         'rows': sorted(rows, key=lambda r: r['period_start_day'])})
    return {'node_id': SITES[division], 'item_id': 'item:' + item, 'uom': output_uom,
            'period_anchor_day': 4, 'period_days': 7, 'repeat_period_days': 0,
            'current_bucket_policy': 'freeze_previous_exclude_current_vintage',
            'missing_future_policy': 'unprovided_not_observed_zero', 'versions': versions}


def prepare_revisions(args):
    """001848/1810 only: dated knowledge experiment, with frozen witnesses."""
    import openpyxl
    reference = args.reference.resolve()
    old = json.loads((reference / 'plan.json').read_text(encoding='utf8'))
    protected(old)
    graph_path = reference / 'shared_components_graph.json'
    graph = json.loads(graph_path.read_text(encoding='utf8'))
    meta = graph['meta']['external_component_demands']
    pair = '001848/1810'
    info = next(r for r in meta['estimation_audit'] if r['pair'] == pair)
    fraction = info['estimated_other_use_fraction']
    workbook = ROOT / 'data/source/Flow_Data_MRP_results.xlsx'
    records, vintages = [], set()
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        for line, row in enumerate(book['Feuille1'].iter_rows(min_row=2, values_only=True), 2):
            known = (row[0] - ORIGIN).days
            if not 0 <= known < 365:
                continue
            vintages.add(known)
            if str(row[1]).zfill(6) != '001848' or str(row[3]) != '1810':
                continue
            canonical, factor = unit(row[5])
            if row[8] is None or not math.isfinite(float(row[8])) or float(row[8]) < 0:
                raise ValueError(f'Invalid requirement I{line}')
            records.append({'known_day': known, 'day': (row[6] - ORIGIN).days,
                            'qty': float(row[8]) * factor, 'unit': canonical, 'row': line})
    finally:
        book.close()
    series = revision_series(records, vintages, pair, fraction)
    meta['rows'] = [r for r in meta['rows'] if (r['node_id'], r['item_id']) != ('M-1810', 'item:001848')]
    meta.update(schema_version=2, versioned_series=[series])
    meta['assumptions'] += [
        '001848/1810 only: successive MRP vintages replace future complementary forecasts.',
        'Its first-plan estimated share stays fixed; later inventories, H, J and K never fit demand.',
        'Current weekly demand freezes from the last vintage strictly before the week starts.',
        'Missing rows are unknown: no supported additional demand is injected; older rows are not carried over.',
        '001848 revisions cover 2025 only and do not repeat. Other 18 shared-use hypotheses are unchanged.',
    ]
    graph_file = args.workdir / 'shared_components_graph.json'
    write(graph_file, graph)
    write(args.workdir / 'scope_audit.json', {'pair': pair, 'fraction': fraction,
          'version_count': len(series['versions']), 'source_rows': records,
          'basis': info, 'independent_stock_fit': False})
    commands = {}
    for name, previous in [('nominal', 'nominal'), ('usages_complementaires', 'usages_complementaires'),
                           ('usages_figes_2025', 'usages_complementaires'),
                           ('revisions_001848', 'usages_complementaires')]:
        command = list(old['commands'][previous])
        command[0] = sys.executable
        command[command.index('--output-dir') + 1] = str(args.workdir / name)
        if name == 'revisions_001848':
            command[command.index('--input') + 1] = str(graph_file)
        if name in ('usages_figes_2025', 'revisions_001848'):
            command[command.index('--days') + 1] = '365'
        commands[name] = command
    inputs = dict(old['protected_inputs'])
    inputs.update({str(p): sha(p) for p in [graph_path, graph_file, workbook]})
    plan = {'schema_version': 2, 'experiment': 'weekly_revisions_001848', 'reference': str(reference),
            'commands': commands, 'protected_inputs': inputs, 'estimator_sha256': sha(Path(__file__)),
            'csv_witnesses': {name: str(reference / name / 'data') for name in ('nominal', 'usages_complementaires')},
            'labels': {'nominal': 'Nominal conservé · calcul sur cinq ans',
                       'usages_complementaires': 'Témoin figé · calcul sur cinq ans',
                       'usages_figes_2025': 'Plan du 5 janvier figé · calcul sur 2025',
                       'revisions_001848': '001848 Avène · besoins révisés chaque semaine'}}
    protected(plan)
    write(args.workdir / 'plan.json', plan)
    print(json.dumps({'pair': pair, 'versions': len(series['versions']), 'fixed_fraction': fraction}), flush=True)


def prepare_sourcing(args):
    """Confirmed 001848 vendor roles; dated backup trigger remains an experiment."""
    import openpyxl
    reference = args.reference.resolve()
    old = json.loads((reference / 'plan.json').read_text(encoding='utf8'))
    if old.get('experiment') != 'weekly_revisions_001848':
        raise ValueError('Sourcing requires the revised 001848 reference study')
    protected(old)
    graph_path = reference / 'shared_components_graph.json'
    graph = json.loads(graph_path.read_text(encoding='utf8'))
    if graph['meta'].get('sourcing_policy'):
        raise ValueError('Reference already contains a sourcing policy')
    workbook = ROOT / 'data/source/268091.xlsx'
    offers = []
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        for line, row in enumerate(book['FIA'].iter_rows(min_row=2, values_only=True), 2):
            if str(row[0]).zfill(6) != '001848':
                continue
            canonical, factor = unit(row[7])
            base = float(row[3])
            if canonical != 'KG' or row[4] != 'EUR' or base <= 0:
                raise ValueError('001848 source price is not comparable in EUR/KG')
            offers.append({'supplier_id': 'SDC-' + str(row[1]),
                           'purchase_unit_cost': float(row[2]) / base / factor,
                           'price_uom': canonical, 'currency': row[4],
                           'source_file': workbook.name, 'source_cells': f'FIA!C{line}:H{line}',
                           'source_delivery_days': float(row[5]),
                           'source_standard_qty': float(row[6]) * factor,
                           'lot_semantics': 'standard_quantity_not_independently_proven_minimum_or_multiple'})
    finally:
        book.close()
    expected = {'SDC-VD0951020A': (1.58, 56., 6000.), 'SDC-VD0519670A': (4.2, 21., 4000.)}
    found = {r['supplier_id']: (r['purchase_unit_cost'], r['source_delivery_days'], r['source_standard_qty'])
             for r in offers}
    if found != expected or len(offers) != 2:
        raise ValueError('Source offers differ from the confirmed 001848 roles; review required')
    graph['meta']['sourcing_policy'] = {'schema_version': 1, 'rows': [{
        'policy_id': '001848_primary_backup', 'node_id': 'M-1810', 'item_id': 'item:001848',
        'uom': 'KG', 'currency': 'EUR', 'policy': 'cheapest_purchase_primary_with_dated_backup',
        'confirmation': 'user_confirmed_roles_only', 'primary_supplier_id': 'SDC-VD0951020A',
        'backup_supplier_ids': ['SDC-VD0519670A'], 'offers': offers,
        'backup_trigger': 'physical_shortage_before_primary',
        'bridge_surplus_policy': 'retain_firm_and_disclose_surplus',
        'trigger_evidence': 'experimental_convention_not_identified_in_industrial_flow',
    }]}
    graph_file = args.workdir / 'sourcing_graph.json'
    write(graph_file, graph)
    commands = {}
    for name, prior in [('nominal', 'nominal'), ('reference_revisions', 'revisions_001848'),
                        ('regles_fournisseurs', 'revisions_001848')]:
        command = list(old['commands'][prior])
        command[0] = sys.executable
        command[command.index('--output-dir') + 1] = str(args.workdir / name)
        if name == 'regles_fournisseurs':
            command[command.index('--input') + 1] = str(graph_file)
        commands[name] = command
    inputs = dict(old['protected_inputs'])
    inputs.update({str(p): sha(p) for p in [graph_path, graph_file, workbook]})
    plan = {'schema_version': 2, 'experiment': 'sourcing_001848', 'reference': str(reference),
            'commands': commands, 'protected_inputs': inputs, 'estimator_sha256': sha(Path(__file__)),
            'csv_witnesses': {'nominal': str(reference / 'nominal/data'),
                              'reference_revisions': str(reference / 'revisions_001848/data')},
            'labels': {'nominal': 'Nominal historique · témoin sur cinq ans',
                       'reference_revisions': 'Achats précédents · besoins révisés · 365 jours',
                       'regles_fournisseurs': '001848 principal et secours · essai sur 365 jours'}}
    protected(plan)
    write(args.workdir / 'plan.json', plan)
    print(json.dumps({'pair': '001848/1810', 'offers': len(offers), 'other_sourcing_rules_unchanged': True}), flush=True)


def prepare_batching(args):
    """001757: isolate forecast revisions from purchase rounding and grouping."""
    import openpyxl
    reference = args.reference.resolve()
    old = json.loads((reference / 'plan.json').read_text(encoding='utf8'))
    if old.get('experiment') != 'sourcing_001848':
        raise ValueError('Batching requires the conserved sourcing study reference')
    protected(old)
    graph_path = reference / 'sourcing_graph.json'
    graph = json.loads(graph_path.read_text(encoding='utf8'))
    if graph['meta'].get('procurement_batching'):
        raise ValueError('The reference already contains a batching experiment')
    meta = graph['meta']['external_component_demands']
    info = next(r for r in meta['estimation_audit'] if r['pair'] == '001757/1810')
    fraction = info['estimated_other_use_fraction']
    workbook = ROOT / 'data/source/Flow_Data_MRP_results.xlsx'
    records, vintages = [], set()
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        for line, row in enumerate(book['Feuille1'].iter_rows(min_row=2, values_only=True), 2):
            known = (row[0] - ORIGIN).days
            if not 0 <= known < 365:
                continue
            vintages.add(known)
            if str(row[1]).zfill(6) != '001757' or str(row[3]) != '1810':
                continue
            canonical, factor = unit(row[5])
            if row[8] is None or not math.isfinite(float(row[8])) or float(row[8]) < 0:
                raise ValueError(f'Invalid 001757 requirement I{line}')
            records.append({'known_day': known, 'day': (row[6] - ORIGIN).days,
                            'qty': float(row[8]) * factor, 'unit': canonical, 'row': line})
    finally:
        book.close()
    series = revision_series(records, vintages, '001757/1810', fraction)
    if any((r['node_id'], r['item_id']) == ('M-1810', 'item:001757')
           for r in meta.get('versioned_series', [])):
        raise ValueError('001757 revisions already present in the reference')
    meta['rows'] = [r for r in meta['rows'] if (r['node_id'], r['item_id']) != ('M-1810', 'item:001757')]
    meta['versioned_series'].append(series)
    meta['assumptions'].append('001757 additionally revised with the same known-vintage rule; its first-plan other-use fraction stays fixed.')
    configurations = [('previsions_actualisees', None, None), ('lot_1000', 1000, 0),
                      ('couverture_14j', 100, 14), ('couverture_28j', 100, 28),
                      ('lot1000_couverture14j', 1000, 14), ('lot1000_couverture28j', 1000, 28)]
    commands, inputs = {}, dict(old['protected_inputs'])
    inputs.update({str(p): sha(p) for p in [graph_path, workbook]})
    for name, previous in [('nominal', 'nominal'), ('reference_figee', 'regles_fournisseurs')]:
        command = list(old['commands'][previous])
        command[0] = sys.executable
        command[command.index('--output-dir') + 1] = str(args.workdir / name)
        commands[name] = command
    for name, multiple, grouping in configurations:
        variant = deepcopy(graph)
        if multiple is not None:
            variant['meta']['procurement_batching'] = {'schema_version': 1, 'rows': [{
                'policy_id': f'001757_round{multiple}_window{grouping}', 'node_id': 'M-1810',
                'item_id': 'item:001757', 'uom': 'KG', 'rounding_multiple_qty': multiple,
                'grouping_days': grouping, 'semantics': 'first_uncovered_requirement_forward_window',
                'status': 'candidate_sensitivity',
                'evidence_note': 'Test hypothesis, not an identified industrial rule. FIA standard is 100 KG; repeated H forecasts mostly use 1000 KG increments.',
            }]}
        graph_file = args.workdir / f'{name}_graph.json'
        write(graph_file, variant)
        inputs[str(graph_file)] = sha(graph_file)
        command = list(old['commands']['regles_fournisseurs'])
        command[0] = sys.executable
        command[command.index('--output-dir') + 1] = str(args.workdir / name)
        command[command.index('--input') + 1] = str(graph_file)
        command[command.index('--days') + 1] = '365'
        commands[name] = command
    labels = {'nominal': 'Nominal historique · témoin sur cinq ans',
              'reference_figee': '001757 · prévisions figées · référence 365 jours',
              'previsions_actualisees': 'Prévisions actualisées seules · pas 100 kg',
              'lot_1000': 'Prévisions actualisées · arrondi 1 000 kg',
              'couverture_14j': 'Prévisions actualisées · 14 jours · pas 100 kg',
              'couverture_28j': 'Prévisions actualisées · 28 jours · pas 100 kg',
              'lot1000_couverture14j': 'Prévisions actualisées · 14 jours · arrondi 1 000 kg',
              'lot1000_couverture28j': 'Prévisions actualisées · 28 jours · arrondi 1 000 kg'}
    plan = {'schema_version': 2, 'experiment': 'batching_001757', 'reference': str(reference),
            'commands': commands, 'labels': labels, 'protected_inputs': inputs,
            'estimator_sha256': sha(Path(__file__)),
            'csv_witnesses': {'nominal': str(reference / 'nominal/data'),
                              'reference_figee': str(reference / 'regles_fournisseurs/data')}}
    protected(plan)
    write(args.workdir / 'scope_audit.json', {'pair': '001757/1810', 'fraction': fraction,
          'version_count': len(series['versions']), 'source_rows': records, 'basis': info,
          'initial_state_safety_and_existing_001848_policy_unchanged': True,
          'configurations': configurations, 'stock_used_to_fit_demand': False})
    write(args.workdir / 'plan.json', plan)
    print(json.dumps({'pair': '001757/1810', 'versions': len(series['versions']),
                      'fixed_fraction': fraction, 'runs': len(commands)}), flush=True)


def prepare_rolling52(args):
    """Observe 2025 while retaining each known forecast's 52-week projection."""
    import openpyxl
    reference = args.reference.resolve()
    old = json.loads((reference / 'plan.json').read_text(encoding='utf8'))
    if old.get('experiment') != 'batching_001757':
        raise ValueError('Rolling horizon comparison requires the conserved batching study')
    protected(old)
    graph_path = reference / 'previsions_actualisees_graph.json'
    graph = json.loads(graph_path.read_text(encoding='utf8'))
    meta = graph['meta']['external_component_demands']
    pairs = {(r['node_id'], r['item_id']) for r in meta['versioned_series']}
    if pairs != {('M-1810', 'item:001757'), ('M-1810', 'item:001848')}:
        raise ValueError('Unexpected revised component scope in the reference')
    records, vintages = defaultdict(list), set()
    workbook = ROOT / 'data/source/Flow_Data_MRP_results.xlsx'
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        for line, row in enumerate(book['Feuille1'].iter_rows(min_row=2, values_only=True), 2):
            known = (row[0] - ORIGIN).days
            if not 0 <= known < 365:
                continue
            vintages.add(known)
            key = (SITES.get(str(row[3])), 'item:' + str(row[1]).zfill(6))
            if key not in pairs:
                continue
            canonical, factor = unit(row[5])
            if row[8] is None or not math.isfinite(float(row[8])) or float(row[8]) < 0:
                raise ValueError(f'Invalid rolling requirement I{line}')
            records[key].append({'known_day': known, 'day': (row[6] - ORIGIN).days,
                                 'qty': float(row[8]) * factor, 'unit': canonical, 'row': line})
    finally:
        book.close()
    extended, audits = [], []
    for old_series in meta['versioned_series']:
        key = (old_series['node_id'], old_series['item_id'])
        pair = key[1].split(':')[1] + '/1810'
        fraction = next(r['estimated_other_use_fraction'] for r in meta['estimation_audit'] if r['pair'] == pair)
        series = revision_series(records[key], vintages, pair, fraction, planning_horizon_days=364)
        extended.append(series)
        audits.append({'pair': pair, 'fixed_fraction': fraction, 'versions': len(series['versions']),
                       'rows': sum(len(v['rows']) for v in series['versions']),
                       'rows_starting_2026': sum(r['period_start_day'] >= 365 for v in series['versions'] for r in v['rows']),
                       'latest_forecast_day': max(r['period_start_day'] + r['period_days'] - 1 for v in series['versions'] for r in v['rows'])})
    commands, inputs = {}, dict(old['protected_inputs'])
    inputs[str(workbook)] = sha(workbook)
    for name, previous in [('nominal', 'nominal'), ('reference_figee', 'previsions_actualisees')]:
        command = list(old['commands'][previous])
        command[0] = sys.executable
        command[command.index('--output-dir') + 1] = str(args.workdir / name)
        commands[name] = command
    for name in old['commands']:
        if name in ('nominal', 'reference_figee'):
            continue
        prior_graph = reference / f'{name}_graph.json'
        variant = json.loads(prior_graph.read_text(encoding='utf8'))
        variant['meta']['mrp_planning_horizon_days'] = 364
        calendar = variant['meta']['external_component_demands']
        calendar['schema_version'] = 3
        calendar['versioned_series'] = deepcopy(extended)
        calendar['assumptions'].append('Keep the 52-week future projection of each known vintage, including 2026; physical execution still stops after 365 days. Other fixed component calendars retain their previous scope.')
        graph_file = args.workdir / f'{name}_graph.json'
        write(graph_file, variant)
        inputs[str(prior_graph)] = sha(prior_graph)
        inputs[str(graph_file)] = sha(graph_file)
        command = list(old['commands'][name])
        command[0] = sys.executable
        command[command.index('--output-dir') + 1] = str(args.workdir / name)
        command[command.index('--input') + 1] = str(graph_file)
        commands[name] = command
    labels = dict(old['labels'])
    labels['reference_figee'] = 'Ancien horizon borné à décembre · prévisions actualisées'
    labels['previsions_actualisees'] = 'Horizon glissant 52 semaines · pas 100 kg'
    for name in commands:
        if name not in ('nominal', 'reference_figee', 'previsions_actualisees'):
            labels[name] = labels[name].replace('Prévisions actualisées', 'Horizon 52 semaines')
    plan = {'schema_version': 2, 'experiment': 'rolling52_001757', 'reference': str(reference),
            'commands': commands, 'labels': labels, 'protected_inputs': inputs,
            'estimator_sha256': sha(Path(__file__)),
            'csv_witnesses': {'nominal': str(reference / 'nominal/data'),
                              'reference_figee': str(reference / 'previsions_actualisees/data')}}
    protected(plan)
    write(args.workdir / 'scope_audit.json', {'planning_horizon_days': 364, 'execution_days': 365,
          'source_versions_known_in': 2025, 'extended_series': audits,
          'other_17_fixed_calendars_unchanged': True, 'no_source_receipt_or_stock_replay': True})
    write(args.workdir / 'plan.json', plan)
    print(json.dumps({'runs': len(commands), 'extended_series': audits}), flush=True)


def prepare_policy(args):
    """Sequential, isolated 693055 policy candidates, then 001757 rounding."""
    import openpyxl
    reference = args.reference.resolve()
    old = json.loads((reference / 'plan.json').read_text(encoding='utf8'))
    if old.get('experiment') != 'rolling52_001757':
        raise ValueError('Policy study requires the conserved rolling52 comparison')
    protected(old)
    graph_path = reference / 'previsions_actualisees_graph.json'
    graph = json.loads(graph_path.read_text(encoding='utf8'))
    if graph['meta'].get('internal_component_policy') or graph['meta'].get('procurement_batching'):
        raise ValueError('Policy reference already contains a candidate override')
    meta = graph['meta']['external_component_demands']
    info = next(r for r in meta['estimation_audit'] if r['pair'] == '693055/1810')
    fraction = info['estimated_other_use_fraction']
    state = next(s for n in graph['nodes'] if n['id'] == 'M-1810'
                 for s in n['inventory']['states'] if s['item_id'] == 'item:693055')
    if state['uom'] != 'G' or graph['meta'].get('mrp_planning_horizon_days') != 364:
        raise ValueError('Expected G stock and rolling52 reference')
    workbook = ROOT / 'data/source/Flow_Data_MRP_results.xlsx'
    inventory = ROOT / 'data/source/Flow_Data_Inventory_and_Replenishment_rules.xlsx'
    records, vintages, receipt_days = [], set(), set()
    book = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        for line, row in enumerate(book['Feuille1'].iter_rows(min_row=2, values_only=True), 2):
            known = (row[0] - ORIGIN).days
            if not 0 <= known < 365:
                continue
            vintages.add(known)
            if str(row[1]).zfill(6) != '693055' or str(row[3]) != '1810':
                continue
            canonical, factor = unit(row[5])
            if row[8] is None or not math.isfinite(float(row[8])) or float(row[8]) < 0:
                raise ValueError(f'Invalid 693055 requirement I{line}')
            receipt_days.add(float(row[4]))
            records.append({'known_day': known, 'day': (row[6] - ORIGIN).days,
                            'qty': float(row[8]) * factor, 'unit': canonical, 'row': line})
    finally:
        book.close()
    if receipt_days != {7.}:
        raise ValueError('693055 reception days are not consistently seven')
    book = openpyxl.load_workbook(inventory, read_only=True, data_only=True)
    try:
        lots = [(line, row) for line, row in enumerate(book['Taile de Lot'].iter_rows(min_row=2, values_only=True), 2)
                if str(row[0]).zfill(6) == '693055' and str(row[2]) == '1450']
    finally:
        book.close()
    if len(lots) != 1 or float(lots[0][1][4]) != 600000:
        raise ValueError('Expected source Gaillac fixed lot600000G')
    if any(s['item_id'] == 'item:693055' and s['node_id'] == 'M-1810' for s in meta['versioned_series']):
        raise ValueError('693055 already revised in reference')
    series = revision_series(records, vintages, '693055/1810', fraction,
                             planning_horizon_days=364, output_uom='G')
    fixed_before = len({(r['node_id'], r['item_id']) for r in meta['rows']})
    meta['rows'] = [r for r in meta['rows'] if (r['node_id'], r['item_id']) != ('M-1810', 'item:693055')]
    meta['versioned_series'].append(series)
    meta['assumptions'].append('693055 uses known source revisions with its unchanged first-plan residual share. Sixteen other fixed complementary calendars remain unchanged; this is not measured industrial consumption.')
    source_refs = {
        'transfer_multiple': f"Flow_Data_Inventory_and_Replenishment_rules.xlsx/Taile de Lot!E{lots[0][0]}: Gaillac600000G; using the same multiple for transfer proposals is a candidate supported by repeated Avene H quantities.",
        'receipt_days': 'Flow_Data_MRP_results.xlsx/Feuille1!E: seven for693055/1810; monday-Friday availability is an explicit candidate convention.',
        'safety': 'Twenty source working days retained; interpreted here as a dated stock floor, not a confirmed ERP anticipation rule.',
        'delivery': 'FIA70day reference and aggregate Gaillac28day boundary unchanged; delivery is not certified physical truck travel time.',
    }
    commands, inputs = {}, dict(old['protected_inputs'])
    for p in (graph_path, workbook, inventory):
        inputs[str(p)] = sha(p)
    for name, previous in [('nominal', 'nominal'), ('reference_figee', 'previsions_actualisees')]:
        command = list(old['commands'][previous])
        command[0] = sys.executable
        command[command.index('--output-dir') + 1] = str(args.workdir / name)
        commands[name] = command
    stages = ['previsions_693055', 'lots_693055', 'reception_693055', 'protection_693055', 'lot_1000']
    for index, name in enumerate(stages):
        variant = deepcopy(graph)
        if index >= 1:
            policy = {'policy_id': name, 'node_id': 'M-1810', 'item_id': 'item:693055',
                      'source_node_id': 'SDC-1450', 'uom': 'G', 'transfer_multiple_qty': 600000,
                      'source_refs': source_refs, 'status': 'candidate_not_inferred_ERP'}
            if index >= 2:
                policy.update(receipt_days=7, receipt_calendar='monday_friday')
            if index >= 3:
                policy['protection_mode'] = 'dated_stock_floor'
            variant['meta']['internal_component_policy'] = {'schema_version': 1, 'rows': [policy]}
        if index == 4:
            variant['meta']['procurement_batching'] = {'schema_version': 1, 'rows': [{
                'policy_id': '001757_round1000_no_fixed_cycle', 'node_id': 'M-1810',
                'item_id': 'item:001757', 'uom': 'KG', 'rounding_multiple_qty': 1000,
                'grouping_days': 0, 'semantics': 'first_uncovered_requirement_forward_window',
                'status': 'candidate_sensitivity',
                'evidence_note': '646/649 positive source H occurrences are multiples of1000KG; weekly aggregates do not prove individual orders. No fixed14/28day cycle imposed.',
            }]}
        graph_file = args.workdir / f'{name}_graph.json'
        write(graph_file, variant)
        inputs[str(graph_file)] = sha(graph_file)
        command = list(old['commands']['previsions_actualisees'])
        command[0] = sys.executable
        command[command.index('--output-dir') + 1] = str(args.workdir / name)
        command[command.index('--input') + 1] = str(graph_file)
        commands[name] = command
    labels = {'nominal': 'Nominal historique · témoin sur cinq ans',
              'reference_figee': 'Référence 52 semaines · anciennes conventions 693055',
              'previsions_693055': '693055 · prévisions actualisées seules',
              'lots_693055': '693055 · prévisions + pas de 600 kg',
              'reception_693055': '693055 · prévisions + 600 kg + réception 7 jours ouvrés',
              'protection_693055': '693055 · conventions précédentes + protection datée candidate',
              'lot_1000': 'Protection datée 693055 + achats 001757 par multiples de 1 000 kg'}
    plan = {'schema_version': 2, 'experiment': 'mrp_policy_693055_001757', 'reference': str(reference),
            'commands': commands, 'labels': labels, 'protected_inputs': inputs,
            'estimator_sha256': sha(Path(__file__)),
            'csv_witnesses': {'nominal': str(reference / 'nominal/data'),
                              'reference_figee': str(reference / 'previsions_actualisees/data')}}
    protected(plan)
    write(args.workdir / 'scope_audit.json', {'pair': '693055/1810', 'fraction': fraction,
          'source_records': records, 'versions': len(series['versions']), 'output_uom': 'G',
          'fixed_calendars_before': fixed_before, 'fixed_calendars_after': fixed_before - 1,
          'source_refs': source_refs, 'stages': stages,
          'initial_stocks_source_safety_and_other_calendars_unchanged': True,
          'nominal_not_promoted': True, 'observed_H_J_K_not_replayed': True})
    write(args.workdir / 'plan.json', plan)
    print(json.dumps({'pair': '693055/1810', 'versions': len(series['versions']),
                      'fixed_fraction': fraction, 'runs': len(commands)}), flush=True)


def prepare(args):
    reference = args.reference.resolve()
    old = json.loads((reference / 'plan.json').read_text(encoding='utf8'))
    graph_path = reference / 'source_forecasts_graph.json'
    graph = json.loads(graph_path.read_text(encoding='utf8'))
    workbook = ROOT / 'data/source/Flow_Data_MRP_results.xlsx'
    known, records = read_first_plan(workbook)
    rows, audit = estimate_first_plan(graph, known, records)
    if not rows:
        raise ValueError('No supported complementary requirements')
    graph['meta']['external_component_demands'] = {
        'schema_version': 1, 'origin': '2025-01-01', 'scenario_id': 'scn:BASE',
        'semantics': 'incremental_non_modelled_component_use', 'repeat_period_days': 365,
        'rows': rows, 'estimation_audit': audit,
        'assumptions': ['First January 5 MRP vintage frozen; no later source version informs this demand.',
                        'Current vintage bucket excluded: old backlog is not identified as new demand.',
                        'Projected component use is treated as a physical demand hypothesis, not observed consumption.',
                        'Missing weeks remain unknown; no additional demand is inferred there.',
                        'Aggregate residual share is not an identified list of other products.',
                        'Upstream residual deducts additional downstream use to avoid double counting.',
                        'Modeled finished products have physical allocation priority.',
                        'Years 2-5 repeat the first-year hypothetical complementary calendar.'],
    }
    graph_file = args.workdir / 'shared_components_graph.json'
    write(graph_file, graph)
    write(args.workdir / 'scope_audit.json', {'schema_version': 1, 'known_day': known,
          'screening_ratio': 1.25, 'pairs': audit, 'rows': len(rows)})
    commands = {}
    for name, reference_name in [('nominal', 'nominal'), ('reference_datee', 'planification_datee'), ('usages_complementaires', 'planification_datee')]:
        command = list(old['commands'][reference_name])
        command[0] = sys.executable
        command[command.index('--output-dir') + 1] = str(args.workdir / name)
        if name == 'usages_complementaires':
            command[command.index('--input') + 1] = str(graph_file)
        commands[name] = command
    inputs = {str(ROOT.parent / path): digest for path, digest in old['protected_inputs'].items()}
    inputs.update({str(p): sha(p) for p in [graph_path, graph_file, workbook]})
    plan = {'schema_version': 1, 'reference': str(reference), 'commands': commands,
            'protected_inputs': inputs, 'estimator_sha256': sha(Path(__file__)),
            'labels': {'nominal': 'Nominal conservé', 'reference_datee': 'MRP daté · sans usages complémentaires',
                       'usages_complementaires': 'MRP daté · usages complémentaires estimés'}}
    protected(plan)
    write(args.workdir / 'plan.json', plan)
    print(json.dumps({'pairs_examined': len(audit), 'pairs_selected': [r['pair'] for r in audit if r['status'] == 'estimated_complement'], 'rows': len(rows)}), flush=True)


def run(args):
    from etudecas import regenerate as regen
    from etudecas.testing.qualification import require_run_invariants
    plan = json.loads((args.workdir / 'plan.json').read_text(encoding='utf8'))
    protected(plan)
    for name in plan['commands']:
        if (args.workdir / name).exists():
            raise ValueError(f'Refuse overwrite: {name}')
    bundle = regen._capture_sources(args.workdir, list(plan['commands'].values()))
    results = []
    for name, command in plan['commands'].items():
        start = time.perf_counter()
        with (args.workdir / f'{name}.log').open('x', encoding='utf8') as log:
            subprocess.run(command, cwd=ROOT.parent, stdout=log, stderr=subprocess.STDOUT, check=True)
        elapsed = time.perf_counter() - start
        regen._record_run('nominal', args.workdir / name, command,
                          source_bundle=regen._check_sources(bundle), elapsed_seconds=elapsed)
        checks = require_run_invariants(args.workdir / name)
        write(args.workdir / name / 'invariants.json', checks)
        result = {'name': name, 'elapsed_seconds': elapsed, 'invariants_ok': checks['ok']}
        if name in plan.get('csv_witnesses', {}) or (plan.get('schema_version') == 1 and name in ('nominal', 'reference_datee')):
            prior = (Path(plan['csv_witnesses'][name]) if name in plan.get('csv_witnesses', {}) else
                     Path(plan['reference']) / ('nominal' if name == 'nominal' else 'planification_datee') / 'data')
            a = {p.name: sha(p) for p in prior.glob('*.csv')}
            b = {p.name: sha(p) for p in (args.workdir / name / 'data').glob('*.csv')}
            result['csv_identity'] = bool(a) and a == b
            result['changed_csv'] = sorted(k for k in a.keys() | b.keys() if a.get(k) != b.get(k))
            if not result['csv_identity']:
                write(args.workdir / f'{name}_regression.json', result)
                raise ValueError(f'Witness regression: {result}')
        results.append(result)
        protected(plan)
        print(json.dumps(result), flush=True)
    write(args.workdir / 'runs.json', {'ok': True, 'results': results})


def render(args):
    from etudecas.visualization.maps.source_comparison import build_comparison_payload
    from etudecas.visualization.maps.portable_diagnostic import build_comparison_map
    plan = json.loads((args.workdir / 'plan.json').read_text(encoding='utf8'))
    protected(plan)
    revised = plan.get('experiment') == 'weekly_revisions_001848'
    sourcing = plan.get('experiment') == 'sourcing_001848'
    rolling = plan.get('experiment') == 'rolling52_001757'
    policy = plan.get('experiment') == 'mrp_policy_693055_001757'
    batching = policy or rolling or plan.get('experiment') == 'batching_001757'
    runs = {name: args.workdir / name for name in plan['commands']}
    if revised:
        # A five-year witness is not a causal comparator for a one-year run:
        # end-of-horizon planning changes even with identical initial forecasts.
        runs = {name: args.workdir / name for name in ('nominal', 'usages_figes_2025', 'revisions_001848')}
        for name in ('usages_figes_2025', 'revisions_001848'):
            summary = json.loads((runs[name] / 'summaries/first_simulation_summary.json').read_text(encoding='utf8'))
            if summary['sim_days'] != 365:
                raise ValueError('The revised comparison requires two 365-day runs')
    if sourcing:
        for name in ('reference_revisions', 'regles_fournisseurs'):
            summary = json.loads((runs[name] / 'summaries/first_simulation_summary.json').read_text(encoding='utf8'))
            if summary['sim_days'] != 365:
                raise ValueError('The sourcing comparison requires two 365-day runs')
    if batching:
        for name, folder in runs.items():
            if name != 'nominal' and json.loads((folder / 'summaries/first_simulation_summary.json').read_text(encoding='utf8'))['sim_days'] != 365:
                raise ValueError('Every active batching comparator must cover 365 days')
        selected = args.compare_run or ('protection_693055' if policy else ('previsions_actualisees' if rolling else 'lot_1000'))
        if selected not in runs or selected in ('nominal', 'reference_figee'):
            raise ValueError('Select an explicit batching candidate or revised-forecast control')
    payload = build_comparison_payload(ROOT / 'data/source/Flow_Data_Inventory_and_Replenishment_rules.xlsx',
                                       ROOT / 'data/source/Flow_Data_MRP_results.xlsx',
                                       runs)
    for name, label in plan['labels'].items():
        if name in payload['runs']:
            payload['runs'][name]['label'] = label
    if revised:
        payload['runs']['nominal']['label'] = 'Nominal historique · calcul sur cinq ans (hors comparaison à horizon égal)'
        payload['runs']['usages_figes_2025']['label'] = 'Plan du 5 janvier figé · calcul sur 2025'
        payload['default_pair'] = '001848/1810'
        payload['default_runs'] = ['usages_figes_2025', 'revisions_001848']
    if sourcing:
        payload['default_pair'] = '001848/1810'
        payload['default_runs'] = ['reference_revisions', 'regles_fournisseurs']
    if batching:
        payload['default_pair'] = '693055/1810' if policy else '001757/1810'
        payload['default_runs'] = ['reference_figee', selected]
    payload['conventions'] += ([
        'Étude isolée : prévisions 693055 actualisées, puis pas de transfert 600 kg, réception 7 jours ouvrés, protection datée candidate et enfin arrondi 001757 à 1 000 kg. Chaque étape est comparée à la précédente ; aucune ne remplace automatiquement le nominal.',
        'Les 20 jours ouvrés sources de sécurité restent inchangés. Leur interprétation en plancher de stock est une hypothèse explicite ; l’autre interprétation possible en anticipation des besoins n’est pas certifiée par cette étude.',
        'Une cible de protection est un niveau à préserver, pas une consommation ni une sortie physique. Un manque de protection n’est pas nécessairement un besoin client non servi.',
        'Les prévisions complémentaires de 693055 utilisent la dernière version connue, avec la part estimée initiale conservée. Les 16 autres calendriers figés et les prévisions actualisées de 001757/001848 sont inchangés.',
        'Les quantités de 600 et 1 000 kg sont des pas candidats de réapprovisionnement étayés par les flux ; elles ne désignent pas des palettes ou des capacités de camion.',
        'Le délai de référence de 70 jours et la frontière agrégée Gaillac restent conservés. La décomposition industrielle entre fabrication, transport et réception reste partiellement inconnue.',
        'Chaque calcul comparé exécute 2025 avec un horizon glissant de 52 semaines. Les H/K industriels restent des projections de référence, jamais des mouvements exécutés injectés dans la simulation.',
        'Les autres panneaux et les deux suivis de lots restent ceux du nominal historique. La protection candidate et les nouveaux résultats sont présentés dans ce panneau comparatif.',
    ] if policy else [
        'Horizon glissant : le MRP regarde les 364 jours futurs à chaque décision, même en décembre. Les stocks et mouvements comparés restent ceux des 365 jours de 2025.',
        'La référence utilisait déjà les prévisions actualisées mais coupait les besoins au 31 décembre ; la correction conserve les semaines 2026 fournies par les plans connus à date.',
        'Les séries complémentaires 001757 et 001848 conservent les mêmes parts estimées ; elles incluent désormais leur projection source après décembre. Les 17 autres calendriers complémentaires restent dans leur périmètre précédent.',
        'Une semaine non fournie reste inconnue. Les prévisions futures ne sont pas des consommations déjà exécutées ; les versions successives se remplacent, sans addition.',
        'Les sécurités sources, les états initiaux, les commandes engagées, les règles fournisseurs et les paramètres des aléas sont conservés. Aucun inventaire ultérieur ou réception H ne force la trajectoire.',
        'Les variantes de regroupement restent des hypothèses ; comparer leur effet au scénario 52 semaines sans regroupement. Ce seul tirage ne constitue pas un classement statistique.',
        'Les produits finis utilisent leur dernière version source connue ; sur les périodes absentes, le profil nominal de secours du moteur reste conservé. Les autres panneaux et les deux suivis de lots restent historiques.',
    ] if rolling else [
        'Étude 001757/Avène : comparer séparément la mise à jour des prévisions, l’arrondi des achats et le regroupement des besoins sur 14 ou 28 jours.',
        'La référence fige les autres usages de 001757 au plan du 5 janvier. Les six variantes utilisent les mêmes versions de prévision, connues à la date de décision, avec la même part complémentaire initiale.',
        'Les arrondis de 1 000 kg et les fenêtres de regroupement sont des hypothèses testées ; la donnée FIA indique une quantité standard de 100 kg. Une entrée H hebdomadaire ne prouve pas un ordre individuel.',
        'La fenêtre commence au premier besoin non couvert, sans cadence fixe d’achat. Les échéances des besoins restent visibles ; les commandes déjà engagées sont déduites une seule fois.',
        'Aucun stock source ultérieur ni réception H n’est imposé à la simulation. Les sécurités, les stocks initiaux, les délais et les paramètres d’aléas sont conservés.',
        'La règle fournisseur expérimentale de 001848 et ses prévisions révisées sont identiques dans tous les comparateurs ; les 17 autres calendriers complémentaires restent figés.',
        'Les sept comparateurs actifs durent 365 jours. Le nominal sur cinq ans vérifie la non-régression ; les autres panneaux et les deux suivis de lots restent historiques.',
        'Les hypothèses ont été suggérées par les plans 2025 : la comparaison des deux semestres teste leur stabilité, pas une validation hors échantillon indépendante.',
    ] if batching else [
        'Essai fournisseur limité à 001848/Avène : principal VD0951020A, secours VD0519670A, rôles confirmés par l’utilisateur.',
        'Les prix, délais et quantités standard viennent des lignes FIA 3 et 4 de 268091.xlsx. Le déclenchement précis du secours reste une convention expérimentale à vérifier.',
        'Les autres matières conservent leurs règles : le croisement MRP, inventaires et carnet initial ne démontre pas une politique générale principal/secours.',
        'Les deux calculs comparés durent 365 jours et partagent les mêmes besoins, états initiaux, sécurités et aléas. Le nominal historique sur cinq ans sert de témoin de non-régression.',
        'Les commandes fermes sont conservées. Un secours anticipant un ferme tardif peut donc créer un surplus ultérieur, à distinguer d’un double comptage.',
        'Les prévisions industrielles H ne sont pas imposées comme des réceptions exécutées. Les versions MRP successives ne s’additionnent pas.',
        'Les autres panneaux et les deux suivis de lots restent ceux du nominal historique.',
    ] if sourcing else [
        'Essai 001848 à Avène : seul son calendrier complémentaire est révisé. Les 18 autres usages estimés restent figés.',
        'La part estimée des autres usages reste celle du 5 janvier ; aucun inventaire ultérieur ne sert à ajuster la consommation.',
        'Les versions MRP remplacent les prévisions futures à leur date de connaissance. Chaque semaine exécutée est figée depuis le dernier plan strictement antérieur à son début.',
        'La colonne I de la semaine courante peut inclure des reports : elle ne crée pas un nouveau besoin. Les retards physiques du modèle sont conservés séparément.',
        'Une semaine absente du dernier plan reste inconnue : aucune demande complémentaire n’y est injectée, sans reprise d’une ancienne ligne omise.',
        'Les commandes engagées, les stocks de départ et les jours de sécurité restent conservés. Les entrées prévues industrielles H ne sont pas imposées à la simulation.',
        'La comparaison active oppose deux calculs de 365 jours, figé et révisé : même horizon de décision. Les témoins sur cinq ans restent dans les résultats conservés et la carte précédente.',
        'Les autres panneaux et les deux suivis de lots restent ceux du nominal historique.',
    ] if revised else [
        'Usages complémentaires : hypothèse figée à partir du seul plan MRP du 5 janvier, sans utiliser les inventaires ultérieurs pour calculer les sorties.',
        'La semaine courante du premier plan est exclue : elle peut contenir des retards antérieurs. Les besoins complémentaires commencent au 12 janvier.',
        'Les besoins complémentaires sont estimés sur le périmètre des autres usages ; ils ne prouvent ni leur identité ni leur consommation exécutée dans les données industrielles.',
        'Le stock global, les sécurités et la nomenclature sont conservés. Les PF étudiés sont servis en priorité ; les besoins complémentaires non servis restent visibles.',
        'Les absences de semaines sources restent inconnues. Aucune consommation supplémentaire n’est inventée pour ces semaines.',
        'Les années 2 à 5 répètent le calendrier complémentaire estimé de la première année ; seule 2025 dispose d’inventaires de comparaison.',
        'Les autres onglets et les deux suivis de lots restent ceux du nominal historique ; ce panneau présente les trois nouveaux calculs.',
    ])
    write(args.workdir / 'comparison.json', payload)
    delivery = build_comparison_map(ROOT / 'resultats/02_carte_lots_recente.html', payload, args.html)
    write(args.workdir / 'delivery.json', delivery)
    impact = []
    before_name = 'reference_figee' if batching else ('reference_revisions' if sourcing else ('usages_figes_2025' if revised else 'reference_datee'))
    after_name = selected if batching else ('regles_fournisseurs' if sourcing else ('revisions_001848' if revised else 'usages_complementaires'))
    for pair, row in payload['pairs'].items():
        old = row['simulations'][before_name]['metrics'].get('previous_day_physical')
        new = row['simulations'][after_name]['metrics'].get('previous_day_physical')
        if old and new:
            impact.append({'pair': pair, 'unit': row['unit'], 'source_mean': old['observed_mean'],
                           'mae_before': old['mae'], 'mae_after': new['mae'], 'mae_change': new['mae'] - old['mae'],
                           'bias_before': old['bias'], 'bias_after': new['bias']})
    comparisons = ({name: row['metrics'].get('previous_day_physical')
                    for name, row in payload['pairs']['001757/1810']['simulations'].items()}
                   if batching else {})
    write(args.workdir / 'impact.json', {'stock_comparisons': impact,
                                        **({'all_001757_scenarios': comparisons, 'display_candidate': selected} if batching else {}),
                                        **({'all_693055_scenarios': {name: row['metrics'].get('previous_day_physical')
                                             for name, row in payload['pairs']['693055/1810']['simulations'].items()}} if policy else {})})
    print(json.dumps({'html': str(args.html), 'pairs_compared': len(impact)}), flush=True)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('phase', choices=['prepare', 'prepare-revisions', 'prepare-sourcing', 'prepare-batching', 'prepare-rolling52', 'prepare-policy', 'run', 'render'])
    parser.add_argument('--reference', type=Path, default=DEFAULT_REFERENCE)
    parser.add_argument('--workdir', type=Path, default=DEFAULT_WORK)
    parser.add_argument('--html', type=Path, default=ROOT / 'resultats/composants_partages_20260929/carte.html')
    parser.add_argument('--compare-run', help='Explicit default candidate for the 001757 batching chart')
    args = parser.parse_args(argv)
    args.workdir = args.workdir.resolve()
    args.workdir.relative_to(ROOT / 'artifacts/testing')
    args.workdir.mkdir(parents=True, exist_ok=True)
    args.html = args.html.resolve()
    args.html.relative_to(ROOT)
    globals()[args.phase.replace('-', '_')](args)


if __name__ == '__main__':
    main()
