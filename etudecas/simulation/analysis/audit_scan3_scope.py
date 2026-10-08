"""Read-only, reproducible reconciliation of Scan3 source workbooks.

No simulator, source rewriting, date restoration, browser or calibration is used.
Missing movement rows remain missing; totals mean sums of the rows supplied.
Run from the repository root with ``python -B -m
etudecas.simulation.analysis.audit_scan3_scope``. Only the final JSON is written.
"""

from __future__ import annotations

import argparse
from collections import defaultdict
from datetime import date, datetime, timedelta
from decimal import Decimal
import hashlib
import json
from pathlib import Path
import statistics

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / "etudecas/data/source"
OUTPUT = ROOT / "etudecas/artifacts/testing/scan3_scope_20261007/source_audit.json"
MOV = "Flow_Data_Inventory_movements.xlsx"
STK = "Flow_Data_Inventory_and_Replenishment_rules.xlsx"
MRP = "Flow_Data_MRP_results.xlsx"
FIRST = date(2024, 12, 29)
LAST = date(2025, 12, 21)
MARKERS = [FIRST + timedelta(weeks=i) for i in range(52)]
UNITS = {"G": ("kg", .001), "KG": ("kg", 1), "ZUN": ("UN", 1),
         "UN": ("UN", 1), "UN.": ("UN", 1), "M": ("m", 1)}
SITES = {"1810": "Avène", "1430": "Gien", "1450": "Gaillac", "1920": "Muret"}


def code(value):
    return str(value).strip().zfill(6)


def day(value):
    if isinstance(value, datetime):
        return value.date()
    return value if isinstance(value, date) else date.fromisoformat(str(value))


def sha256(path):
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def ref(name, sheet, cells):
    return f"etudecas/data/source/{name}!{sheet}!{cells}"


def clean(value):
    if isinstance(value, dict):
        return {str(k): clean(v) for k, v in value.items()}
    if isinstance(value, (list, tuple)):
        return [clean(v) for v in value]
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, float):
        return round(value, 9)
    return value


def stats(values):
    return {"n": len(values), "mean": statistics.mean(values),
            "min": min(values), "max": max(values),
            "relative_range_pct": 100 * (max(values) - min(values)) / statistics.mean(values)
            if statistics.mean(values) else None}


def read_sources():
    tables, provenance = {}, []
    requested = {MOV: ["Feuille1"], STK: ["Stocks"],
                 "268091.xlsx": ["BOM"], "268967.xlsx": ["BOM"], "773474.xlsx": ["BOM"],
                 "Extract_En_cours.xlsx": ["Sheet1"], "Data_poc.xlsx": ["BOM"], MRP: ["Feuille1"]}
    for name, sheets in requested.items():
        path = SOURCE / name
        before = sha256(path)
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in sheets:
            tables[name, sheet] = list(enumerate(workbook[sheet].iter_rows(values_only=True), 1))
        workbook.close()
        after = sha256(path)
        if before != after:
            raise RuntimeError(f"Source changed while reading: {name}")
        provenance.append({"path": path.relative_to(ROOT).as_posix(), "sha256": after,
                           "sheets": {s: len(tables[name, s]) for s in sheets}})
    return tables, provenance


def movement_rows(tables):
    movements = defaultdict(dict)
    for row, values in tables[MOV, "Feuille1"][1:]:
        article, label, typ, site, marker, raw_unit, *flows = values
        unit, factor = UNITS[raw_unit]
        key, marker = (code(article), str(site)), day(marker)
        if marker in movements[key]:
            raise ValueError(f"Duplicate movement article/site/week: {key}, {marker}")
        if any(v is None for v in flows):
            raise ValueError(f"Blank flow is not an explicit zero: row {row}")
        movements[key][marker] = {"row": row, "label": label, "article_type": typ,
                                "unit": unit, "raw_unit": raw_unit,
                                **{field: float(v) * factor for field, v in zip("GHIJ", flows)}}
    return movements


def stock_rows(tables):
    stocks = defaultdict(dict)
    for row, values in tables[STK, "Stocks"][1:]:
        article, typ, site, site_name, quantity, value, raw_unit, snapshot = values
        unit, factor = UNITS[raw_unit]
        key, snapshot = (code(article), str(site)), day(snapshot)
        if snapshot in stocks[key]:
            raise ValueError(f"Duplicate stock article/site/date: {key}, {snapshot}")
        stocks[key][snapshot] = {"quantity": float(quantity) * factor, "unit": unit,
                                "source": ref(STK, "Stocks", f"E{row}:H{row}")}
    return stocks


def period_balance(rows, photos, start, end):
    # First Jan-1 snapshot is compared with the entire boundary week, explicitly.
    first_marker = FIRST if start == date(2025, 1, 1) else start - timedelta(days=1)
    selected = {d: r for d, r in rows.items() if first_marker <= d < end - timedelta(days=1)}
    totals = {c: sum(r[c] for r in selected.values()) for c in "GHIJ"}
    opening, closing = photos.get(start), photos.get(end)
    projected = opening["quantity"] + sum(totals.values()) if opening else None
    return {"opening_date": start, "closing_date": end, "opening": opening, "closing": closing,
            "first_marker": first_marker, "last_marker": end - timedelta(days=8),
            "boundary_week_overlaps_2024": start == date(2025, 1, 1),
            "present_rows": len(selected), "signed_flows": totals,
            "computed_closing": projected,
            "residual_computed_minus_snapshot": projected - closing["quantity"]
            if projected is not None and closing else None,
            "flow_source_rows": [r["row"] for r in selected.values()]}


def balances(movements, stocks):
    result = []
    for key in sorted(set(movements) | set(stocks)):
        rows, photos = movements.get(key, {}), stocks.get(key, {})
        selected = {d: r for d, r in rows.items() if FIRST <= d <= LAST}
        annual = period_balance(rows, photos, date(2025, 1, 1), date(2025, 12, 29))
        common = period_balance(rows, photos, date(2025, 1, 6), date(2025, 12, 29))
        weekly_anomalies, missing_with_stock_change = [], []
        closed, missing_unchanged, checked = 0, 0, 0
        for start, opening in sorted(photos.items()):
            end = date(2025, 1, 6) if start == date(2025, 1, 1) else start + timedelta(weeks=1)
            if end not in photos:
                continue
            marker = FIRST if start == date(2025, 1, 1) else start - timedelta(days=1)
            record = rows.get(marker)
            change = photos[end]["quantity"] - opening["quantity"]
            if record is None:
                if abs(change) <= 1e-6:
                    missing_unchanged += 1
                else:
                    missing_with_stock_change.append({"start": start, "end": end, "stock_change": change})
                continue
            checked += 1
            residual = sum(record[c] for c in "GHIJ") - change
            if abs(residual) <= 1e-6:
                closed += 1
            else:
                weekly_anomalies.append({"start": start, "end": end, "marker": marker,
                    "signed_flows": {c: record[c] for c in "GHIJ"}, "stock_change": change,
                    "residual": residual, "opening": opening, "closing": photos[end],
                    "movement_source": ref(MOV, "Feuille1", f"G{record['row']}:J{record['row']}")})
        result.append({"article": key[0], "site": key[1], "site_name": SITES.get(key[1], key[1]),
            "label": next(iter(rows.values()))["label"] if rows else None,
            "unit": next(iter(photos.values()))["unit"] if photos else next(iter(rows.values()))["unit"],
            "coverage": {"expected_markers": 52, "present_markers": len(selected),
                "first_present_marker": min(selected) if selected else None,
                "last_present_marker": max(selected) if selected else None,
                "missing_markers": [d for d in MARKERS if d not in rows],
                "excluded_markers": [d for d in rows if d < FIRST or d > LAST],
                "explicit_zero_J_rows": sum(r["J"] == 0 for r in selected.values()),
                "positive_I_rows": sum(r["I"] > 0 for r in selected.values()),
                "positive_I_total": sum(r["I"] for r in selected.values() if r["I"] > 0),
                "negative_I_total": sum(r["I"] for r in selected.values() if r["I"] < 0),
                "stock_snapshot_count": len(photos)},
            "annual_boundary_balance": annual, "complete_weeks_balance": common,
            "weekly_check": {"present_intervals": checked, "closed_intervals": closed,
                "anomalies": weekly_anomalies,
                "missing_movement_unchanged_stock_intervals": missing_unchanged,
                "missing_movement_changed_stock_intervals": missing_with_stock_change},
            "movement_source": ref(MOV, "Feuille1", "A1:J1320")})
    return result


def bom_products(tables, movements):
    products = []
    for product, site in [("268091", "1810"), ("268967", "1430"), ("773474", "1450")]:
        components = []
        for row, values in tables[f"{product}.xlsx", "BOM"][1:]:
            if values[0] is None:
                continue
            if product == "773474":
                _, label, base, article, typ, quantity, raw_unit = values
                output_unit, output_factor = UNITS["G"]
            else:
                _, base, article, typ, quantity, raw_unit = values[:6]
                output_unit, output_factor = UNITS["UN"]
            unit, factor = UNITS[raw_unit]
            coefficient = float(quantity) * factor / (float(base) * output_factor)
            key = (code(article), site)
            rows = movements.get(key, {})
            selected = {d: r for d, r in rows.items() if FIRST <= d <= LAST}
            if any(r["unit"] != unit for r in selected.values()):
                raise ValueError(f"Incompatible BOM/movement units: {key}")
            components.append({"article": key[0], "type": typ, "unit": unit,
                "bom_quantity_raw": quantity, "bom_unit_raw": raw_unit,
                "bom_base_raw": base, "bom_output_unit_raw": "G" if product == "773474" else "UN",
                "coefficient_per_output_unit": coefficient, "output_unit": output_unit,
                "source": ref(f"{product}.xlsx", "BOM", f"A{row}:{'G' if product == '773474' else 'F'}{row}"),
                "I_signed": sum(r["I"] for r in selected.values()),
                "J_signed": sum(r["J"] for r in selected.values()),
                "J_consumption_equivalent": -sum(r["J"] for r in selected.values()) / coefficient,
                "I_opposite_signed_equivalent": -sum(r["I"] for r in selected.values()) / coefficient,
                "present_weeks": len(selected),
                "weekly": [{"marker": d, "week_start": d + timedelta(days=1),
                    "week_end": d + timedelta(days=7), "present": d in selected,
                    "I_signed": selected[d]["I"] if d in selected else None,
                    "J_signed": selected[d]["J"] if d in selected else None,
                    "J_equivalent": -selected[d]["J"] / coefficient if d in selected else None,
                    "source": ref(MOV, "Feuille1", f"I{selected[d]['row']}:J{selected[d]['row']}")
                    if d in selected else None} for d in MARKERS]})
        products.append({"product": product, "site": site, "output_unit": output_unit,
                         "components": components})
    return products


def group_and_lag(products):
    groups, pairs = [], []
    for product in products:
        by_article = {r["article"]: r for r in product["components"]}
        if product["product"] == "268091":
            articles = [r["article"] for r in product["components"] if r["type"] == "MP"]
            comparisons = [("049371", "338929"), ("049371", "338928"), ("338928", "338929")]
        elif product["product"] == "268967":
            articles = ["038005", "042342", "773474"]
            comparisons = [("773474", "333362"), ("042342", "333362"), ("038005", "333362")]
        else:
            continue
        series = {a: {v["marker"]: v["J_equivalent"] for v in by_article[a]["weekly"]
                      if v["present"]} for a in by_article}
        common = sorted(set.intersection(*(set(series[a]) for a in articles)))
        active = [d for d in common if any(series[a][d] for a in articles)]
        weekly_dispersion = [{"marker": d, **stats([series[a][d] for a in articles])} for d in active]
        groups.append({"product": product["product"], "articles": articles,
            "annual_equivalent_stats": stats([by_article[a]["J_consumption_equivalent"] for a in articles]),
            "all_articles_present_weeks": len(common), "active_common_weeks": len(active),
            "common_week_equivalent_totals": {a: sum(series[a][d] for d in common) for a in articles},
            "weekly_dispersion": weekly_dispersion})
        for upstream, downstream in comparisons:
            first, second = series[upstream], series[downstream]
            diagnostics = []
            for lag in range(-26, 27):
                aligned = [(d, first[d], second[d + timedelta(weeks=lag)]) for d in first
                           if d + timedelta(weeks=lag) in second]
                x, y = [v[1] for v in aligned], [v[2] for v in aligned]
                if not aligned:
                    continue
                correlation = statistics.correlation(x, y) if len(x) > 1 and len(set(x)) > 1 and len(set(y)) > 1 else None
                diagnostics.append({"lag_weeks": lag, "paired_present_weeks": len(aligned),
                    "upstream_total_equivalent": sum(x), "downstream_total_equivalent": sum(y),
                    "ratio": sum(x) / sum(y) if sum(y) else None, "pearson_correlation": correlation})
            eligible = [d for d in diagnostics if d["paired_present_weeks"] >= 20 and d["pearson_correlation"] is not None]
            full_upstream = by_article[upstream]["J_consumption_equivalent"]
            full_downstream = by_article[downstream]["J_consumption_equivalent"]
            quarterly = []
            for quarter in range(1, 5):
                dates = [d for d in set(first) & set(second) if (d + timedelta(days=1)).year == 2025
                         and ((d + timedelta(days=1)).month - 1) // 3 + 1 == quarter]
                a, b = sum(first[d] for d in dates), sum(second[d] for d in dates)
                quarterly.append({"quarter": quarter, "paired_present_weeks": len(dates),
                                  "upstream": a, "downstream": b, "ratio": a / b if b else None})
            pairs.append({"product": product["product"], "upstream_article": upstream,
                "downstream_article": downstream, "annual_upstream_equivalent": full_upstream,
                "annual_downstream_equivalent": full_downstream,
                "annual_ratio": full_upstream / full_downstream if full_downstream else None,
                "conditional_intermediate_stock_change_equivalent": full_upstream - full_downstream,
                "conditional_stock_note": "If identical exclusive scope, same BOM, no losses/other exits and complete flows; not an observed stock.",
                "quarterly_common_present_weeks": quarterly,
                "best_correlation_at_minimum_20_paired_weeks": max(eligible, key=lambda d: d["pearson_correlation"]) if eligible else None,
                "lag_diagnostics": diagnostics})
    return groups, pairs


def mrp_identity(tables):
    groups = defaultdict(list)
    for row, values in tables[MRP, "Feuille1"][1:]:
        groups[(day(values[0]), code(values[1]), str(values[3]), values[5])].append((row, values))
    largest, bad, later_j, gaps, total = Decimal(0), [], 0, 0, 0
    for key, rows in groups.items():
        previous, previous_date = Decimal(0), None
        for index, (row, values) in enumerate(sorted(rows, key=lambda rv: rv[1][6])):
            h, i, j, k = (Decimal(str(v)) for v in values[7:11])
            error = k - (previous + h + j - i)
            largest = max(largest, abs(error))
            total += 1
            if abs(error) > Decimal("0.01"):
                bad.append({"group": key, "row": row, "error_native_unit": str(error)})
            later_j += int(index > 0 and j != 0)
            gaps += int(previous_date is not None and (day(values[6]) - previous_date).days > 7)
            previous, previous_date = k, day(values[6])
    return {"formula": "K = previous_present_row_K + H + J - I; previous_K=0 for first group row",
        "grouping": ["date_MRP", "article", "division", "native_unit"], "rows": total,
        "groups": len(groups), "max_absolute_error_native_unit": str(largest),
        "tolerance_native_unit": "0.01", "outside_tolerance": bad,
        "nonzero_J_after_first_row": later_j, "gaps_longer_than_7_days": gaps,
        "source": ref(MRP, "Feuille1", "A1:K53399"),
        "limitation": "Identity on present rows only; gaps are not zero needs. J is a dated stock contribution, not proven physical delivery."}


def auxiliary(tables, products, movements):
    opening = []
    for row, values in tables["Extract_En_cours.xlsx", "Sheet1"][1:]:
        if code(values[0]) not in {"049371", "268091", "268967", "773474", "693055"}:
            continue
        unit, factor = UNITS[values[5]]
        opening.append({"article": code(values[0]), "planning_element": values[1], "site": str(values[2]),
            "quantity": float(values[4]) * factor, "unit": unit, "delivery_date": day(values[6]),
            "receipt_days_source": values[7], "entry_date": day(values[8]),
            "source": ref("Extract_En_cours.xlsx", "Sheet1", f"A{row}:I{row}")})
    poc = {}
    for row, values in tables["Data_poc.xlsx", "BOM"][1:]:
        if values[6] is not None:
            unit, factor = UNITS[values[11]]
            poc[(code(values[6]), code(values[9]))] = {"coefficient": float(values[10]) * factor / float(values[7]),
                "unit": unit, "source": ref("Data_poc.xlsx", "BOM", f"G{row}:L{row}")}
    bom_crosscheck = []
    for product in products[:2]:
        for component in product["components"]:
            other = poc.get((product["product"], component["article"]))
            bom_crosscheck.append({"product": product["product"], "article": component["article"],
                "detail_source": component["source"], "poc": other,
                "same_coefficient_and_unit": other is not None and other["unit"] == component["unit"]
                and abs(other["coefficient"] - component["coefficient_per_output_unit"]) < 1e-12})
    transfers = []
    for article in ["773474", "693055"]:
        receiving = "1430" if article == "773474" else "1810"
        out = {d: r for d, r in movements[article, "1450"].items() if FIRST <= d <= LAST}
        incoming = {d: r for d, r in movements[article, receiving].items() if FIRST <= d <= LAST}
        transfers.append({"article": article, "from_site": "1450", "to_site": receiving,
            "I_signed_origin": sum(r["I"] for r in out.values()),
            "J_signed_origin": sum(r["J"] for r in out.values()),
            "H_destination": sum(r["H"] for r in incoming.values()), "unit": "kg",
            "origin_rows": [r["row"] for r in out.values()], "destination_rows": [r["row"] for r in incoming.values()],
            "conclusion": "Equal annual quantities support a transfer candidate; no shipment/order IDs prove individual linkage."})
    return {"opening_orders": opening, "bom_data_poc_crosscheck": bom_crosscheck,
        "data_poc_unmatched_line": {"product": "268091", "article": "693710", **poc["268091", "693710"],
            "note": "Detailed BOM uses 007923 at same 3248 G per 1000 UN; no automatic identity substitution."},
        "transfer_candidates": transfers}


def build_audit():
    tables, sources = read_sources()
    movements, stocks = movement_rows(tables), stock_rows(tables)
    products = bom_products(tables, movements)
    groups, lags = group_and_lag(products)
    stock_balances = balances(movements, stocks)
    computed = [r for r in stock_balances if r["annual_boundary_balance"]["residual_computed_minus_snapshot"] is not None]
    report = {"schema_version": 1, "status": "source_audit_completed_with_unresolved_scope",
        "script": {"path": Path(__file__).relative_to(ROOT).as_posix(), "sha256": sha256(Path(__file__))},
        "sources": sources,
        "definitions": {"movement_I": tables[MOV, "Feuille1"][0][1][8],
            "movement_J": tables[MOV, "Feuille1"][0][1][9],
            "stock_balance": "opening + signed_G + signed_H + signed_I + signed_J - closing",
            "equivalent": "-signed_J / (normalized_BOM_component_quantity / normalized_BOM_output_base)",
            "units": "G converted once to kg; ZUN/UN./UN preserved as units; M preserved as m. Equivalent PF may be fractional and are not physical productions.",
            "week": "Sunday marker labels following Monday-Sunday; 2024-12-29 includes 2024-12-30 to 2025-01-05. 2025-12-28 marker excluded.",
            "missing": "Missing movement week is null, never an explicit zero. Totals sum available rows; unchanged stock supports only zero net change, not zero gross movements.",
            "lag": "Positive lag compares upstream week t with downstream week t+lag. Only pairs with both rows present; correlation cannot prove causality or matching quantities."},
        "coverage": {"first_marker": FIRST, "last_marker": LAST, "represented_start": date(2024, 12, 30),
            "represented_end": date(2025, 12, 28), "complete_2025_weeks_start": date(2025, 1, 6),
            "complete_2025_weeks_end": date(2025, 12, 28), "complete_2025_week_count": 51},
        "summary": {"article_site_count": len(stock_balances), "annual_balances_with_both_photos": len(computed),
            "annual_balances_closed_at_1e_6": sum(abs(r["annual_boundary_balance"]["residual_computed_minus_snapshot"]) <= 1e-6 for r in computed),
            "initial_photo_missing": [f"{r['article']}/{r['site']}" for r in stock_balances if r not in computed]},
        "article_site_balances": stock_balances, "bom_products": products,
        "group_consistency": groups, "lag_diagnostics": lags,
        "mrp_stock_identity": mrp_identity(tables), **auxiliary(tables, products, movements),
        "limitations": ["No measured bulk/intermediate stock register or complete OF-to-material-to-SKU genealogy is supplied.",
            "No other finished-format BOM is in the three detailed BOM workbooks; scope sharing remains an unproved hypothesis.",
            "No mass multiplier, I/J inversion, source correction or new simulation is authorized by these arithmetic findings.",
            "The initial photo dated 2025-01-01 and the first full source week have different calendar boundaries; exact daily 2025 allocation is unavailable.",
            "Source inconsistencies remain signed residuals, including 773474/Gaillac; no balancing correction is inserted.",
            "MRP plans and opening orders are commitments/plans, not execution or proof that material was already incorporated.",
            "Calibration of C8R and the industrial meaning of Scan3 are not certified."]}
    # Read-only stability check: no mutation probes, timestamp changes or restoration.
    for source in sources:
        if sha256(ROOT / source["path"]) != source["sha256"]:
            raise RuntimeError(f"Source changed during audit: {source['path']}")
    return clean(report)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args(argv)
    output = args.output.resolve()
    allowed = (ROOT / "etudecas/artifacts/testing").resolve()
    if not output.is_relative_to(allowed):
        parser.error("Output must remain below etudecas/artifacts/testing")
    report = build_audit()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(json.dumps({"status": report["status"], "report": str(output), "summary": report["summary"]}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
